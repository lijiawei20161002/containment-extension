"""Shared helpers and theme for the Medicare agent incident deck.

Pure python-pptx; no model calls. Content compiled from public sources.
"""
from __future__ import annotations
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x11, 0x11, 0x11)
BLUE = RGBColor(0x36, 0x5E, 0x9E)
BLUE_PALE = RGBColor(0xED, 0xF4, 0xFB)
BLUE_LIGHT = RGBColor(0xBD, 0xD8, 0xED)
GRAY = RGBColor(0x66, 0x66, 0x66)
LGRAY = RGBColor(0xF3, 0xF3, 0xF3)
LINE = RGBColor(0xC9, 0xC9, 0xC9)
RED = RGBColor(0xB3, 0x48, 0x31)
RED_PALE = RGBColor(0xFB, 0xEF, 0xEC)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)
GREEN_PALE = RGBColor(0xEA, 0xF5, 0xEF)
ORANGE = RGBColor(0xB8, 0x73, 0x2A)
ORANGE_PALE = RGBColor(0xFD, 0xF3, 0xE7)
PURPLE = RGBColor(0x6A, 0x4C, 0x93)
PURPLE_PALE = RGBColor(0xF1, 0xEC, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

SW, SH = Inches(13.333), Inches(7.5)
M = Inches(0.5)
FOOTER = ("Medicare Statistics Portal incident - reconstruction from public sources as of 16 Oct 2026 - "
          "no agent logs or ASD findings have been released")


def _run(p, text, size, color=INK, bold=False, italic=False):
    r = p.add_run(); r.text = text
    f = r.font; f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return r


def rich(p, text, size, color=INK, bold_all=False):
    for i, part in enumerate(text.split("**")):
        if part:
            _run(p, part, size, color, bold=bold_all or (i % 2 == 1))


def textbox(s, x, y, w, h, text="", size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, italic=False, ls=1.1):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        if "**" in ln and not bold:
            rich(p, ln, size, color)
        else:
            _run(p, ln, size, color, bold, italic)
    return tb


def bullets(s, x, y, w, h, items, size=13, color=INK, gap=4, ls=1.06):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = (item, 0)
        text, level = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = ls; p.space_after = Pt(gap)
        sz = size if level == 0 else size - 1
        marker = "-  " if level == 0 else "       -  "
        _run(p, marker, sz, color)
        rich(p, text, sz, color)
    return tb


def hrule(s, x, y, w, color=BLUE, weight=Pt(1.5)):
    """A horizontal rule drawn as a thin rectangle (connector-free; PowerPoint-safe)."""
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, weight)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def chevron(s, x, y, size, color=GRAY):
    """A small right-pointing triangle (connector-free)."""
    shp = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, size, size)
    shp.rotation = 90
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def rect(s, x, y, w, h, fill=WHITE, line=LINE, width=0.75, rounded=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        shp.adjustments[0] = 0.05
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(width)
    shp.shadow.inherit = False; shp.text_frame.text = ""
    return shp


def card(s, x, y, w, h, title, fill=BLUE_PALE, line=BLUE_LIGHT, tcolor=BLUE, items=None, tsize=13, bsize=11):
    rect(s, x, y, w, h, fill, line)
    textbox(s, x + Inches(0.08), y + Inches(0.05), w - Inches(0.16), Inches(0.35), title, tsize, tcolor, bold=True)
    if items:
        bullets(s, x + Inches(0.08), y + Inches(0.42), w - Inches(0.16), h - Inches(0.48), items, bsize, gap=3)


def table(s, x, y, w, h, rows, col_w, size=10, header_fill=BLUE, zebra=True, first_bold=True):
    t = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = cw
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03); cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            if i == 0:
                _run(p, val, size, WHITE, bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                rich(p, val, size, INK, bold_all=(first_bold and j == 0))
                cell.fill.solid(); cell.fill.fore_color.rgb = LGRAY if (zebra and i % 2 == 0) else WHITE
    return t
