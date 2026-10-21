"""Ten-slide research narrative. No hosted inference; toy calculations are exact.

Build: /Users/lijiawei/Desktop/.venv/bin/python slides/containment-decision-story/build_deck.py
Uses the repository's editable native-shape vocabulary without changing older decks.
"""
from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE.parent / 'decision-theory-deck'))
import deck_style as ds
from deck_style import (
    PRS, PAGES, SOURCES, DEEP, NAVY, DARKBOX, INK, MUTED, BLUE, TEAL, ORANGE,
    RED, PLUM, GOLD, BG, WHITE, PALE, MINT, SAND, ROSE, LILAC, CREAM, LINE, LIGHT,
    ML, MR, CW, text, box, line, dashed, arrow, circle, dot, badge, label, chip,
    card, panel, page, notes, link,
)
from pptx.util import Inches

OUT = HERE / 'containment-decision-story.pptx'
QA = HERE / 'qa'
QA.mkdir(exist_ok=True)
PRS.core_properties.title = 'When the task fails: a decision theory of agent boundaries'
PRS.core_properties.subject = 'Incidents, formal model, algorithms, experiments, and real-world implications'
PRS.core_properties.keywords = 'containment-extension; decision theory; authorization; research proposal'
SOURCES.clear()
SOURCES.update({
    'medicare': ('Australian Government briefing, 24 September 2026; incident in June',
        'https://ministers.finance.gov.au/financeminister/transcript/2026/09/24/press-conference-sydney'),
    'workbook': ('OpenAI, Unauthorized communication via temporary file hosting services, 16 September 2026; incident 14 April',
        'https://alignment.openai.com/misalignment-reports/unauthorized-communication-via-temporary-file-hosting-services/'),
    'wallet': ('OpenAI, Using a funded test cryptocurrency wallet without authorization, 15 October 2026; incident 26 June',
        'https://alignment.openai.com/misalignment-reports/using-a-funded-test-cryptocurrency-wallet-without-authorization/'),
    'pomdp': ('Kaelbling, Littman & Cassandra (1998), Planning and Acting in Partially Observable Stochastic Domains',
        'https://cs.brown.edu/courses/csci2951-k/papers/kaelbling98.pdf'),
    'voi': ('Hay, Russell, Tolpin & Shimony (2012), Selecting Computations: Theory and Applications',
        'https://arxiv.org/abs/1207.5879'),
    'cpo': ('Achiam et al. (2017), Constrained Policy Optimization',
        'https://arxiv.org/abs/1705.10528'),
    'camel': ('Debenedetti et al. (2025), Defeating Prompt Injections by Design (CaMeL)',
        'https://arxiv.org/abs/2503.18813'),
    'impossible': ('Zhong, Raghunathan & Carlini (2025), ImpossibleBench, sections 5-6',
        'https://arxiv.org/html/2510.20270v1'),
    'anthropic': ('Anthropic (9 September 2026), An alignment assessment of recent cybersecurity incidents',
        'https://www.anthropic.com/research/alignment-assessment-cybersecurity-incidents'),
    'pilot': ('Project: 160-run incident study, completed 17 September 2026',
        '../../docs/incident-study-results.md'),
    'design': ('Project: From reported incidents to testable questions',
        '../../docs/trust4rsi-incident-to-experiment.md'),
    'optimization': ('Project: Incident-aware optimization of agent control policies',
        '../../docs/incident-aware-optimization.md'),
    'next': ('Project: Next experiments, including silent and conflicting grants',
        '../../docs/trust4rsi-next-experiments.md'),
    'followup': ('Project: Live follow-up results; calibration and selection limits',
        '../../docs/trust4rsi-live-followup-results.md'),
})

def footer_links(s, refs):
    x = ML
    for key, caption, width in refs:
        link(s, x, 7.00, width, caption, SOURCES[key][1], MUTED, 9, .23)
        x += width + .16

def new(title, section, subtitle='', refs=(), dark=False):
    s = page(title, section, subtitle, source=' ', dark=dark, accent=TEAL)
    if refs:
        footer_links(s, refs)
    return s

def takeaway(s, value, color=TEAL, y=6.25):
    box(s, ML, y, CW, .60, color)
    text(s, ML+.22, y+.03, CW-.44, .54, value, 18, WHITE, True)

def block(s, x, y, w, h, title, body, color=TEAL, size=17):
    panel(s, x, y, w, h, color)
    text(s, x+.20, y+.18, w-.40, .43, title, 19, color, True)
    text(s, x+.20, y+.75, w-.40, h-.88, body, size, INK)

def para(s, x, y, w, h, value, size=17, color=INK, bold=False):
    return text(s, x, y, w, h, value, size, color, bold)

# Exact two-stage toy. The request is a contingent policy: ask, use the alternative
# only if granted, otherwise stop. The request is independent of local feasibility.
def toy_values(p, request_cost):
    ask = .6*10-request_cost
    return {
        'STOP': 0.,
        'RECOVER': 10*p-2,
        'ASK': ask,
        'PROBE': -.5 + p*max(8, ask, 0) + (1-p)*max(-2, ask, 0),
    }

def best(p, d):
    values = toy_values(p, d)
    # Prefer terminal actions on exact ties, so probes need positive net value.
    order = ('STOP', 'RECOVER', 'ASK', 'PROBE')
    return max(order, key=lambda key: values[key])

TOY_COLOR = {'STOP': MUTED, 'RECOVER': TEAL, 'ASK': BLUE, 'PROBE': GOLD}

# 01: A question with a visual fork, not a generic agenda slide.
s = new('', 'Containment-extension / Research outline / 21 October 2026', dark=True)
PAGES[-1]['title'] = 'When the task fails, what may the agent change?'
para(s, ML, 1.05, 6.5, 2.05, 'When the task fails,\nwhat may the\nagent change?', 40, WHITE, True)
box(s, ML, 3.37, 1.05, .06, TEAL, rounded=False)
para(s, ML, 3.66, 6.10, 1.10,
     'A decision theory of recovery, evidence,\nauthority and stopping.', 23, LIGHT)
para(s, ML, 5.15, 6.1, .82,
     'Real incidents → a formal problem → two algorithms\n→ experiments → implications for deployed agents', 16, LIGHT)
box(s, 7.25, 1.18, 5.42, .78, DARKBOX)
para(s, 7.46, 1.28, 4.95, .55, 'A useful task hits a blocked route.', 21, WHITE, True)
arrow(s, [(9.96, 1.97), (9.96, 2.46)], LIGHT)
box(s, 7.25, 2.52, 5.42, 1.40, TEAL)
para(s, 7.50, 2.68, 4.90, .43, 'Change the plan', 25, WHITE, True)
para(s, 7.50, 3.20, 4.90, .43, 'Recover · inspect · ask · report a blocker', 16, WHITE)
box(s, 7.25, 4.16, 5.42, 1.40, DARKBOX)
para(s, 7.50, 4.33, 4.90, .43, 'Change the permission?', 25, WHITE, True)
para(s, 7.50, 4.86, 4.90, .43, 'Only a trusted authority can grant it.', 17, LIGHT)
para(s, ML, 6.38, CW, .43, 'Research question: can agents stay useful while respecting this distinction?', 23, WHITE, True)
notes(s, 'Opening prompt: a research agent gets an error while retrieving useful data. Ask the audience what the next move should be. '
      'The talk treats containment-extension as a study of delegated action boundaries, not a claim to solve arbitrary sandbox escape. '
      'The model and algorithms are proposed. Slide 4 is an exact, explicitly hypothetical calculation. Slide 9 is archived hosted-model evidence. '
      'No new hosted-model experiments were run to create this deck. The new research bet is that an explicit model of recovery, '
      'information and grant transitions yields testable behavior and better held-out safety-utility tradeoffs.', ('design',))

# 02: Only source-supported incident facts; interpretation gets its own band.
s = new('Three incidents, one decision point to study.', '01 / Motivation from reported incidents',
        'A legitimate objective met a resource boundary. The next action crossed it.',
        [('medicare', 'Government briefing · 24 Sep 2026', 3.75),
         ('workbook', 'OpenAI workbook report · 16 Sep 2026', 3.90),
         ('wallet', 'OpenAI wallet report · 15 Oct 2026', 3.80)])
cases = [
    ('01', 'Retrieve statistics', 'Medicare portal / June 2026',
     'Request denied', 'Unauthorized portal access',
     'Aggregate statistics; the government reports no individual medical data accessed.', BLUE),
    ('02', 'Share a workbook', 'Training collaboration / April 2026',
     'Local transfer failed', 'Workbook put on public hosting',
     'The task specified local files. A working transfer did not satisfy that boundary.', TEAL),
    ('03', 'Deploy contracts', 'Test-network wallet / June 2026',
     'No supplied funded wallet', 'A found account paid deployment fees',
     'The model treated public keys and test tokens as grounds to spend without a grant.', ORANGE),
]
for i, (n, title, date, obstacle, effect, detail, col) in enumerate(cases):
    x = ML + i*4.10
    card(s, x, 2.45, 3.89, 3.38, col)
    badge(s, x+.22, 2.63, n, col, .44)
    para(s, x+.82, 2.59, 2.83, .50, title, 20, col, True)
    para(s, x+.22, 3.12, 3.42, .35, date, 12, MUTED)
    para(s, x+.22, 3.65, 3.42, .50, obstacle, 19, INK, True)
    arrow(s, [(x+.42, 4.18), (x+.42, 4.51)], col)
    para(s, x+.22, 4.50, 3.42, .68, effect, 19, RED, True)
    para(s, x+.22, 5.25, 3.42, .42, detail, 12, MUTED)
takeaway(s, 'Our hypothesis to test: technical failure changes the plan, but does not expand authority.')
notes(s, 'These are three reported cases, not a representative incident sample or observations generated by our fixture. '
      'Medicare: the government described a denied request followed by unauthorized access, and specifically said no individual medical data was accessed. '
      'Workbook: the local transfer path failed; the report documents public hosting despite a local-only task. '
      'Wallet: a deployment succeeded while charging fees to an account the task did not supply. '
      'We do not infer that every legitimate route was impossible, or that one causal mechanism explains all three. '
      'The common decision-point interpretation is ours; the reports supply motivation, not a randomized counterfactual.',
      ('medicare', 'workbook', 'wallet'))

# 03: Separate epistemic uncertainty from authority; avoid a soft penalty licensing violations.
s = new('Model recovery as constrained sequential choice.', '02 / Formal problem',
        'The agent is uncertain about feasibility. Its authority comes from a trusted task contract.',
        [('pomdp', 'Belief-state planning: Kaelbling et al. 1998', 4.40),
         ('cpo', 'Reward + constraints: Achiam et al. 2017', 4.35)])
for x, title, body, col in [
    (ML, 'Belief b_{t}', 'What routes might still work?\nUpdated from observations.', BLUE),
    (4.74, 'Grant g_{t}', 'Allowed resources and effects.\nChanged by trusted authority.', TEAL),
    (8.86, 'Budget k_{t}', 'How much time and tool use remain?\nReduced by each decision.', GOLD),
]:
    block(s, x, 2.40, 3.85, 1.48, title, body, col, 14)
box(s, ML, 4.08, CW, 1.20, NAVY)
para(s, ML+.25, 4.18, CW-.50, .44,
     'Choose π to maximize  E[u(outcome) − Σ cost(a_{t})]', 26, WHITE, True)
para(s, ML+.25, 4.73, CW-.50, .42,
     'subject to  a_{t} in A(g_{t})   and   total resource use ≤ B', 22, WHITE)
para(s, ML, 5.49, CW, .44,
     'Actions: execute an allowed plan  /  probe  /  request a grant  /  stop with evidence', 19, INK, True)
takeaway(s, 'A failed tool call updates b_{t}. It does not update g_{t}.')
notes(s, 'Finite-horizon constrained POMDP formulation. Let x_t be the hidden environment state (including route feasibility), '
      'g_t the trusted, versioned authorization contract, k_t the remaining resource budget, and h_t the observation/action history. '
      'The controller state is z_t=(b_t,g_t,k_t), with b_t=P(x_t|h_t). Transition and observation kernels are T and O; '
      'Bayesian filtering yields b_(t+1). A request action can lead to a trusted grant update or denial; q is a response probability, '
      'not the probability that an ungranted action is secretly authorized. Mere reachability and tool success never update g_t. '
      'A(g) includes permitted execution, permitted diagnostics, permitted requests and termination. Missing scope is not filled by the agent: '
      'choose an already permitted alternative, request clarification, or stop. The closed fixture can define exact effect-based permissions; '
      'an open environment requires validated effect prediction and reference policy. Utility u must be fixed in advance, including the '
      'values of valid completion, useful partial output and justified termination; c is utility cost and resource use is a separate κ. '
      'For an implemented policy measure I(τ)=1 if any executed effect violates the grant in force at that time, R(π)=P(I=1), and '
      'U(π)=P(successful completion with no violation). The ideal action constraint is a specification, not a guarantee that an LLM follows it. '
      'This uses established decision theory; CPO is background for constrained optimization, not the algorithm implemented or a source '
      'of guarantees for this proposal.', ('pomdp', 'cpo', 'design'))

# 04: Policy phase diagram, computed from exact payoffs, with an easy calculation.
s = new('Sometimes the best next action is to learn.', '03 / A worked decision problem',
        'Hypothetical utility units; one optional diagnostic, followed by one final commitment.',
        [('voi', 'Value-of-information background: Hay et al. 2012', 5.20)])
px, py, pw, ph = 1.10, 2.92, 5.25, 2.72
para(s, .78, 2.40, 5.9, .35, 'Optimal action as belief and request cost change', 16, NAVY, True)
para(s, 1.10, 2.74, 5.9, .25, 'Request cost d', 10, MUTED)
N = 100
for row in range(64):
    d = 8*(1-(row+.5)/64)
    current, start = None, 0
    for col in range(N+1):
        action = best((col+.5)/N, d) if col < N else None
        if action != current:
            if current is not None:
                box(s, px+pw*start/N, py+ph*row/64, pw*(col-start)/N+.002,
                    ph/64+.002, TOY_COLOR[current], rounded=False)
            current, start = action, col
line(s, px, py+ph, px+pw, py+ph, INK)
line(s, px, py, px, py+ph, INK)
for p in (0, .5, 1):
    para(s, px+pw*p-.20, py+ph+.04, .40, .25, str(p), 11, MUTED)
for d in (0, 4, 8):
    para(s, px-.43, py+ph*(1-d/8)-.15, .32, .30, str(d), 11, MUTED)
para(s, px+.65, py+ph+.32, pw-1, .27, 'Belief p that local recovery works', 13, MUTED)
for x, y, w, value, size in [(1.105, 3.18, .32, 'stop', 8), (1.20, 4.87, .55, 'ASK', 14),
                         (2.50, 3.70, 1.6, 'PROBE', 16), (5.18, 3.18, 1.10, 'RECOVER', 12)]:
    para(s, x, y, w, .35, value, size, WHITE, True)
# Highlight the actual worked case.
hx, hy = px+pw*.3, py+ph*(1-2/8)
circle(s, hx-.075, hy-.075, .15, WHITE, NAVY)
para(s, hx+.12, hy-.05, 1.80, .30, 'p=.3, d=2', 12, WHITE, True)
card(s, 6.90, 2.44, 5.81, 3.57, TEAL)
para(s, 7.15, 2.57, 5.30, .44, 'At p = 0.30 and d = 2', 22, TEAL, True)
rows = [
    ('Recover', '10p − 2', '1.0'),
    ('Ask, then obey the reply', '0.6 × 10 − d', '4.0'),
    ('Stop', '0', '0.0'),
    ('Probe, then decide', '−0.5 + 0.3×8 + 0.7×4', '4.7'),
]
for i, (a, formula, val) in enumerate(rows):
    y = 3.21+i*.58
    if i==3:
        box(s, 7.10, y-.04, 5.37, .55, MINT)
    para(s, 7.19, y, 2.26, .42, a, 14, INK, i==3)
    para(s, 9.46, y, 2.30, .42, formula, 12, MUTED)
    para(s, 11.91, y, .48, .42, val, 18, TEAL if i==3 else INK, True)
para(s, 7.17, 5.45, 5.18, .49, 'Success=10; recovery cost=2; grant chance=.6.\nPerfect probe costs .5; it changes no permission.', 11.5, MUTED)
takeaway(s, 'Cheap, informative checks can reduce premature stopping and unnecessary requests.')
notes(s, 'Exact toy calculation, not measured model performance. Hidden feasibility θ is Bernoulli(p). '
      'Direct authorized recovery yields +10 on success, 0 on failure, and costs 2 in either case, so Q_recover=10p−2. '
      'A request costs d, independently has probability q=0.6 of obtaining a valid grant for an alternative route, and '
      'then yields 10 if granted or 0 if denied. It includes compliance with the response; no alternative is executed without a grant. '
      'Direct recovery and request are terminal commitments; failure or denial ends the episode. Stop pays 0. '
      'One permitted, non-mutating, perfect diagnostic costs 0.5, reveals θ, and leaves the final commitment available. '
      'Q_probe=−0.5+p max(8,6−d,0)+(1−p)max(−2,6−d,0). '
      'At p=.3,d=2, it is optimal to diagnose, recover if feasible, and request if infeasible: value 4.7. '
      'Other checks: p=.05,d=2 gives ASK; p=.95,d=2 gives RECOVER; p=.01,d=7 gives STOP. '
      'The chart evaluates these formulas over a grid; its boundaries are hypotheses for behavior under this model. '
      'A richer horizon, noisy signals, correlated grant decisions or different utilities will change the regions. '
      'Hay et al. supplies the metalevel value-of-information background, not these invented task payoffs.', ('voi',))

# 05: A concrete algorithm and a precise separation of planning / enforcement / observation.
s = new('Algorithm A: plan within the grant, one step at a time.', '04 / Online decision algorithm / Proposed',
        'Start with exact dynamic programming in the finite fixture; approximate only after that reference works.',
        [('pomdp', 'Belief-state planning: Kaelbling et al. 1998', 4.40),
         ('camel', 'Capability enforcement precedent: CaMeL 2025', 4.70)])
stages = [
    ('1', 'Read the contract', 'Resource, owner, recipient,\noperation and grant version.', BLUE),
    ('2', 'Build choices', 'Recover, inspect, request,\nor report a blocker.', TEAL),
    ('3', 'Compare futures', 'Estimate outcomes/costs;\nlook ahead over evidence.', GOLD),
    ('4', 'Execute + update', 'Broker checks the call;\njournal records its effects.', PLUM),
]
for i, (n, head, body, col) in enumerate(stages):
    x = ML+i*3.08
    card(s, x, 2.48, 2.86, 1.92, col)
    badge(s, x+.18, 2.65, n, col, .39, 12)
    para(s, x+.18, 3.13, 2.50, .38, head, 18, col, True)
    para(s, x+.18, 3.61, 2.50, .60, body, 14, INK)
    if i<3:
        arrow(s, [(x+2.88, 3.40), (x+3.06, 3.40)], MUTED)
box(s, ML, 4.67, 7.32, 1.20, NAVY)
para(s, ML+.22, 4.80, 6.90, .43, 'V(z) = max  Q(z, a),   a in A(g)', 25, WHITE, True)
para(s, ML+.22, 5.31, 6.90, .31, 'Q = immediate net value + expected continuation value', 15, LIGHT)
para(s, 8.25, 4.72, 4.41, 1.13,
     'Budgeted lookahead for larger tasks.\nCalibrate estimates on development data.\nRequest when a required grant is missing.', 16, INK)
takeaway(s, 'Plan for utility; enforce scope before execution; measure effects independently.')
notes(s, 'Proposed online algorithm. Input: trusted contract g, task evidence h, transition/observation model, utility definition, and remaining budget. '
      'Normalize each candidate tool action into effect-relevant fields; use an existing grant checker to form the candidate set. '
      'Always include justified termination and available authorized requests. For finite fixtures, calculate '
      'V_h(b,g,k)=max_{a in A(g),κ(a)≤k} E[r(x,a)−c(a)+V_(h−1)(b′,g′,k−κ(a))], with explicit terminal payoffs. '
      'Observed grant transitions may occur only through trusted responses. For larger action spaces use a fixed-budget beam search '
      'or Monte Carlo lookahead; compare to a one-step VOI ablation. Execute one step, record actual effects, update belief and repeat. '
      'Learn transition/observation estimates from development data and examine calibration and sensitivity; an LLM confidence statement '
      'is not a calibrated probability. The unbrokered planner arm measures behavioral compliance. The otherwise identical brokered arm '
      'checks the action before dispatch and measures effects separately from attempted violations. Full prevention in the toy requires '
      'complete mediation and a correct effect/authorization model; the effect journal is post-action evidence, not prevention of that action. '
      'A simple broker allowlist is an engineering baseline, not a CaMeL reproduction. The proposed planner has not been evaluated on LLM agents.',
      ('pomdp', 'camel', 'design'))

# 06: An identifiable offline selection experiment, not an optimized vs bare-system comparison.
s = new('Algorithm B: select policies using actual effects.', '05 / Offline policy selection / Proposed',
        'Does better incident evidence change which controller we select, and improve its behavior on new tasks?',
        [('optimization', 'Project: incident-aware policy optimization', 4.50),
         ('impossible', 'Closest task/monitor overlap: ImpossibleBench §§5–6', 5.30)])
box(s, ML, 2.43, 3.14, 2.36, NAVY)
para(s, ML+.22, 2.62, 2.70, .56, 'Fixed policy pool', 21, WHITE, True)
para(s, ML+.22, 3.39, 2.70, 1.12, 'Same model and tasks.\nSame tools and budgets.\nShared development runs.', 16, LIGHT)
for i, (title, sub, col) in enumerate([
    ('T', 'Transcript judge', BLUE),
    ('T + J', 'Same judge + effect journal', TEAL),
    ('J*', 'Exact local effect oracle', PLUM),
]):
    y = 2.43+i*.88
    arrow(s, [(3.78, 3.61), (4.06, 3.61), (4.06, y+.34), (4.40, y+.34)], col)
    box(s, 4.43, y, 3.56, .70, {BLUE:PALE,TEAL:MINT,PLUM:LILAC}[col])
    para(s, 4.60, y+.10, .98, .46, title, 19, col, True)
    para(s, 5.65, y+.10, 2.15, .46, sub, 13, INK)
arrow(s, [(8.04, 3.64), (8.60, 3.64)], MUTED)
card(s, 8.66, 2.43, 4.05, 2.36, TEAL)
para(s, 8.89, 2.64, 3.59, .54, 'Select → freeze → test', 21, TEAL, True)
para(s, 8.89, 3.42, 3.59, 1.02, 'Select the lowest incident estimate\namong utility-eligible candidates.\nFresh tasks; one common observer.', 16, INK)
para(s, ML, 5.16, CW, .46,
     'Eligibility: authorized utility ≥ baseline − δ; fixed cost and review limits.', 21, NAVY, True)
para(s, ML, 5.73, CW, .34,
     'T → T+J tests evidence access. T+J → oracle tests the labeling method. A tie is a valid result.', 16, MUTED)
takeaway(s, 'Useful evidence must improve a held-out decision, not merely produce a better-looking label.')
notes(s, 'Algorithm B is an offline selection procedure. Fix K candidate controllers, a development set, actor model, all rollout budgets, '
      'utility definition and tie-breaks. Reuse the same trajectories across selection-signal arms. '
      'For each candidate compute common authorized utility and cost, plus an incident estimate from the selected evidence arm. '
      'Among candidates meeting the same development utility margin and cost/review limits, choose the lowest estimated incident rate '
      '(or a prespecified conservative bound); if none qualify, retain the reference and make no improvement claim. '
      'Freeze selected candidates before untouched test tasks. T is an actor-visible transcript; J is an independent service-effect journal. '
      'The T+J judge uses the same model/rubric as T, with extra evidence. The oracle uses explicit local grant and effect semantics and is '
      'a privileged reference, not a real-world general truth oracle. The common eligibility calculation is held fixed, so the contrast '
      'specifically tests the incident ranking signal. Match total search/labeling budget, model access and candidate trials; log unused resources. '
      'Use disjoint development, calibration and final test lineages. Freeze a repaired judge after development calibration; the prior failed '
      'calibration in this repository is not a qualified gate. Repeat complete selections to assess optimizer variability. '
      'Current feedback-selection gains remain unmeasured. ImpossibleBench already studies prompting, abort options and monitoring; '
      'this proposal must establish an incremental empirical contribution rather than claim those components as novel.',
      ('optimization', 'impossible', 'followup'))

# 07: The experiment that decouples difficulty from permission.
s = new('Same reachable resource. Different permission.', '06 / The discriminating experiment / Proposed',
        'Cross repair availability with trusted authorization; keep the alternate resource technically accessible.',
        [('design', 'Project: repair × authorization design', 4.00),
         ('next', 'Extension: silent and conflicting grants', 4.20)])
gx, gy, gw, gh = 2.45, 3.12, 4.96, 1.12
para(s, gx, 2.50, gw, .43, 'ALTERNATE FORBIDDEN', 17, RED, True)
para(s, gx+gw+.22, 2.50, gw, .43, 'ALTERNATE EXPLICITLY GRANTED', 17, TEAL, True)
para(s, ML, gy+.22, 1.62, .65, 'Local repair\nworks', 18, INK, True)
para(s, ML, gy+gh+.45, 1.62, .65, 'Local repair\nfails', 18, INK, True)
matrix = [
    (0,0,'Recover locally', 'An authorized route still exists.', MINT, TEAL),
    (1,0,'Use either permitted route', 'The alternative is legitimate here.', PALE, BLUE),
    (0,1,'Verify the blocker, then stop', 'Closed fixture: no allowed solution remains.', ROSE, RED),
    (1,1,'Use the granted alternative', 'Local failure does not make the task impossible.', MINT, TEAL),
]
for col,row,title,sub,fill,accent in matrix:
    x,y=gx+col*(gw+.22),gy+row*(gh+.22)
    box(s,x,y,gw,gh,fill)
    para(s,x+.20,y+.15,gw-.4,.42,title,20,accent,True)
    para(s,x+.20,y+.67,gw-.4,.28,sub,13,MUTED)
para(s, ML, 5.86, CW, .31,
     'Follow-up: silent scope → clarification/access rates; conflicting scope → a frozen authority rule.', 15, MUTED)
takeaway(s, 'When permission changes, the next action should change even if access does not.')
notes(s, 'The matrix is a normative reference, not observed results. Pair each underlying task instance across local-repair available/unavailable '
      'and alternate forbidden/explicitly granted. Keep contents, endpoint reachability, action schemas and budget fixed. '
      'Recoverable and unavailable variants share the initial failure history; an allowed probe reveals availability. '
      'Trusted permission is the treatment, not a claim inside tool output. In the lower-right cell the local route fails but the task is feasible. '
      'For this first factorial evaluation, grants are fixed and no new permissions arrive during the episode; this isolates compliance. '
      'A separate dynamic-request study introduces grant/deny responses and costs. The matrix states allowable behavior, not a universal '
      'requirement to finish when time cost exceeds reward. Utility-matched experimental instances should make completion worthwhile. '
      'The lower-left justified-impossibility label relies on a closed action set and checked environment specification; in open-world tasks '
      'the agent should report bounded uncertainty rather than assert no solution exists. Silent scope needs a reference authorization policy '
      'before incident rates can be assigned. Without one, report reads, requests and responses rather than manufacture violation labels. '
      'Conflicting scope requires an authority precedence rule frozen before collection.', ('design', 'next'))

# 08: Experiments, ablations, and decisive failure criteria.
s = new('Three tests that can prove the proposal wrong.', '07 / Experiments and falsifiable predictions / Proposed',
        'Measure better decisions, preserved capability and executed effects under the same resource limits.',
        [('design', 'Factorial task design', 2.50), ('optimization', 'Matched selection protocol', 3.10),
         ('voi', 'Decision-value background', 3.10)])
tests = [
    ('E1', 'Does the decision model predict choices?',
     'Sweep feasibility, diagnostic accuracy,\nrequest delay and grant probability.',
     'Exact DP vs one-step VOI vs LLM planner.\nMeasure value regret, calibration and requests.', BLUE),
    ('E2', 'Does permission govern the next action?',
     'Run the 2×2 design in service recovery\nand artifact-transfer tasks.',
     'Original / fixed scope / stop-or-ask /\nplanner / planner + broker.', TEAL),
    ('E3', 'Does incident feedback select better?',
     'Freeze T-, T+J- and oracle-selected\npolicies; evaluate on fresh task lineages.',
     'Require lower incident risk with authorized\nutility within a preregistered margin.', PLUM),
]
for i,(n,title,body,method,col) in enumerate(tests):
    y=2.41+i*1.11
    badge(s,ML,y+.17,n,col,.52,14)
    para(s,1.34,y,5.55,.40,title,19,col,True)
    para(s,1.34,y+.43,5.55,.59,body,14,INK)
    box(s,7.24,y,5.47,.99,{BLUE:PALE,TEAL:MINT,PLUM:LILAC}[col])
    para(s,7.45,y+.13,5.05,.72,method,15,INK)
para(s,ML,5.90,CW,.27,
     'Pilot: 480 rollouts/model; pair by lineage; report utility, effects, attempts, false stops, unknowns and cost.',14,MUTED)
takeaway(s,'Falsifiers: only the broker helps; refusals explain the gain; feedback selects the same policy.')
notes(s, 'All studies are proposed. E1: run a small finite simulator with known kernels, then noisy and misspecified variants; '
      'the exact finite-horizon DP is the normative reference, not a clairvoyant policy. Compare one-step VOI, budgeted lookahead and '
      'LLM planner choices using expected value regret V*−Vπ. Sweep p, diagnostic accuracy, diagnostic cost, request latency/cost and q. '
      'Evaluate calibration using proper scoring rules and reliability plots. An event log must not reveal hidden feasibility before the permitted probe. '
      'E2: service recovery and local artifact transfer, five arms as on the slide, matched models/tools/caps. '
      'An initial precision pilot could use 2 repair × 2 grant × 2 families × 5 arms × 12 fresh task instances per family = 480 rollouts per model. '
      'Twelve seeds of one template are not twelve independent task structures: use distinct task lineages and report template diversity. '
      'This is a proposed pilot allocation, not a powered confirmatory claim or an executed budget. Use its paired discordances to choose '
      'confirmatory sample size for a prespecified useful risk reduction and utility noninferiority margin, before collecting a new test cohort. '
      'Estimate permission effects and the repair×permission interaction within each policy. Randomize order; cluster paired uncertainty by lineage. '
      'Include a disjoint task family and evaluate each model separately for transfer. Check original-task competence before any coding extension '
      'such as ImpossibleBench. E3: shared development pool, fresh final tasks and repeated full selections; measure selected-policy differences '
      'in risk, utility, coverage, review burden and search cost. Primary comparison T+J vs T; oracle is diagnostic. '
      'For every assigned run report any unauthorized executed effect, attempted violations, authorized completion, false stopping, probes, '
      'requests, cost and unknown status. Report bounds for incomplete outcomes and do not count them as safe. '
      'Success requires a confidence interval supporting lower risk and retained utility, with cost/coverage constraints satisfied. '
      'No reduction, no feasible policy, identical selections, or unacceptable utility loss are valid negative results.',
      ('design','optimization','voi','impossible'))

# 09: Ground the proposal in actual archive data and clearly delimit its claims.
summary_path=ROOT/'results/incident-study-v1/summary.json'
historical=json.loads(summary_path.read_text())
counts=Counter((r['stop']=='scorer_success',bool(r['any_effect'])) for r in historical['records'])
assert len(historical['records'])==160
assert (counts[True,False],counts[True,True],counts[False,False],counts[False,True])==(96,13,49,2)
s = new('The pilot establishes a measurement problem.', '08 / Evidence we already have',
        '160 archived hosted-model continuations in a local service fixture; multiple models and instruction arms.',
        [('pilot','Archived study + trace evidence · September 2026',5.10),
         ('followup','Follow-up limits: calibration and selection',4.80)])
para(s,ML,2.45,6.15,.36,'Task score and authorization tell different stories',18,NAVY,True)
tx,ty,tw,th=2.26,3.25,2.12,.91
para(s,tx,2.91,tw,.28,'NO INCIDENT',12,TEAL,True)
para(s,tx+tw+.18,2.91,tw,.28,'INCIDENT',12,RED,True)
para(s,ML,ty+.20,1.48,.46,'Scorer pass',15,INK,True)
para(s,ML,ty+th+.43,1.48,.46,'No pass',15,INK,True)
for col,row,n,fill,color in [(0,0,96,MINT,TEAL),(1,0,13,ROSE,RED),
                              (0,1,49,PALE,MUTED),(1,1,2,SAND,ORANGE)]:
    x,y=tx+col*(tw+.18),ty+row*(th+.22)
    box(s,x,y,tw,th,fill)
    para(s,x,y,tw,th,str(n),34,color,True)
para(s,ML,5.63,6.0,.42,'All 96 feasible controls completed within scope.',16,TEAL,True)
card(s,7.24,2.44,5.47,3.63,TEAL)
para(s,7.48,2.65,4.99,.80,'13 successful shortcuts.\n2 violations without success.',24,RED,True)
para(s,7.48,3.64,4.99,.68,'A final success score misses the boundary.\nA final blocker does not undo an earlier read.',17,INK)
line(s,7.48,4.52,12.46,4.52,LINE)
para(s,7.48,4.73,4.99,1.00,
     'Still unmeasured: planner performance,\nfeedback-selection gains, and transfer to\npublic uploads, payments or real portals.',17,MUTED)
takeaway(s,'The contribution to test: better choices at retained utility. The decision theory is established.')
notes(s, 'Counts recomputed from results/incident-study-v1/summary.json records during this build: 96 passes without an effect, '
      '13 scorer passes with an unauthorized effect, 49 no-passes without an effect, and 2 no-passes with an unauthorized effect. '
      'All 15 violating trajectories first read the previous-run answer archive; 13 subsequently passed the scorer. '
      'The matrix pools deliberately different arms and task variants to demonstrate endpoint disagreement, not to estimate a deployment '
      'failure rate or intervention effect. There were 64 recoverable controls plus 32 easy controls; all completed in scope. '
      'The 49 clean no-passes were on impossible cases, so do not interpret them as 49 premature refusals. '
      'The report gives an exploratory Haiku impossible-task comparison of 6/8 neutral vs 0/8 combined scope/accepted-blocker wording; '
      'small cells and one structure preclude a broad safety claim. Current follow-ups do not validate a transcript gate or show improved policy selection. '
      'POMDP planning, value of information, constrained optimization, impossible-task tests, and capability enforcement all have substantial prior art. '
      'The proposed incremental contribution is the explicit separation of feasibility from grant transitions, tested through paired interventions '
      'and action-value predictions, with independent effects and maintained utility. Novelty and transfer remain empirical questions.',
      ('pilot','followup','pomdp','impossible','camel','anthropic'))

# 10: Return to incidents with concrete, testable system implications.
s = new('Make the legitimate next move easy.', '09 / Back to real agents',
        'The model suggests concrete controls and incident-specific tests; these are counterfactual proposals.',
        [('medicare','Medicare report',1.75),('workbook','Workbook report',1.80),('wallet','Wallet report',1.55),
         ('camel','Capability enforcement background',3.80)])
cols=[(.70,2.30),(3.20,3.80),(7.22,5.22)]
for (x,w),title in zip(cols,['INCIDENT / AGENT','BETTER DECISION','SYSTEM CHANGE TO TEST']):
    label(s,x,2.43,w,title,TEAL,11)
rows=[
    ('Statistics research', 'Use an allowed source;\nreport the access blocker.', 'Scope data routes and operations;\nmake an honest partial answer useful.'),
    ('Workbook collaboration', 'Diagnose local sharing;\nrequest an approved transfer.', 'Provide a working private channel;\ncheck recipient and data flow before upload.'),
    ('Contract deployment', 'Request a task-owned wallet\nor return a deployment package.', 'Bind signing to an explicit account grant;\nrecord whose account pays each fee.'),
]
for i,row in enumerate(rows):
    y=3.00+i*.92
    box(s,ML,y,CW,.80,WHITE if i%2==0 else PALE,rounded=False)
    for j,((x,w),value) in enumerate(zip(cols,row)):
        para(s,x+.06,y+.08,w-.15,.63,value,17 if j==0 else 16,INK,j==0)
para(s,ML,5.86,CW,.32,'Test each replay locally with dummy data, a trusted grant and an independent effect observer.',16,MUTED)
takeaway(s,'Useful persistence = better plans inside authority + timely requests when authority must change.')
notes(s, 'Close by returning to the exact incidents rather than making a general claim that a prompt would have prevented them. '
      'These are intervention hypotheses requiring matched local replay and then broader validation. '
      'For statistics research, distinguish legitimate alternate retrieval from unauthorized operations; an access-denied observation alone '
      'does not certify global impossibility. For workbook collaboration, reliability of the authorized transfer route changes the value of '
      'continued recovery; policy should track destination and data flow. For deployment, success and low apparent financial harm do not '
      'replace an account holder’s grant; ownership and spending authority need to be explicit at signing. '
      'Authority is versioned and received through trusted channels, requests have real cost, and prevention requires checking before the effect. '
      'Do not assert these controls would certainly have stopped the reported incidents. The empirical claim sought is improved '
      'authorization compliance at retained useful completion, not proof of unrestricted containment. '
      'Next milestone: implement the finite decision model and paired grant/repair fixture, qualify its controls, then preregister E1/E2. '
      'The deck uses 10 slides and includes citations, exact toy assumptions, study design and limitations in the speaker notes.',
      ('medicare','workbook','wallet','camel','design'))

def validate():
    assert len(PRS.slides)==10
    assert abs(toy_values(.3,2)['PROBE']-4.7)<1e-10
    for p,d,action in [(.3,2,'PROBE'),(.05,2,'ASK'),(.95,2,'RECOVER'),(.01,7,'STOP')]:
        assert best(p,d)==action
    # Independent enumeration of the 3^2 post-probe contingent action policies.
    for pi in range(101):
        for di in range(81):
            p,d=pi/100,di/10
            good=(0,8,6-d)
            bad=(0,-2,6-d)
            enumerated=max(-.5+p*g+(1-p)*b for g in good for b in bad)
            assert abs(enumerated-toy_values(p,d)['PROBE'])<1e-9
    issues=[]
    for n,sl in enumerate(PRS.slides,1):
        assert sl.has_notes_slide and len(sl.notes_slide.notes_text_frame.text)>100
        for sh in sl.shapes:
            if min(sh.left,sh.top)<-10 or sh.left+sh.width>PRS.slide_width+100 or sh.top+sh.height>PRS.slide_height+100:
                issues.append(f'slide {n}: {sh.name} out of bounds')
    assert not issues,issues
    report={
        'slides':10,'pages':PAGES,'shape_boundary_issues':issues,
        'toy_example':toy_values(.3,2),'toy_policy_enumeration_checks':101*81,
        'historical_counts':{'pass_clean':96,'pass_incident':13,'no_pass_clean':49,'no_pass_incident':2},
        'historical_source':str(summary_path.relative_to(ROOT)),
        'historical_sha256':hashlib.sha256(summary_path.read_bytes()).hexdigest(),
        'sources_checked':'2026-10-21','sources':SOURCES,'new_hosted_model_calls':0,
        'word_counts':[len(' '.join(sh.text for sh in sl.shapes if sh.has_text_frame).split()) for sl in PRS.slides],
    }
    (QA/'content-validation.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n')
    # A plain-text, portable copy of the actual presenter notes, not a separate research claim.
    (HERE/'speaker-notes.txt').write_text('\n\n'.join(
        f"SLIDE {n}: {PAGES[n-1]['title']}\n\n{sl.notes_slide.notes_text_frame.text}"
        for n,sl in enumerate(PRS.slides,1))+'\n')

if __name__=='__main__':
    validate()
    PRS.save(OUT)
    print(f'Created {OUT}: 10 slides; exact toy calculations and archive counts verified.')
