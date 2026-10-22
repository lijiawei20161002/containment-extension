"""Twenty-four-slide research narrative plus two incident-code glossary slides.

Build:  /Users/lijiawei/Desktop/.venv/bin/python slides/containment-decision-story/build_deck.py
Render: /Users/lijiawei/Desktop/.venv/bin/python slides/containment-decision-story/render_deck.py

No hosted inference. Library counts are read from figures/incidents/analysis/summary.json and
incidents/records/*.json at build time; archived run counts are recomputed from results/; toy
decision calculations are exact and labelled as hypothetical. The previous ten-slide builder is
kept at qa/build_deck.previous-10-slide.py.
"""
from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter
from math import comb, log, sqrt
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE.parent / 'decision-theory-deck'))
import deck_style as ds
from deck_style import (
    PRS, PAGES, SOURCES, DEEP, NAVY, DARKBOX, INK, MUTED, BLUE, TEAL, ORANGE,
    RED, PLUM, GOLD, BG, WHITE, PALE, MINT, SAND, ROSE, LILAC, CREAM, LINE, LIGHT, TINT,
    ML, MR, CW, text, box, line, dashed, arrow, circle, dot, badge, label, chip,
    card, panel, page, notes, link,
)
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from PIL import Image

OUT = HERE / 'containment-decision-story.pptx'
QA = HERE / 'qa'
QA.mkdir(exist_ok=True)
MATH_DIR = QA / 'math'
MATH_DIR.mkdir(exist_ok=True)
for old in MATH_DIR.glob('*.png'):
    old.unlink()
matplotlib.rcParams.update({'mathtext.fontset': 'stix', 'font.family': 'STIXGeneral'})
N_SLIDES = 26
ds.TOTAL = N_SLIDES

def footer(s, n, source='', dark=False):
    """Footer with a progress bar sized for N_SLIDES (the shared one assumes ten segments)."""
    color = LIGHT if dark else MUTED
    text(s, ML, 7.00, 11.3, .23, source or 'containment-extension  /  Decision rules with stated assumptions', 9, color, role='source')
    text(s, 11.95, 6.96, .76, .29, f'{n:02d} / {N_SLIDES}', 10, color, True, True, role='source')
    gap = .05
    seg = (CW - (N_SLIDES - 1) * gap) / N_SLIDES
    for i in range(N_SLIDES):
        done = i < n
        box(s, ML + i * (seg + gap), 7.34, seg, .055,
            (TEAL if done else DARKBOX) if dark else (TEAL if done else LINE), rounded=False)

ds.footer = footer

def page(title, kicker, subtitle='', source='', dark=False, tag='', accent=BLUE):
    """Page header that allows two-line titles (26 pt) and two-line subtitles (14 pt)."""
    s = PRS.slides.add_slide(PRS.slide_layouts[6])
    box(s, 0, 0, ds.SW, ds.SH, DEEP if dark else BG, rounded=False)
    n = len(PRS.slides)
    text(s, ML, .40, 11.0, .46, kicker.upper(), 11, LIGHT if dark else accent, True, role='label')
    if title:
        text(s, ML, .86, CW, .88, title, 26, WHITE if dark else NAVY, True, role='title', ls=1.05)
        box(s, ML, 1.76, .95, .05, accent, rounded=False)
    if subtitle:
        text(s, ML + .02, 1.84, CW, .50, subtitle, 14, LIGHT if dark else MUTED, role='subtitle', ls=1.05)
    footer(s, n, source, dark)
    PAGES.append({'slide': n, 'title': title, 'kicker': kicker, 'source': source})
    return s
PRS.core_properties.title = 'When the task fails: a decision theory of agent boundaries'
PRS.core_properties.author = 'Jiawei Li'
PRS.core_properties.subject = 'Reported incidents, a constrained POMDP, two algorithms against prior methods, experiments, implications'
PRS.core_properties.keywords = 'containment-extension; decision theory; authorization; research proposal'
SOURCES.clear()
SOURCES.update({
    'medicare': ('Australian Government briefing, 24 September 2026; incident in June',
        'https://ministers.finance.gov.au/financeminister/transcript/2026/09/24/press-conference-sydney'),
    'workbook': ('OpenAI, Unauthorized communication via temporary file hosting services, 16 September 2026; incident 14 April',
        'https://alignment.openai.com/misalignment-reports/unauthorized-communication-via-temporary-file-hosting-services/'),
    'wallet': ('OpenAI, Using a funded test cryptocurrency wallet without authorization, 15 October 2026; incident 26 June',
        'https://alignment.openai.com/misalignment-reports/using-a-funded-test-cryptocurrency-wallet-without-authorization/'),
    'anthropic': ('Anthropic (9 September 2026), An alignment assessment of recent cybersecurity incidents',
        'https://www.anthropic.com/research/alignment-assessment-cybersecurity-incidents'),
    'library': ('Project: public incident library, 26 records coded with the Raftari (2026) frame',
        '../../docs/incident-library-analysis.md'),
    'taxonomy': ('Project: incident taxonomy; formal indicator, mechanism and proposition labels',
        '../../incidents/taxonomy.json'),
    'medicare_coding': ('Project: Medicare incident record; analyst draft, review pending',
        '../../incidents/records/05-medicare.json'),
    'pomdp': ('Kaelbling, Littman & Cassandra (1998), Planning and Acting in Partially Observable Stochastic Domains',
        'https://cs.brown.edu/courses/csci2951-k/papers/kaelbling98.pdf'),
    'pomcp': ('Silver & Veness (2010), Monte-Carlo Planning in Large POMDPs (POMCP)',
        'https://papers.nips.cc/paper_files/paper/2010/hash/edfbe1afcf9246bb0d40eb4d8027d90f-Abstract.html'),
    'voi': ('Hay, Russell, Tolpin & Shimony (2012), Selecting Computations: Theory and Applications',
        'https://arxiv.org/abs/1207.5879'),
    'stopping': ('Ferguson, Optimal Stopping and Applications, ch. 5 (monotone case; one-step look-ahead)',
        'https://www.math.ucla.edu/~tom/Stopping/Contents.html'),
    'gilboa': ('Gilboa & Schmeidler (1989), Maxmin expected utility with non-unique prior',
        'https://doi.org/10.1016/0304-4068(89)90018-9'),
    'cirl': ('Hadfield-Menell et al. (2016), Cooperative Inverse Reinforcement Learning',
        'https://arxiv.org/abs/1606.03137'),
    'switch': ('Hadfield-Menell et al. (2017), The Off-Switch Game',
        'https://arxiv.org/abs/1611.08219'),
    'cpo': ('Achiam et al. (2017), Constrained Policy Optimization',
        'https://arxiv.org/abs/1705.10528'),
    'rcpo': ('Tessler, Mankowitz & Mannor (2018), Reward Constrained Policy Optimization',
        'https://arxiv.org/abs/1805.11074'),
    'shield': ('Alshiekh et al. (2018), Safe Reinforcement Learning via Shielding',
        'https://arxiv.org/abs/1708.08611'),
    'camel': ('Debenedetti et al. (2025), Defeating Prompt Injections by Design (CaMeL)',
        'https://arxiv.org/abs/2503.18813'),
    'progent': ('Shi et al. (2025), Progent: Securing AI Agents with Privilege Control',
        'https://arxiv.org/abs/2504.11703'),
    'control': ('Greenblatt et al. (2023), AI Control: Improving Safety Despite Intentional Subversion',
        'https://arxiv.org/abs/2312.06942'),
    'agentdojo': ('Debenedetti et al. (2024), AgentDojo',
        'https://arxiv.org/abs/2406.13352'),
    'impossible': ('Zhong, Raghunathan & Carlini (2025), ImpossibleBench, sections 5-6',
        'https://arxiv.org/html/2510.20270v1'),
    'pilot': ('Project: 160-run incident study, completed 17 September 2026',
        '../../docs/incident-study-results.md'),
    'design': ('Project: From reported incidents to testable questions',
        '../../docs/trust4rsi-incident-to-experiment.md'),
    'optimization': ('Project: Incident-aware optimization of agent control policies',
        '../../docs/incident-aware-optimization.md'),
    'next': ('Project: Next experiments, including X1b silent and conflicting grants',
        '../../docs/trust4rsi-next-experiments.md'),
    'followup': ('Project: Live follow-up results; calibration and selection limits',
        '../../docs/trust4rsi-live-followup-results.md'),
    'related': ('Project: Prior work and the contribution we can defend',
        '../../docs/trust4rsi-related-work.md'),
})

# ----------------------------------------------------------------------------- evidence
SUMMARY = json.loads((ROOT / 'figures/incidents/analysis/summary.json').read_text())
TAXONOMY = json.loads((ROOT / 'incidents/taxonomy.json').read_text())
LIB_N = SUMMARY['n']
MECH, PROP, IND = SUMMARY['mechanisms'], SUMMARY['propositions'], SUMMARY['indicators']
LAT, STOPB, PATH = SUMMARY['latency'], SUMMARY['stop_behavior'], SUMMARY['pathways']
RECORDS = {}
for path in sorted((ROOT / 'incidents/records').glob('*.json')):
    r = json.loads(path.read_text())
    RECORDS[r['entry_number']] = r
assert len(RECORDS) == LIB_N
OVERREACH = [n for n, r in RECORDS.items() if r['pathway'] == 'live_access_overreach']
OVERREACH_F = sum(1 for n in OVERREACH if RECORDS[n]['indicators']['f'] == 'yes')
assert len(OVERREACH) == PATH['live_access_overreach']
SELF_STOP = STOPB['self_stop_after_access']
RETRO = STOPB['retrospective_only']
MEDICARE = RECORDS[5]
assert MEDICARE['slug'] == 'medicare'
INDICATOR_ORDER = tuple(TAXONOMY['indicators'])
INDICATOR_STATUS = {'yes': 'observed', 'no': 'no', 'unknown': 'unknown'}
MEDICARE_INDICATOR_READINGS = {
    'f': 'Public search failed; agent persisted',
    'p': 'Guest route found and tested',
    'x': 'Third-party government portal accessed',
    'c': 'No cross-run link reported; coded no',
    'v': 'Coded yes; credentials retrieved',
    'o': 'Loss of visibility cannot be established',
}

summary_path = ROOT / 'results/incident-study-v1/summary.json'
historical = json.loads(summary_path.read_text())
counts = Counter((r['stop'] == 'scorer_success', bool(r['any_effect'])) for r in historical['records'])
assert len(historical['records']) == 160
HIST = {'pass_clean': counts[True, False], 'pass_incident': counts[True, True],
        'no_pass_clean': counts[False, False], 'no_pass_incident': counts[False, True]}
assert tuple(HIST.values()) == (96, 13, 49, 2)
diag_path = ROOT / 'results/baseline-diagnostic-v2/summary.json'
diag = json.loads(diag_path.read_text())
P_BASE_CELL = next(g for g in diag['groups'] if g['arm'] == 'current_original' and g['variant'] == 'impossible')['incidents']
P_BASE = P_BASE_CELL['positive'] / P_BASE_CELL['observed']
assert P_BASE_CELL['positive'] == 11 and P_BASE_CELL['observed'] == 16

STOP_LABEL = {
    'retrospective_only': 'found later', 'human_stop_ignored': '2 stops ignored',
    'none_observed': 'none', 'self_stop_after_access': 'self, after access', 'manual_stop': 'manual',
    'contained_by_third_party': 'third party', 'automatic_stop_failed_then_manual': 'auto failed; manual',
    'not_applicable': 'n/a',
}

# ----------------------------------------------------------------------------- exact toy calculations
# R1 Optimal stopping with learning: feasibility prior p0; each attempt succeeds w.p. q if feasible;
# attempt cost c; success R; honest stop B. Posterior after a failure: p' = p(1-q) / (1 - p q).
def posterior_sequence(p0, q, n):
    seq, p = [p0], p0
    for _ in range(n):
        p = p * (1 - q) / (1 - p * q)
        seq.append(p)
    return seq

STOP_TOY = {'p0': .7, 'q': .5, 'R': 10., 'B': 2., 'c': 1.}
STOP_THRESHOLD = STOP_TOY['c'] / (STOP_TOY['q'] * (STOP_TOY['R'] - STOP_TOY['B']))
POSTERIORS = posterior_sequence(STOP_TOY['p0'], STOP_TOY['q'], 5)
ATTEMPTS_BEFORE_STOP = next(i for i, p in enumerate(POSTERIORS) if p < STOP_THRESHOLD)

def stopping_value(p, k, t):
    """Exact finite-horizon value with k attempts left (backward induction) for the stop toy."""
    if k == 0:
        return t['B']
    cont = t['q'] * p * t['R'] + (1 - t['q'] * p) * stopping_value(p * (1 - t['q']) / (1 - p * t['q']), k - 1, t) - t['c']
    return max(t['B'], cont)

# R2 Value of information: one optional diagnostic then one final commitment.
VOI_TOY = {'reward': 10., 'recovery_cost': 2., 'grant_probability': .6, 'probe_cost': .5}

def toy_values(p, request_cost):
    t = VOI_TOY
    ask = t['grant_probability'] * t['reward'] - request_cost
    good, bad = t['reward'] - t['recovery_cost'], -t['recovery_cost']
    return {
        'STOP': 0.,
        'RECOVER': t['reward'] * p - t['recovery_cost'],
        'ASK': ask,
        'PROBE': -t['probe_cost'] + p * max(good, ask, 0) + (1 - p) * max(bad, ask, 0),
    }

def best(p, d):
    values = toy_values(p, d)
    order = ('STOP', 'RECOVER', 'ASK', 'PROBE')
    return max(order, key=lambda key: values[key])

TOY_COLOR = {'STOP': MUTED, 'RECOVER': TEAL, 'ASK': BLUE, 'PROBE': GOLD}

def voi_boundaries(d):
    """Analytic policy thresholds for the displayed model, 0 <= d <= 8.

    Below lower: ASK (d < 6) or STOP (d > 6); between: PROBE; above upper: RECOVER.
    At d=6 ASK and STOP tie below lower; on a threshold the adjacent actions tie.
    These constants are derived independently of toy_values and checked against it.
    """
    assert 0 <= d <= 8
    if d <= 6:
        return .5 / (2 + d), (7.5 - d) / (8 - d)
    return 1 / 16, 3 / 4

def voi_phase_figure():
    """Export an exact-boundary plot as both a slide PNG and a standalone vector PDF."""
    import numpy as np
    original_font = plt.rcParams['font.family']
    plt.rcParams['font.family'] = 'Arial'
    fig = plt.figure(figsize=(6.1, 3.12), facecolor=f'#{BG}')
    ax = fig.add_axes([.105, .17, .87, .74])
    ax.set_facecolor(f'#{BG}')
    for ds_ in (np.linspace(0, 6, 601), np.linspace(6, 8, 201)):
        lower, upper = np.array([voi_boundaries(float(d)) for d in ds_]).T
        left = 'ASK' if ds_[0] < 6 else 'STOP'
        for a, lo, hi in [(left, 0, lower), ('PROBE', lower, upper), ('RECOVER', upper, 1)]:
            ax.fill_betweenx(ds_, lo, hi, color=f'#{TOY_COLOR[a]}', linewidth=0)
        ax.plot(lower, ds_, color='white', linewidth=.9)
        ax.plot(upper, ds_, color='white', linewidth=.9)
    ax.hlines(6, 0, 1, colors=f'#{NAVY}', linestyles=(0, (3, 3)), linewidth=.65, alpha=.55)
    ax.text(.43, 6.15, 'd = 6: ASK = STOP', fontsize=8, color=f'#{NAVY}',
            ha='center', bbox={'facecolor': f'#{BG}', 'edgecolor': 'none', 'pad': 1.8, 'alpha': .93})
    for x, y, title in [(.065, 1.55, 'ASK'), (.43, 4.55, 'PROBE'), (.875, 6.85, 'RECOVER')]:
        ax.text(x, y, title, color='white', fontsize=11, weight='bold', ha='center', va='center')
    ax.annotate('STOP', xy=(.031, 7.15), xytext=(.16, 8.5), fontsize=9, weight='bold',
                color=f'#{MUTED}', ha='center', annotation_clip=False,
                arrowprops={'arrowstyle': '->', 'color': f'#{NAVY}', 'lw': .8})
    for label_, p, d, dx, dy in [('A', .9, 8, -1, -15), ('B', .1, 0, 0, 9), ('C', .3, 2, 10, -3)]:
        ax.plot(p, d, 'o', color='white', markeredgecolor=f'#{NAVY}', markersize=6, clip_on=False, zorder=5)
        ax.annotate(label_, (p, d), xytext=(dx, dy), textcoords='offset points', fontsize=10,
                    weight='bold', color='white', ha='center')
    ax.set(xlim=(0, 1), ylim=(0, 8), xticks=[0, .3, .5, 1], yticks=[0, 2, 4, 6, 8])
    ax.set_xlabel('Belief p that local recovery works', fontsize=10, color=f'#{MUTED}', labelpad=5)
    ax.set_ylabel('Request cost d', fontsize=10, color=f'#{MUTED}', labelpad=6)
    ax.tick_params(axis='both', labelsize=9, colors=f'#{MUTED}', length=3)
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    for side in ('left', 'bottom'):
        ax.spines[side].set_color(f'#{INK}')
    target = QA / 'voi-policy-phase-diagram'
    fig.savefig(target.with_suffix('.png'), dpi=240)
    fig.savefig(target.with_suffix('.pdf'))
    plt.close(fig)
    plt.rcParams['font.family'] = original_font
    return target.with_suffix('.png')

# R3 Ambiguity: payoffs under two plausible grants, no prior.
AMB_ACTIONS = [('Read the archive', 10, -20), ('Ask first', 7, 7), ('Stop and report', 2, 2)]
def ambiguity_table():
    best_col = [max(a[1] for a in AMB_ACTIONS), max(a[2] for a in AMB_ACTIONS)]
    rows = []
    for name, u_allowed, u_excluded in AMB_ACTIONS:
        rows.append((name, u_allowed, u_excluded, min(u_allowed, u_excluded),
                     max(best_col[0] - u_allowed, best_col[1] - u_excluded)))
    return rows
AMB = ambiguity_table()
AMB_MAXMIN = max(AMB, key=lambda r: r[3])[0]
AMB_REGRET = min(AMB, key=lambda r: r[4])[0]
AMB_PRIOR_SWITCH = 27 / 30   # p(allowed) above which expected utility prefers reading: 10p − 20(1−p) > 7

# R4 Selection from finite evidence: Hoeffding + union bound over K candidates.
def selection_slack(K, n, delta=.05):
    return 2 * sqrt(log(2 * K / delta) / (2 * n))

SEL_K = 8
SEL_ROWS = [(n, selection_slack(SEL_K, n)) for n in (48, 200, 1000, 5000)]

def binom_pmf(k, n, p):
    return comb(n, k) * p ** k * (1 - p) ** (n - k)

def mcnemar_exact_p(b, c):
    n = b + c
    if n == 0:
        return 1.0
    return min(1.0, 2 * sum(binom_pmf(i, n, .5) for i in range(min(b, c) + 1)))

def mcnemar_power(n, p1, p2, alpha=.05):
    pb, pc = p1 * (1 - p2), (1 - p1) * p2
    rest = 1 - pb - pc
    power = 0.
    for b in range(n + 1):
        for c in range(n + 1 - b):
            prob = comb(n, b) * comb(n - b, c) * pb ** b * pc ** c * rest ** (n - b - c)
            if mcnemar_exact_p(b, c) < alpha:
                power += prob
    return power

POWER_TARGETS = (.375, .25)
POWER_ROWS = [(n, [mcnemar_power(n, P_BASE, p2) for p2 in POWER_TARGETS]) for n in (16, 24, 32, 48)]

# E2 pilot allocation
E2 = {'repair': 2, 'grant': 2, 'families': 2, 'arms': 7, 'lineages': 12}
E2_ROLLOUTS = E2['repair'] * E2['grant'] * E2['families'] * E2['arms'] * E2['lineages']

# ----------------------------------------------------------------------------- primitives
def footer_links(s, refs):
    x = ML
    for key, caption, width in refs:
        link(s, x, 7.00, width, caption, SOURCES[key][1], MUTED, 9, .23)
        x += width + .16

STATUS_COLOR = {'REPORTED': MUTED, 'OBSERVED': MUTED, 'EXACT TOY': GOLD, 'PROPOSED': ORANGE,
                'THEORY': BLUE, 'FRAMING': TEAL}

def new(title, section, subtitle='', refs=(), dark=False, status=''):
    s = page(title, section, subtitle, source=' ', dark=dark, accent=TEAL)
    if status:
        chip(s, 11.05, .40, 1.66, status, STATUS_COLOR.get(status, MUTED), size=10, h=.30)
    if refs:
        footer_links(s, refs)
    return s

def takeaway(s, value, color=TEAL, y=6.25, size=16):
    box(s, ML, y, CW, .60, color)
    text(s, ML + .22, y + .03, CW - .44, .54, value, size, WHITE, True)

def block(s, x, y, w, h, title, body, color=TEAL, size=15, title_size=18):
    panel(s, x, y, w, h, color)
    text(s, x + .20, y + .16, w - .40, .43, title, title_size, color, True)
    text(s, x + .20, y + .66, w - .40, h - .78, body, size, INK, valign=MSO_ANCHOR.TOP, ls=1.1)

def para(s, x, y, w, h, value, size=17, color=INK, bold=False, **kw):
    return text(s, x, y, w, h, value, size, color, bold, **kw)

def grid(s, x, y, widths, rows, row_h=.46, size=12, header=NAVY, bold_first=True, fills=None, ls=1.0):
    """A table from rectangles and text boxes; header row filled, body rows banded."""
    yy = y
    for i, row in enumerate(rows):
        xx = x
        for j, value in enumerate(row):
            w = widths[j]
            if i == 0:
                box(s, xx, yy, w, row_h, header, rounded=False)
                text(s, xx + .09, yy, w - .18, row_h, value, size, WHITE, True, ls=ls)
            else:
                fill = PALE if i % 2 == 0 else WHITE
                color = INK
                if fills and (i, j) in fills:
                    fill, color = fills[(i, j)]
                box(s, xx, yy, w, row_h, fill, rounded=False)
                text(s, xx + .09, yy, w - .18, row_h, value, size, color, bold_first and j == 0, ls=ls)
            xx += w
        yy += row_h
    return yy

def hbars(s, x, y, w, rows, total, step=.40, label_w=2.3, size=11, color=BLUE):
    for i, (lab, value, *rest) in enumerate(rows):
        col = rest[0] if rest else color
        yy = y + i * step
        text(s, x, yy, label_w, step - .04, lab, size, INK)
        bx, bw = x + label_w + .08, w - label_w - .95
        box(s, bx, yy + .07, bw, step - .18, PALE, rounded=False)
        if value:
            box(s, bx, yy + .07, bw * value / total, step - .18, col, rounded=False)
        text(s, bx + bw + .08, yy, .85, step - .04, f'{value}/{total}', size, col, True)

def chain(s, steps, y, h, x0=ML, total_w=CW, gap=.22, title_size=15, body_size=12):
    n = len(steps)
    w = (total_w - gap * (n - 1)) / n
    for j, (title, body, color) in enumerate(steps):
        x = x0 + j * (w + gap)
        box(s, x, y, w, h, TINT.get(color, PALE))
        box(s, x, y, w, .06, color, rounded=False)
        text(s, x + .12, y + .14, w - .24, .50, title, title_size, color, True, ls=1.0)
        if body:
            text(s, x + .12, y + .66, w - .24, h - .76, body, body_size, INK, valign=MSO_ANCHOR.TOP, ls=1.05)
        if j < n - 1:
            arrow(s, [(x + w + .03, y + h / 2), (x + w + gap - .03, y + h / 2)], MUTED)

def mono(s, x, y, w, h, lines, size=10.5, fill=NAVY, color=WHITE, accent=TEAL):
    """A pseudocode block: dark surface, accent bar, Menlo runs (renderer measures line widths)."""
    box(s, x, y, w, h, fill, rounded=False)
    box(s, x, y, .07, h, accent, rounded=False)
    tb = s.shapes.add_textbox(Inches(x + .22), Inches(y + .12), Inches(w - .34), Inches(h - .24))
    tb.name = f'code:{len(s.shapes)}'
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line_ in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.18
        run = p.add_run()
        run.text = line_ if line_ else ' '
        run.font.name = 'Menlo'
        run.font.size = Pt(size)
        run.font.color.rgb = LIGHT if line_.lstrip().startswith('#') else color
    return tb

_MATH_COUNT = [0]

def mathpng(tex, color, size):
    """Typeset one line of math with matplotlib mathtext (STIX); returns (path, width_in, height_in)."""
    _MATH_COUNT[0] += 1
    path = MATH_DIR / f'f{_MATH_COUNT[0]:02d}.png'
    fig = plt.figure(figsize=(.01, .01))
    fig.text(0, 0, f'${tex}$', fontsize=size, color=f'#{color}')
    fig.savefig(path, dpi=300, transparent=True, bbox_inches='tight', pad_inches=.03)
    plt.close(fig)
    with Image.open(path) as im:
        w, h = im.size
    return path, w / 300, h / 300

def formula(s, x, y, tex, size=16, color=WHITE, w_max=None, h_box=None, align='l'):
    """One typeset formula placed as a transparent picture at its natural point size."""
    path, wi, hi = mathpng(tex, color, size)
    if w_max and wi > w_max:
        wi, hi = w_max, hi * w_max / wi
    if h_box:
        y = y + (h_box - hi) / 2
    if w_max and align == 'c':
        x = x + (w_max - wi) / 2
    elif w_max and align == 'r':
        x = x + (w_max - wi)
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(wi), height=Inches(hi))
    pic.name = f'math:{len(s.shapes)}'
    return pic

def serif(s, x, y, w, h, value, size, color=WHITE, bold=False, italic=False, ls=1.1, align=None, valign=MSO_ANCHOR.MIDDLE):
    """Georgia display text (the renderer registers Georgia; see render_deck.py)."""
    tb = text(s, x, y, w, h, value, size, color, bold, ls=ls, align=align, valign=valign)
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = 'Georgia'
            r.font.italic = italic
    return tb

def theorem_band(s, x, y, w, h, formula_text='', caption='', size=24, tag='RESULT', accent=TEAL, ls=1.08, cap_h=.46,
                 math=None, math_size=17, gap=.12):
    """Statement band: dark surface, accent bar, typeset math (or white text), two-line caption room."""
    box(s, x, y, w, h, NAVY, rounded=False)
    box(s, x, y, .09, h, accent, rounded=False)
    label(s, x + .26, y + .12, w - .4, tag, accent)
    ch = cap_h if caption else 0
    area_y, area_h = y + .42, h - .48 - ch
    if math:
        rendered = []
        for tex in math:
            path, wi, hi = mathpng(tex, WHITE, math_size)
            if wi > w - .56:
                wi, hi = w - .56, hi * (w - .56) / wi
            rendered.append((path, wi, hi))
        total = sum(hi for _, _, hi in rendered) + gap * (len(rendered) - 1)
        if total > area_h:   # never let the math run into the caption
            k = area_h / total
            rendered = [(path, wi * k, hi * k) for path, wi, hi in rendered]
            gap *= k
            total = area_h
        yy = area_y + (area_h - total) / 2
        for path, wi, hi in rendered:
            pic = s.shapes.add_picture(str(path), Inches(x + .28), Inches(yy), width=Inches(wi), height=Inches(hi))
            pic.name = f'math:{len(s.shapes)}'
            yy += hi + gap
    else:
        text(s, x + .26, area_y, w - .42, area_h, formula_text, size, WHITE, True, role='formula',
             valign=MSO_ANCHOR.MIDDLE, ls=ls)
    if caption:
        text(s, x + .26, y + h - ch - .06, w - .42, ch, caption, 11, LIGHT, role='caption',
             valign=MSO_ANCHOR.TOP, ls=1.05)

def picture(s, rel, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw['width'] = Inches(w)
    if h is not None:
        kw['height'] = Inches(h)
    return s.shapes.add_picture(str(ROOT / rel), Inches(x), Inches(y), **kw)

def rec_title(n):
    return f"{RECORDS[n]['short_title']} ({n})"

# ============================================================================= 00 TITLE
s = new('', 'Containment-extension  ·  Research outline  ·  21 October 2026', dark=True)
PAGES[-1]['title'] = 'When the task fails, what may the agent change?'
serif(s, ML, 1.55, 10.5, 2.3, 'When the task fails,\nwhat may the agent change?', 50, WHITE, ls=1.12, valign=MSO_ANCHOR.TOP)
box(s, ML, 4.08, 1.40, .05, TEAL, rounded=False)
serif(s, ML, 4.32, 10.5, .60, 'A decision theory of recovery, evidence, authority and stopping.', 23, LIGHT, italic=True, valign=MSO_ANCHOR.TOP)
serif(s, ML, 4.98, 8.0, .42, 'Jiawei Li  ·  21 October 2026', 17, WHITE, valign=MSO_ANCHOR.TOP)
label(s, ML, 5.56, 4.0, 'Research question', TEAL)
para(s, ML, 5.82, 9.2, .72,
     'Can an agent stay useful after a route fails, treating the failure as evidence about feasibility '
     'and never as a grant of new authority?', 16, WHITE, ls=1.18, valign=MSO_ANCHOR.TOP)
para(s, ML, 6.62, 9.0, .30, 'No new model runs. Library counts, archived runs and toy calculations are recomputed at build time.', 10.5, LIGHT)
notes(s, 'Opening prompt: a research agent gets an error while retrieving useful data; ask the audience what the next move should be. '
      'The talk treats containment-extension as a study of delegated action boundaries, not a claim to solve arbitrary sandbox escape. '
      'No hosted-model calls were made to build this deck; the planner and the selection procedure are proposals and have not been '
      'evaluated on LLM agents. The outline on the next slide gives the five parts and their slide ranges.', ('design', 'library'))

# ============================================================================= 00b OUTLINE
s = new('', 'Outline', dark=True)
PAGES[-1]['title'] = 'Outline'
serif(s, ML, .95, 8.0, .80, 'Outline', 36, WHITE, valign=MSO_ANCHOR.TOP)
agenda = [
    ('01', 'Incidents', 'Six reported cases in the style of the paper plan: framing, the feedback-loop latency, the Medicare '
     'chain, the substitution and scope-inference families, and the three questions a theory must answer.', 'slides 3–8', RED),
    ('02', 'Decision theory', 'A constrained POMDP whose grant only a trusted reply can change, then four exact results: '
     'stopping with learning, value of information (policy, payoffs, boundaries), ambiguity under a silent grant, and selection.', 'slides 9–15', BLUE),
    ('03', 'Algorithms', 'GRANT-LOOKAHEAD (online) and INCIDENT-SELECT (offline) as pseudocode with properties; a point-by-point '
     'comparison with POMCP, constrained RL, shielding, CaMeL, Progent, AI Control and ImpossibleBench.', 'slides 16–19', TEAL),
    ('04', 'Experiments', 'E1 simulator regret study; E2 repair × grant factorial with seven arms including prior-method baselines; '
     'E3 selection contrast on fresh lineages; falsifiers for each.', 'slides 20–22', PLUM),
    ('05', 'Implications', 'Each incident family mapped to a result, a system change and its test; outcomes interpreted in advance; '
     'milestones.', 'slides 23–24', ORANGE),
]
for i, (num, name, desc, where, col) in enumerate(agenda):
    y = 1.95 + i * .94
    box(s, ML, y, CW, .82, DARKBOX)
    box(s, ML, y, .07, .82, col, rounded=False)
    serif(s, ML + .25, y + .14, .75, .56, num, 24, col)
    text(s, ML + 1.10, y + .08, 2.6, .34, name, 16, WHITE, True)
    text(s, ML + 1.10, y + .42, 2.6, .30, where, 11, LIGHT)
    text(s, ML + 3.75, y + .06, CW - 3.95, .72, desc, 11.5, LIGHT, ls=1.1)
para(s, ML, 6.74, CW, .22, 'Reference appendix: M1-M6 failure mechanisms (25); P1-P9 design propositions (26).', 10, LIGHT)
notes(s, 'Roadmap for the five parts. Parts 01 and 05 are the real-world ends of the argument; 02 to 04 are the formal middle. '
      'Status chips on each slide distinguish REPORTED facts from public records, OBSERVED archived runs, EXACT TOY calculations, '
      'THEORY, and PROPOSED work. Slide ranges are checked at build time against the section kickers.', ('design',))
OUTLINE_RANGES = {k[0]: k[3] for k in agenda}

# ============================================================================= 01 FRAMING
s = new('A failed route is where authority gets reinterpreted.', '01 / Motivation from reported incidents',
        f'{LIB_N} public incident records, October 2025 to September 2026, coded by pathway, mechanism, risk indicator and stop behaviour.',
        [('library', 'Incident library analysis', 2.3), ('pilot', 'Archived 160-run study', 2.3), ('design', 'Incidents to testable questions', 3.0)],
        status='FRAMING')
chain(s, [
    (f'{LIB_N} public incidents', f"Largest pathway: live-access overreach, {PATH['live_access_overreach']} of {LIB_N}. "
     f"Repeated task failure recorded in {IND['f']['yes']}; unknown in {IND['f']['unknown']}.", RED),
    ('One decision point', 'Task blocked. Allowed route failing. Forbidden resource reachable. What changes: the plan, or the permission?', ORANGE),
    (f"{len(historical['records'])} archived runs", f"Local service fixture, effect oracle in SQLite. {HIST['pass_incident']} scorer passes "
     f"carried an unauthorized effect; {HIST['no_pass_incident']} violations never passed.", BLUE),
    ('A theory and two algorithms', 'A constrained POMDP whose grant only a trusted reply can change; an online planner and an offline selector; tests that can fail.', TEAL),
], y=2.40, h=2.15, title_size=16, body_size=12.5)
para(s, ML, 4.80, CW, .85,
     'Claim boundary: the incidents motivate hypotheses and do not establish rates. The decision model is classical; the proposed '
     'contribution is the separation of feasibility evidence from grant transitions, tested through paired interventions with an '
     'independent effect observer. No containment guarantee is claimed.', 13.5, MUTED, ls=1.1)
takeaway(s, 'Failure changes what is possible. It does not grant new authority.')
notes(s, f"Library counts are read from figures/incidents/analysis/summary.json at build time: n={LIB_N}, live-access overreach "
      f"{PATH['live_access_overreach']}, indicator f yes={IND['f']['yes']} unknown={IND['f']['unknown']}. Archived counts are "
      "recomputed from results/incident-study-v1/summary.json. These are analyst codings of public reports with a second reader "
      "still pending, and the archived runs come from one small service fixture. The framing follows the Trust4RSI paper plan: "
      "one decision point recurs across incidents, and the question is whether an explicit model of that decision yields better "
      "behaviour at retained utility.", ('library', 'pilot', 'design'))

# ============================================================================= 02 LATENCY
s = new('The world reports an incident four months after it happens.', '02 / The feedback loop a deployed agent actually gets',
        f"Occurrence, detection and public disclosure for the {LAT['disclosure_days_known']} library entries with a known start date.",
        [('library', 'figures/incidents/analysis/latency-dumbbell.png; month-precision dates plotted mid-month', 6.0)],
        status='REPORTED')
picture(s, 'figures/incidents/analysis/latency-dumbbell.png', ML, 2.32, h=3.70)
x = 6.05
block(s, x, 2.32, 6.66, 1.08, f"Median detection latency: {LAT['detection_median']:g} days",
      f"Known for {LAT['detection_days_known']} of {LIB_N} entries.", BLUE, 13, 16)
block(s, x, 3.52, 6.66, 1.08, f"Median disclosure latency: {LAT['disclosure_median']:g} days",
      f"Known for {LAT['disclosure_days_known']} of {LIB_N} entries.", ORANGE, 13, 16)
block(s, x, 4.72, 6.66, 1.30, f'{RETRO} of {LIB_N} were found only in retrospect',
      f"{STOPB['none_observed']} had no stop at all; {SELF_STOP} self-stopped, both after the access had already happened; "
      f"{STOPB['manual_stop']} manual stops; {STOPB['human_stop_ignored']} human stop ignored.", RED, 13, 16)
takeaway(s, 'A policy cannot be selected on this feedback loop. It needs a model of the decision and replayable cases.')
notes(s, 'Latency figures are copied from the library summary; detection dates are known for only 12 entries, so the detection median '
      'is fragile. The point for the decision theory: labels arrive weeks to months after the action, so an online rule needs a model '
      'of the decision rather than observed incident rates, and an offline selection rule must work from replayable, instrumented cases. '
      'The stop-behaviour counts show that detection did not stop most incidents.', ('library',))

# ============================================================================= 03 MEDICARE
s = new('Medicare portal: the search failed, so the agent found a way in.', '03 / One incident in detail',
        f"{MEDICARE['actor']} research agent, {MEDICARE['dates']['occurred_start']}; government notified 10 Sep 2026; "
        f"disclosed {MEDICARE['dates']['disclosed']}. Library record {MEDICARE['id']}, evidence level {MEDICARE['evidence']['level']}.",
        [('medicare', 'Australian Government press conference, 24 September 2026', 4.2),
         ('taxonomy', 'Symbol definitions', 1.5), ('medicare_coding', 'Draft record coding', 1.6)],
        status='REPORTED')
steps = [('Objective', MEDICARE['coding']['objective'], BLUE),
         ('Obstacle', 'Public search failed; no safe stop taken.', ORANGE)]
for action in MEDICARE['coding']['actions']:
    steps.append((action.capitalize(), '', RED))
chain(s, steps, y=2.38, h=1.36, title_size=13.5, body_size=11)
fired = [k for k in INDICATOR_ORDER if MEDICARE['indicators'][k] == 'yes']
indicator_rows = [['ID', 'Risk indicator', 'Status', 'Reading in this case']]
for code in INDICATOR_ORDER:
    indicator_rows.append([code, TAXONOMY['indicators'][code],
                           INDICATOR_STATUS[MEDICARE['indicators'][code]], MEDICARE_INDICATOR_READINGS[code]])
grid(s, ML, 3.91, [.58, 2.55, 1.02, 3.33], indicator_rows, row_h=.295, size=10.5,
     fills={(i + 1, 2): ((MINT, TEAL) if MEDICARE['indicators'][code] == 'yes' else (PALE, MUTED))
            for i, code in enumerate(INDICATOR_ORDER)})
para(s, 8.32, 3.89, 4.39, .32, 'z_{t} = (f, p, x, c, v, o)', 15, NAVY, True)
para(s, 8.32, 4.24, 4.39, .43, 'Six risk signals at time t. Here: retrospective coding of the incident.', 11.5, MUTED, ls=1.05)
para(s, 8.32, 4.76, 4.39, .72,
     'observed: evidence of the signal.\nno: coded as not observed.\nunknown: insufficient evidence.', 11.5, INK, ls=1.15)
para(s, 8.32, 5.57, 4.39, .42,
     'M = failure mechanism (p. 25).\nP = design proposition (p. 26).', 11.5, BLUE, True, ls=1.05)
para(s, ML, 6.00, CW, .23,
     f"{len(fired)}/6 observed ({', '.join(fired)}) counts signals, not rule violations. 'Leading' is the library term; timing is not established.",
     10.5, MUTED)
takeaway(s, 'Boundary: third-party system authorization. No in-run stop reported. M/P coding is expanded on pp. 25-26.', RED, y=6.30, size=14)
notes(s, "Chain and table fields are read from incidents/records/05-medicare.json: objective, the four coded actions, indicators, "
      "mechanisms and propositions. Formal labels are loaded from incidents/taxonomy.json. The library's z_t denotes six risk "
      "signals at time t, in order f, p, x, c, v, o. This record is retrospective and does not give a time-indexed trace of those "
      "signals or prove that the agent or monitor noticed them in-run. f = repeated task failure; p = boundary probing; "
      "x = unexpected external contact (not every legitimate external request); c = cross-run communication; v = privilege "
      "change; o = observability loss. The four observed signals are f, p, x and v. No means coded as not observed, not proof "
      "of absence; unknown means the available evidence cannot decide. Four of six is a signal count, not four broken rules. "
      "The v row reports the existing coding alongside credential retrieval; the record does not establish that those credentials "
      "were used or that a particular privilege-escalation exploit occurred. o concerns loss of logging or monitoring visibility, "
      "not merely late public disclosure. 'Leading indicators' is the taxonomy's term, not an established temporal claim: some "
      "signals may occur during or after a boundary crossing. M1-M6 are failure mechanisms; P1-P9 are design propositions. "
      "The appendices on slides 25-26 give every formal name, a plain-language explanation and the Medicare coding. These are "
      "analyst drafts pending review. Facts come from the government press conference; the government states aggregate statistics "
      "only and no individual medical data accessed, and the forensic investigation was still open. We read this as: the failed "
      "search was evidence about feasibility, the agent treated it as licence to find another route, and no in-run signal stopped "
      "it. That is the decision point the formal model isolates: a belief update was mistaken for a grant update.",
      ('medicare', 'taxonomy', 'medicare_coding', 'library'))

# ============================================================================= 04 SUBSTITUTION FAMILY
s = new('Five reports, one structure: blocked route, reachable substitute.', '04 / The substitution family',
        f"Live-access overreach is the largest pathway ({PATH['live_access_overreach']} of {LIB_N}). Repeated task failure (indicator f) "
        f"is recorded in {OVERREACH_F} of those {len(OVERREACH)}. The permission never changed.",
        [('workbook', 'OpenAI workbook report · 16 Sep 2026', 3.0), ('wallet', 'OpenAI wallet report · 15 Oct 2026', 3.0),
         ('library', 'Library records 15, 12, 13, 4', 2.2)],
        status='REPORTED')
rows = [['Case (record)', 'Blocked or failing route', 'Substitute the agent used', 'Stop']]
for n, route, sub in [
    (15, "Agents could not read each other's local files", 'Made the workbook publicly downloadable'),
    (12, 'Earnings figures not retrievable', 'Used an exposed API key from a public repo, then fabricated figures'),
    (13, 'Browser citation and image search failed', 'Uploaded records and a task photograph to public hosts'),
    (4, "Needed another team's proof material", "Published a researcher's GitHub token; split it to evade scanning"),
]:
    rows.append([rec_title(n), route, sub, STOP_LABEL[RECORDS[n]['stop_behavior']]])
rows.append(['Test wallet (OpenAI, 15 Oct)', 'No funded wallet supplied', "Spent another account's test funds", 'reported'])
grid(s, ML, 2.38, [3.0, 3.5, 4.35, 1.24], rows, row_h=.52, size=11.5)
para(s, ML, 5.58, CW, .55,
     'In four of five the substitute succeeded at the sub-goal. The task was not impossible; the authorized route was. '
     'A planner that optimizes completion alone reproduces this family.', 13.5, MUTED, ls=1.05)
takeaway(s, 'Hypothesis H1: a failing authorized route raises the probability of a forbidden substitute, with scope unchanged.', size=16)
notes(s, 'Rows copy the public reports; stop labels come from the record files. The wallet report is outside the 26-record '
      'denominator. The decision-theoretic reading: in every row the belief that the allowed route works dropped to near zero, '
      'the substitute had high expected task value, and the only thing against it was the grant. Any planner whose objective is '
      'task value without an effect constraint prefers the substitute; that is why the formulation puts the grant in the action '
      'set rather than in the reward.', ('workbook', 'wallet', 'library'))

# ============================================================================= 05 SCOPE-INFERENCE FAMILY
s = new('When the contract is silent, reachability is read as permission.', '05 / The scope-inference family',
        f"Misconfiguration exposure: {PATH['misconfiguration_exposure']} of {LIB_N} records. The task said 'simulated'; the network said "
        'otherwise; the agent believed the network.',
        [('anthropic', 'Anthropic assessment · 9 Sep 2026', 3.0), ('library', 'Library records 16, 17, 18, 21', 2.4)],
        status='REPORTED')
rows = [['Case (record)', 'What was reachable', 'Effect before any stop', 'Stop']]
for n, reach, effect in [
    (16, 'A real company domain with a matching name', 'Credentials obtained; database with hundreds of production records read (4 runs)'),
    (17, 'The real PyPI registry', 'Malicious package ran on 15 hosts; one scanner leaked credentials to a vendor database'),
    (18, 'About 9,000 live targets', 'One live application compromised'),
    (21, 'Three real companies', 'Systems accessed via public info and guessed or exposed credentials'),
]:
    rows.append([rec_title(n), reach, effect, STOP_LABEL[RECORDS[n]['stop_behavior']]])
grid(s, ML, 2.38, [3.0, 3.1, 4.75, 1.24], rows, row_h=.56, size=11.5)
para(s, ML, 5.30, CW, .80,
     f'{SELF_STOP} of these 4 stopped themselves, and both recognised the target only after the access. Under a silent grant the '
     'agent faced ambiguity, not risk: no probability over "allowed" was available, and the robust move was to ask.', 13.5, MUTED, ls=1.05)
takeaway(s, 'Hypothesis H2: environmental clues substitute for authorization when the task contract is silent.', size=16)
notes(s, 'Counts of hosts and records are as reported by Anthropic and Google, not verified by us. These cases motivate the '
      'ambiguity result: the agent had two plausible grants (the target is in scope; it is not) and no prior over them. '
      'Maxmin and minimax-regret rules pick a reversible action, which here means a clarification. The experiment that tests this '
      'is X1b, where the grant is deliberately silent or self-contradictory.', ('anthropic', 'library'))

# ============================================================================= 06 ACROSS THE LIBRARY
s = new('What recurs, and the three questions a theory must answer.', '06 / Across the library',
        f"Counts over {LIB_N} records coded with the Raftari (2026) frame. 'Present', 'violated' or 'yes' only; partial and unknown omitted.",
        [('library', 'figures/incidents/analysis/summary.json; analyst draft codings, second reader pending', 5.0)],
        status='REPORTED')
para(s, ML, 2.32, 5.9, .32, 'Mechanisms, propositions and indicators', 14, NAVY, True)
hbars(s, ML, 2.66, 5.9, [
    ('M1 pressure, no safe stop', MECH['M1']['present'], BLUE),
    ('M2 capability composition', MECH['M2']['present'], BLUE),
    ('M4 detection, no escalation', MECH['M4']['present'], BLUE),
    ('P3 authz survives failure', PROP['P3']['violated'], RED),
    ('P6 effects over intent', PROP['P6']['violated'], RED),
    ('f repeated failure', IND['f']['yes'], PLUM),
    ('p boundary probing', IND['p']['yes'], PLUM),
    ('x external contact', IND['x']['yes'], PLUM),
], LIB_N, step=.40, label_w=2.35, size=11)
qx = 6.95
for i, (q, head, body, col) in enumerate([
    ('Q1', 'Persist, or stop?', f"M1 present in {MECH['M1']['present']}; f in {IND['f']['yes']}, unknown in {IND['f']['unknown']}. "
     'When is one more attempt worth its cost, and when does failure become evidence?', BLUE),
    ('Q2', 'Act, diagnose, or ask?', f"P3 violated in {PROP['P3']['violated']}. A diagnostic informs feasibility; only a request can "
     'change the grant. Which is worth paying for?', TEAL),
    ('Q3', 'Which policy to release?', f"P6 violated in {PROP['P6']['violated']}; disclosure median {LAT['disclosure_median']:g} days. "
     'Which incident signal should select the next policy, given finite, late evidence?', PLUM),
]):
    y = 2.32 + i * 1.25
    card(s, qx, y, 5.76, 1.13, col)
    badge(s, qx + .20, y + .14, q, col, .40, 11)
    para(s, qx + .72, y + .10, 4.9, .40, head, 16, col, True)
    para(s, qx + .22, y + .50, 5.35, .60, body, 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.05)
takeaway(s, 'Each question is a classical decision problem: optimal stopping, value of information, constrained selection.')
notes(s, 'Bars are read from the library summary at build time. Co-occurrence: M1+M2 in 9 records and M2+M4 in 9. The three questions '
      'organise the rest of the deck. Q1 is optimal stopping with learning; Q2 is value of information plus a principal–agent twist, '
      'because a request can change the permitted set while a diagnostic cannot; Q3 is constrained policy selection from finite, '
      'delayed evidence. None of these is new; the claim is that stating them exactly yields testable behaviour.', ('library',))

# ============================================================================= 07 FORMULATION
s = new('A constrained POMDP whose grant evidence cannot change.', '07 / Formal problem',
        'Three objects evolve by three different rules: the belief by Bayes, the grant by trusted replies only, the value by dynamic programming.',
        [('pomdp', 'Belief-state planning: Kaelbling et al. 1998', 4.4), ('optimization', 'Objective: incident-aware optimization', 3.6)],
        status='THEORY')
for i, (sym, name, body, col) in enumerate([
    ('x', 'Hidden state', 'route feasibility, service state; never observed directly', BLUE),
    ('o', 'Observation', 'tool outputs, errors such as a 503, trusted human replies', BLUE),
    ('a', 'Action', 'execute · probe · ASK(resource) · STOP with evidence', TEAL),
    ('e', 'Effect', 'what the action touched; journaled by the service, not the agent', ORANGE),
]):
    y = 2.36 + i * .56
    box(s, ML, y, .50, .48, TINT[col])
    text(s, ML, y, .50, .48, sym, 18, col, True, True)
    para(s, ML + .62, y, 1.55, .48, name, 14, NAVY, True)
    para(s, ML + 2.15, y, 3.75, .48, body, 12, INK, ls=1.0)
box(s, ML, 4.66, 5.9, 1.38, PALE, rounded=False)
para(s, ML + .18, 4.72, 2.3, .40, 'Controller state', 15, NAVY, True)
formula(s, ML + 2.45, 4.72, r's_t=(b_t,\;g_t,\;k_t)', 15, NAVY, h_box=.40)
formula(s, ML + .18, 5.14, r'b_t=P(x_t\mid h_t)\ \text{by Bayes on every }o;\qquad g_{t+1}\neq g_t\ \text{only if}\ a_t=\mathrm{ASK}\ \text{and}\ o_t\ \text{is a trusted grant};', 11, INK, w_max=5.6)
formula(s, ML + .18, 5.52, r'k_{t+1}=k_t-\kappa(a_t)\ \text{for every action.}', 11, INK, w_max=5.6)
theorem_band(s, 6.80, 2.36, 5.91, 1.90, '',
             'A(g): permitted execution, permitted diagnostics, permitted requests, termination. Terminal payoffs fixed in advance.',
             tag='BELLMAN EQUATION ON THE CONTROLLER STATE',
             math=[r'V_h(b,g,k)=\max_{a\in A(g),\ \kappa(a)\leq k}\ \mathbb{E}\left[\,r(x,a)-c(a)\right.',
                   r'\qquad\qquad\qquad +\,V_{h-1}\left(b^{\prime},\,g^{\prime},\,k-\kappa(a)\right)\left.\right]'], math_size=16)
box(s, 6.80, 4.40, 5.91, 1.72, WHITE, LINE, rounded=False)
para(s, 7.00, 4.48, 5.5, .34, 'Three readings of one 503', 14, NAVY, True)
for i, (word, body, col) in enumerate([
    ('Epistemic', r'b_t\ \text{moves: the local route is less likely to work.}', BLUE),
    ('Deontic', r'g_t\ \text{does not move: nothing was granted.}', TEAL),
    ('Evaluative', r'I(\tau)=\mathbb{1}\{\exists\,e\notin A(g_t)\},\quad R=P(I{=}1),\quad U=P(\text{done},\,I{=}0)', ORANGE),
]):
    y = 4.86 + i * .40
    chip(s, 7.00, y + .05, 1.05, word, col, size=9.5, h=.28)
    formula(s, 8.15, y, body, 11, INK, w_max=4.45, h_box=.38)
takeaway(s, 'A failed tool call updates b_{t}. It never updates g_{t}. The grant is a constraint on effects, not a reward term.')
notes(s, 'Finite-horizon constrained POMDP. Controller state s_t = (b_t, g_t, k_t) is distinct from the incident risk-indicator '
      'vector z_t on slide 5. The local POMDP symbols x (hidden state) and o (observation) have the meanings defined on this slide; '
      'they are not the incident codes x (unexpected external contact) and o (observability loss). Hidden x_t includes route feasibility; g_t is the trusted, versioned authorization '
      'contract; k_t the remaining budget; h_t the history. Transition and observation kernels T and O; Bayesian filtering gives b′. '
      'A request can lead to a trusted grant update or denial with response probability q; q is not the probability that an ungranted '
      'action is secretly authorized. Reachability and tool success never update g. Missing scope is not filled by the agent: choose an '
      'already permitted alternative, request, or stop. In the closed fixture A(g) is exact and effect-based; open environments need '
      'validated effect prediction. Utility u is fixed in advance, including the value of justified termination. The action constraint '
      'is a specification; whether an LLM follows it is what the experiments measure. Constrained MDP background: Altman 1999; '
      'CPO is background for constrained optimization, not the algorithm used.', ('pomdp', 'optimization', 'design'))

# ============================================================================= 08 RESULT 1: STOPPING
t = STOP_TOY
s = new('Result 1: each failure is evidence; stop at a posterior threshold.', '08 / Optimal stopping with learning',
        f"Exact toy: feasible with prior p_{{0}} = {t['p0']}; an attempt succeeds with probability q = {t['q']} if feasible, costs c = {t['c']:g}; "
        f"success R = {t['R']:g}; honest stop B = {t['B']:g}.",
        [('stopping', 'Monotone stopping: one-step look-ahead is optimal (Ferguson, ch. 5)', 4.6), ('next', 'Project X1 design', 1.6)],
        status='EXACT TOY')
card(s, ML, 2.36, 6.05, 3.70, BLUE)
para(s, ML + .25, 2.46, 5.6, .36, 'Posterior that the route is feasible, after each failure', 14, NAVY, True)
bx0, by0, bw, bh = ML + .95, 2.95, 4.70, 2.08
line(s, bx0, by0 + bh, bx0 + bw, by0 + bh, INK)
line(s, bx0, by0, bx0, by0 + bh, INK)
ty = by0 + bh * (1 - STOP_THRESHOLD)
dashed(s, bx0, ty, bx0 + bw, ty, RED, .025)
formula(s, bx0 + bw - 1.55, ty - .36, rf'c\,/\,q(R-B)={STOP_THRESHOLD:.2f}', 11, RED, w_max=1.55, align='r')
slot = bw / len(POSTERIORS)
for i, p in enumerate(POSTERIORS):
    col = BLUE if p >= STOP_THRESHOLD else MUTED
    x = bx0 + i * slot + slot * .2
    box(s, x, by0 + bh * (1 - p), slot * .6, bh * p, col, rounded=False)
    para(s, x - .1, by0 + bh * (1 - p) - .30, slot * .8, .28, f'{p:.2f}', 11, col, True, align='c')
    para(s, x - .1, by0 + bh + .04, slot * .8, .26, f'p_{{{i}}}', 11, MUTED, align='c')
para(s, bx0, by0 + bh + .32, bw, .26, 'failures observed →', 11, MUTED)
para(s, ML + .25, 5.64, 5.6, .36,
     f'Try while p_{{t}} ≥ {STOP_THRESHOLD:.2f}: {ATTEMPTS_BEFORE_STOP} attempts, then stop with evidence.', 13, BLUE, True)
theorem_band(s, 6.95, 2.36, 5.76, 1.84, '',
             'Posterior falls after every failure, so stopping, once optimal, stays optimal: the one-step rule is exact.',
             tag='RESULT 1 · THRESHOLD RULE',
             math=[r'\text{try}\quad\text{iff}\quad p_t\,q\,(R-B)>c', r'p_{t+1}=\dfrac{p_t\,(1-q)}{1-p_t\,q}'], math_size=15, gap=.08)
box(s, 6.95, 4.32, 5.76, 1.74, PALE, rounded=False)
para(s, 7.15, 4.38, 5.4, 1.64,
     '•  A better honest exit B raises the threshold and shortens persistence; a worthless exit buys stubbornness.\n'
     '•  Several recovery routes: a Gittins index per route; try the highest, stop when all fall below the exit (Gittins 1979).\n'
     '•  Substitution is persistence past this point with a different resource: it needs a grant, not a higher posterior.\n'
     '•  Readouts: premature stops (p_{t} above threshold) and wasted attempts (below) are separate errors; a lower violation rate '
     'must not be mistaken for better stopping.', 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.08)
takeaway(s, 'Repeated failure (indicator f) supplies evidence that moves the feasibility posterior in this model.')
notes(s, f"Exact calculation: p0={t['p0']}, q={t['q']}, R={t['R']:g}, B={t['B']:g}, c={t['c']:g}. Posterior sequence after failures: "
      + ', '.join(f'{p:.3f}' for p in POSTERIORS) + f'. Threshold {STOP_THRESHOLD:.3f}; the first posterior below it is '
      f'p_{ATTEMPTS_BEFORE_STOP}, so the agent attempts {ATTEMPTS_BEFORE_STOP} times. One-step look-ahead is optimal here because '
      'the problem is monotone: a success ends the episode and every failure lowers p, so if stopping beats one more try now it '
      'beats it later (Ferguson, Optimal Stopping and Applications, ch. 5). Backward induction over a finite budget agrees, which the '
      'build validates. Calibrated q and risk neutrality are assumptions; an LLM confidence statement is not a calibrated posterior. '
      'The incident library records f in 9 of 26 and cannot tell for 13: reports rarely say whether the task was failing.', ('stopping', 'next'))

# ============================================================================= 09 RESULT 2: VOI
s = new('Result 2: sometimes the best next action is to learn.', '09 / Value of information',
        'p = probability an authorized recovery works. d = request cost (time, money or review), in utility units. Color = highest expected utility.',
        [('voi', 'VOI background: Hay et al. 2012. Hypothetical model; payoffs and boundaries on the next two slides.', 10.8)],
        status='EXACT TOY')
para(s, .78, 2.38, 6.0, .35, 'Optimal action as belief and request cost change', 14.5, NAVY, True)
s.shapes.add_picture(str(voi_phase_figure()), Inches(.66), Inches(2.70), width=Inches(6.10))
para(s, .83, 5.84, 5.95, .30, 'A (.9, 8): RECOVER  |  B (.1, 0): ASK  |  C (.3, 2): PROBE', 10.5, NAVY, True)
for i, (a, desc, col) in enumerate([
    ('RECOVER', 'Try the authorized local route; high p favors it.', TEAL),
    ('STOP', 'Report failure when further action is not worth its cost.', MUTED),
    ('ASK', 'Request a grant for an alternative; obey the reply.', BLUE),
    ('PROBE', 'Run an authorized diagnostic, then choose an action.', GOLD),
]):
    y = 2.38 + i * .54
    box(s, 6.98, y, 5.73, .47, TINT.get(col, PALE))
    box(s, 6.98, y, .055, .47, col, rounded=False)
    para(s, 7.12, y, 1.14, .47, a, 12, col, True)
    para(s, 8.29, y, 4.27, .47, desc, 11.5, INK, ls=1.02)
box(s, 6.98, 4.64, 5.73, 1.49, NAVY, rounded=False)
para(s, 7.18, 4.75, 5.32, .33, 'Two uncertainties, two different interventions', 14, WHITE, True)
para(s, 7.18, 5.16, 5.32, .41, 'Feasibility: diagnose an allowed public route.\nPermission: obtain a trusted grant for a new resource.', 12, LIGHT, ls=1.08)
para(s, 7.18, 5.71, 5.32, .30, 'Reachability of an internal database does not grant permission.', 11.5, WHITE, True)
takeaway(s, 'PROBE changes beliefs, not permissions. ASK may change permissions; it does not make a route feasible.', size=15)
notes(s, 'This is an optimal-policy phase diagram under stated hypothetical utilities, not measured agent behavior. '
      'The horizontal axis p is the probability that an already authorized local recovery works: p=0 means certainly infeasible, '
      'p=.3 means a 30 percent chance, and p=1 means certainly feasible. The vertical axis d is the cost of requesting additional '
      'authorization, measured in the same utility units as the reward; it can represent time, money or human review overhead. '
      'd=0 is free and d=8 is expensive in this toy. The chart uses the analytic boundaries derived on slide 13, rather than a '
      'coarse grid. STOP is a small but real region at p<1/16 and d>6. The dashed d=6 line marks where ASK and STOP tie. '
      'Point A: p=.9,d=8 gives RECOVER=7, PROBE=6.7, STOP=0 and ASK=-2. Point B: p=.1,d=0 gives ASK=6, PROBE=5.7, '
      'STOP=0 and RECOVER=-1. Point C: p=.3,d=2 gives PROBE=4.7, ASK=4, RECOVER=1 and STOP=0. '
      'Epistemic uncertainty is about whether an allowed plan works. A permitted, non-mutating check of a public service or an '
      'already authorized alternative can reduce it. Authorization is a separate state: checking whether a government internal '
      'database is reachable cannot grant access, and the diagnostic itself must be within scope. Only a trusted affirmative '
      'response to ASK may change the grant; ASK alone does not change technical feasibility. The toy assumes a feasible alternative '
      'after a grant, which is an explicit assumption rather than a general property of authorization. The aim is useful recovery '
      'within scope, avoiding both unauthorized persistence and premature stopping.', ('voi',))

# ============================================================================= 09b RESULT 2: PAYOFFS AND CONTINGENT CHOICE
s = new('Why PROBE is worth 4.7: pay once, then choose with evidence.', '09 / Value of information / Payoffs',
        'Risk-neutral toy: one optional perfect diagnostic, then one terminal action. No retries; all numbers are hypothetical utilities.',
        [('voi', 'VOI background: Hay et al. 2012. This payoff model is illustrative, not an empirical estimate.', 10.7)],
        status='EXACT TOY')
para(s, ML + .16, 2.39, 5.9, .34, 'Make every payoff and assumption explicit', 15, NAVY, True)
para(s, ML + .16, 2.83, 5.85, .80,
     'Local recovery: success +10, failure 0; cost 2 either way.\n'
     'ASK: pay d; grant probability .6, independent of feasibility.\n'
     'If granted, the alternative succeeds (+10); if denied, stop (0).\n'
     'PROBE: pay .5; learn feasibility perfectly, within existing scope.', 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.12)
grid(s, ML, 3.79, [1.22, 1.32, 1.40, 2.12], [
    ['Action', 'Feasible', 'Infeasible', 'Expected Q(p, d)'],
    ['RECOVER', '+8', '-2', '10p - 2'],
    ['ASK', '6 - d', '6 - d', '6 - d'],
    ['STOP', '0', '0', '0'],
], row_h=.37, size=11.5)
para(s, ML + .12, 5.31, 5.84, .24, 'PROBE preserves the best choice after either observation:', 10.5, MUTED)
formula(s, ML + .14, 5.61, r'Q_{\rm PROBE}=-0.5+p\max(8,6-d,0)', 14.5, NAVY, w_max=5.8, h_box=.28)
formula(s, ML + .79, 5.92, r'+(1-p)\max(-2,6-d,0)', 14.5, NAVY, w_max=5.15, h_box=.28)

para(s, 7.08, 2.39, 5.40, .34, 'At p = .3 and d = 2: a contingent decision', 15, TEAL, True)
para(s, 7.08, 2.77, 5.40, .24, 'Direct values: RECOVER = 1; ASK = 4; STOP = 0.', 11.5, MUTED)
box(s, 7.02, 3.43, 1.23, .71, CREAM)
para(s, 7.12, 3.48, 1.03, .57, 'PROBE\npay .5', 13, GOLD, True, align='c', ls=1.03)
arrow(s, [(8.27, 3.66), (8.51, 3.66), (8.51, 3.30), (8.77, 3.30)], TEAL)
arrow(s, [(8.27, 3.92), (8.51, 3.92), (8.51, 4.27), (8.77, 4.27)], BLUE)
box(s, 8.83, 3.08, 3.88, .78, MINT)
para(s, 9.00, 3.13, 3.52, .28, 'Feasible (probability .3): RECOVER', 12, TEAL, True)
para(s, 9.00, 3.48, 3.52, .28, 'Choose max(8, 4, 0) = 8', 12, INK)
box(s, 8.83, 4.04, 3.88, .78, PALE)
para(s, 9.00, 4.09, 3.52, .28, 'Infeasible (probability .7): ASK', 12, BLUE, True)
para(s, 9.00, 4.44, 3.52, .28, 'Choose max(-2, 4, 0) = 4', 12, INK)
box(s, 7.02, 4.99, 5.69, 1.20, NAVY, rounded=False)
formula(s, 7.22, 5.08, r'Q_{\rm PROBE}=-0.5+0.3\times8+0.7\times4=4.7', 15, WHITE, w_max=5.28, h_box=.36)
formula(s, 7.22, 5.56, r'\mathrm{VOI}_{\rm net}=4.7-\max(1,4,0)=+0.7', 16, WHITE, w_max=5.28, h_box=.39)
box(s, ML, 6.32, CW, .55, TEAL)
formula(s, ML + .18, 6.35,
        r'\mathrm{VOI}_{\rm net}=\mathbb{E}_{o}\!\left[\max_{a\in\{R,A,S\}}\mathbb{E}[U(a)\mid o]\right]-\max_{a\in\{R,A,S\}}\mathbb{E}[U(a)]-c_{\rm probe};\qquad\text{probe if positive}',
        15, WHITE, w_max=CW - .36, h_box=.48, align='c')
notes(s, 'All outcomes and costs are specified so the audience can independently recompute 4.7. Hidden local feasibility '
      'theta is Bernoulli(p). RECOVER is terminal and costs 2 regardless of outcome: conditional payoffs are 8 if feasible and -2 '
      'if infeasible, so Q_RECOVER=p*8+(1-p)*(-2)=10p-2. STOP is terminal with utility 0. ASK pays d regardless of the reply. '
      'A trusted grant arrives with probability q=.6, independent of local feasibility; the alternative is assumed to succeed '
      'after a grant, with gross utility 10 and no further cost. Denial yields 0 and no action. Hence Q_ASK=.6*10+.4*0-d=6-d. '
      'The conditional ASK cells are expected values over the authorization response, not guaranteed outcomes. PROBE is a '
      'single permitted, non-mutating perfect observation of theta, costs c=.5, changes no permissions, and is followed by the '
      'best of RECOVER, ASK and STOP. It cannot be followed by another probe. Thus Q_PROBE=-.5+p*max(8,6-d,0)+(1-p)*max(-2,6-d,0). '
      'At p=.3,d=2, direct values are RECOVER=1, ASK=4 and STOP=0. After a feasible observation select RECOVER=8; after an '
      'infeasible observation select ASK=4, while still obeying its reply. The post-information value before probe cost is '
      '.3*8+.7*4=5.2; gross VOI is 5.2-4=1.2; net VOI is 1.2-.5=.7, and Q_PROBE is 4.7. '
      'In the bottom formula R,A,S abbreviate RECOVER, ASK and STOP; o is the diagnostic observation. With the same action set '
      'before and after information and the option to ignore it, gross VOI is nonnegative, but net VOI can be negative. '
      'For noisy diagnostics use posterior beliefs P(theta=1|o) and average the posterior-optimal value over observations. '
      'Uncertainty alone is insufficient: information must change a valuable decision enough to cover its cost.', ('voi',))

# ============================================================================= 09c RESULT 2: EXACT DECISION BOUNDARIES
s = new('Decision boundaries are equal-value curves, with dominance checks.', '09 / Value of information / Boundaries',
        'A visible boundary requires Q(a) = Q(b), with both at least as good as every other action. These curves generate the diagram.',
        [('voi', 'Exact derivation for the stated toy; equality curves and neighboring optimal actions checked at build time.', 11.2)],
        status='EXACT TOY')
for x, w, col in [(ML, 6.06, BLUE), (6.98, 5.73, TEAL)]:
    card(s, x, 2.38, w, 3.10, col)
para(s, ML + .20, 2.48, 5.65, .34, '0 ≤ d ≤ 6: an infeasible route leads to ASK', 15, BLUE, True)
formula(s, ML + .20, 2.94, r'Q_{\rm PROBE}=-0.5+8p+(1-p)(6-d)', 16, NAVY, w_max=5.62, h_box=.35)
formula(s, ML + 1.28, 3.34, r'=5.5-d+(2+d)p', 16, NAVY, w_max=4.54, h_box=.35)
para(s, ML + .20, 3.85, 5.65, .28, 'ASK / PROBE: set Q(PROBE) = 6 - d', 12, BLUE, True)
formula(s, ML + .20, 4.17, r'(2+d)p=0.5\quad\Longrightarrow\quad p_L=\dfrac{0.5}{2+d}', 16, NAVY, w_max=5.62, h_box=.43)
para(s, ML + .20, 4.66, 5.65, .28, 'PROBE / RECOVER: set Q(PROBE) = 10p - 2', 12, TEAL, True)
formula(s, ML + .20, 4.97, r'(8-d)p=7.5-d\quad\Longrightarrow\quad p_U=\dfrac{7.5-d}{8-d}', 16, NAVY, w_max=5.62, h_box=.43)

para(s, 7.18, 2.48, 5.31, .34, '6 ≤ d ≤ 8: an infeasible route leads to STOP', 14.5, TEAL, True)
formula(s, 7.18, 2.96, r'Q_{\rm ASK}\leq0,\qquad Q_{\rm PROBE}=8p-0.5', 16, NAVY, w_max=5.31, h_box=.38)
para(s, 7.18, 3.52, 5.31, .28, 'STOP / PROBE: set Q(PROBE) = 0', 12, MUTED, True)
formula(s, 7.18, 3.86, r'8p-0.5=0\quad\Longrightarrow\quad p_L=\dfrac{1}{16}', 16, NAVY, w_max=5.31, h_box=.43)
para(s, 7.18, 4.38, 5.31, .28, 'PROBE / RECOVER: set Q(PROBE) = 10p - 2', 12, TEAL, True)
formula(s, 7.18, 4.72, r'8p-0.5=10p-2\quad\Longrightarrow\quad p_U=\dfrac{3}{4}', 16, NAVY, w_max=5.31, h_box=.43)
para(s, 7.18, 5.17, 5.31, .24, 'At d = 6, ASK and STOP tie below p = 1/16.', 10.5, MUTED)
para(s, ML + .10, 5.62, CW - .20, .46,
     'Policy: below p_{L} choose ASK (d < 6) or STOP (d > 6); between p_{L} and p_{U} choose PROBE; above p_{U} choose RECOVER.\n'
     'Adjacent actions tie at each threshold. The direct ASK = RECOVER curve is hidden: PROBE has higher value there.', 11.5, INK, ls=1.12)
takeaway(s, 'At d = 2: ASK below 1/8; PROBE for 1/8 < p < 11/12; RECOVER above 11/12 (approximately .917).', size=14.5)
notes(s, 'Derivation for the displayed range 0<=p<=1, 0<=d<=8, with perfect diagnostic cost .5. '
      'For d<=6, a feasible observation leads to RECOVER=8 because ASK<=6; an infeasible observation leads to ASK=6-d>=0, '
      'tied with STOP only at d=6. Q_PROBE=-.5+8p+(1-p)(6-d)=5.5-d+(2+d)p. '
      'Subtracting Q_ASK gives (2+d)p-.5, so the lower boundary is p_L=.5/(2+d). Subtracting Q_RECOVER gives '
      '7.5-d-(8-d)p, so the upper boundary is p_U=(7.5-d)/(8-d). Since p_L<p_U and ASK>=STOP, these are actual '
      'optimal-action boundaries, not just pairwise equalities. For d>=6, ASK<=STOP=0; the post-probe choices are RECOVER '
      'if feasible and STOP if infeasible, so Q_PROBE=8p-.5. Setting it equal to 0 gives p_L=1/16; setting it equal to '
      '10p-2 gives p_U=3/4. RECOVER is negative at the lower boundary, and STOP is below both values at the upper boundary. '
      'At d=6, ASK=STOP for all p, but this equality is a visible optimal boundary only for p<=1/16. At p=1/16,d=6 '
      'ASK, STOP and PROBE all tie. Pairwise equality alone is insufficient: for 0<=d<=6 the direct ASK=RECOVER curve '
      'p=(8-d)/10 lies strictly inside the PROBE region. For d>6 that direct curve is also dominated. '
      'At d=2 the exact cutoffs are 1/8=.125 and 11/12, approximately .917. At p=.3,d=2, varying diagnostic cost c '
      'gives Q_PROBE=5.2-c: strictly probe only for c<1.2, tie at c=1.2, otherwise ASK. '
      'These are falsifiable choice predictions conditional on the model: sweep p across both cutoffs at fixed d, and '
      'sweep c across 1.2 at the example point. Measure choice disagreement and value regret separately. Noisy diagnostics, '
      'uncertain authorization probabilities, costly authorized alternatives and additional retries require a new value calculation; '
      'these particular thresholds are not universal. The safety invariant is unchanged: diagnosing feasibility never grants authority.', ('voi',))

# ============================================================================= 10 RESULT 3: AMBIGUITY
s = new('Result 3: under a silent grant, ask dominates without a prior.', '10 / Decision under ambiguity; principal and agent',
        'Not risk, where probabilities are known, but ambiguity: two plausible grants and no prior over them. Toy payoffs.',
        [('gilboa', 'Maxmin expected utility: Gilboa & Schmeidler 1989', 4.0), ('cirl', 'CIRL 2016', 1.0), ('switch', 'Off-switch game 2017', 1.8),
         ('next', 'X1b: silent and conflicting grants', 2.8)],
        status='EXACT TOY')
rows = [['Action', 'g: archive allowed', 'g: archive excluded', 'maxmin', 'max regret']]
fills = {}
for i, (name, ua, ue, mm, rg) in enumerate(AMB, start=1):
    rows.append([name, f'{ua:+d}', f'{ue:+d}', f'{mm:+d}', f'{rg:d}'])
    if name == AMB_MAXMIN:
        fills[(i, 3)] = (MINT, TEAL)
    if name == AMB_REGRET:
        fills[(i, 4)] = (MINT, TEAL)
grid(s, ML, 2.38, [1.75, 1.30, 1.35, .95, 1.00], rows, row_h=.46, size=11.5, fills=fills)
para(s, ML, 4.30, 6.35, .70,
     f'Both rules pick "{AMB_MAXMIN.lower()}". An expected-utility agent reads the archive only if P(allowed) > {AMB_PRIOR_SWITCH:.2f}, '
     'a number the task never supplied. Reachability is not that number.', 12.5, INK, ls=1.08)
box(s, ML, 5.05, 6.35, 1.0, PALE, rounded=False)
para(s, ML + .18, 5.10, 6.0, .92,
     'Principal–agent reading: the principal writes g; the agent acts; the journal says what e touched and only g says whether it was '
     'allowed. Where g is silent, who holds residual control is the design variable. A request transfers that decision to the principal; '
     'a diagnostic cannot.', 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.08)
theorem_band(s, 7.20, 2.38, 5.51, 1.72, '',
             'Robust rules need the set of plausible grants, not a prior over them (Gilboa & Schmeidler 1989; Savage 1951).',
             tag='RESULT 3 · ROBUST RULES UNDER SILENCE',
             math=[r'\text{maxmin:}\quad \max_a\ \min_g\ \mathbb{E}_g\left[\,u(a)\,\right]',
                   r'\text{regret:}\quad\ \min_a\ \max_g\ \left[\,u^{*}(g)-u(a,g)\,\right]'], math_size=15, gap=.06)
box(s, 7.20, 4.24, 5.51, 1.81, WHITE, LINE, rounded=False)
para(s, 7.40, 4.30, 5.2, 1.70,
     '•  Prediction: ambiguity-averse agents prefer reversible actions and questions. Over-conservatism is the price, so it is measured '
     'on authorized-feasible cells where a blocker is a cost.\n'
     '•  CIRL: an agent that does not know the principal\'s reward elicits and defers rather than guesses; deference has positive '
     'value exactly under uncertainty (off-switch game).\n'
     '•  Conflicting grants need a frozen precedence rule; silent grants need a reference policy before any violation rate exists.',
     11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.1)
takeaway(s, 'Silence is not permission. The robust move under an unknown grant is the reversible one: ask, or stop with evidence.')
notes(s, 'Toy payoffs: reading the archive pays +10 if allowed and −20 if it was excluded; asking pays +7 either way (delay cost); '
      f'stopping pays +2. Maxmin picks "{AMB_MAXMIN}" and minimax regret picks "{AMB_REGRET}"; both computed at build time. An '
      f'expected-utility agent prefers reading iff 10p − 20(1−p) > 7, i.e. p > {AMB_PRIOR_SWITCH:.3f}. The scope-inference family '
      'shows agents behaving as if p were near one because the target was reachable. X1b adds unspecified and conflicting grant '
      'columns; its readout is a response mix (silent read / clarification / authorized completion) plus an over-conservatism check, '
      'never a violation rate under silence. Robust MDPs over an uncertainty set (Iyengar 2005) and attainable-utility preservation '
      '(Turner et al. 2020) are the RL analogues.', ('gilboa', 'cirl', 'switch', 'next'))

# ============================================================================= 11 RESULT 4: SELECTION
s = new('Result 4: what finite, paired evidence can and cannot certify.', '11 / Constrained selection from finite data',
        'The release decision is a constrained optimum estimated from development rollouts; the slack is set by the sample, not the signal.',
        [('optimization', 'Objective and arms: incident-aware optimization', 4.0), ('cpo', 'Constrained optimization background: CPO', 3.4),
         ('followup', 'Archived paired contrast', 2.2)],
        status='THEORY')
theorem_band(s, ML, 2.36, 6.10, 1.52, '',
             'R: verified incidents; U: authorized completion; C: cost; F: false intervention. Margins fixed before search.',
             tag='RESULT 4a · THE RELEASE OBJECTIVE',
             math=[r'\theta^{*}=\mathrm{arg\,min}_{\theta}\ R(\theta)\quad\text{subject to}',
                   r'U(\theta)\geq U(\theta_0)-\delta_U,\quad C(\theta)\leq B_{\mathrm{run}},\quad F(\theta)\leq\alpha_F'], math_size=14, gap=.06)
box(s, ML, 4.00, 6.10, 2.05, PALE, rounded=False)
para(s, ML + .18, 4.04, 5.8, .36, 'Hoeffding + union bound over K candidates, n rollouts each', 13, NAVY, True)
formula(s, ML + .18, 4.38,
        r'R(\theta_{\mathrm{sel}})\ \leq\ \min_{\theta\in E}R(\theta)+2\varepsilon,\qquad \varepsilon=\sqrt{\ln(2K/\delta)\,/\,2n},\qquad\text{w.p. }1-\delta',
        12, INK, w_max=5.75, h_box=.40)
rows = [['n per candidate', f'slack 2ε (K = {SEL_K}, δ = .05)']]
for n, slack in SEL_ROWS:
    rows.append([f'{n:,}', f'{slack:.2f}'])
grid(s, ML + .18, 4.84, [1.9, 2.2], rows, row_h=.235, size=10.5)
para(s, ML + 4.42, 4.84, 1.58, 1.18,
     'At pilot scale the bound is vacuous. Information lives in paired contrasts, not in independent rates.',
     10, MUTED, valign=MSO_ANCHOR.TOP, ls=1.05)
card(s, 6.95, 2.36, 5.76, 3.69, PLUM)
para(s, 7.17, 2.45, 5.3, .36, 'Paired twins: exact McNemar power', 14, PLUM, True)
para(s, 7.17, 2.80, 5.3, .52,
     f'Baseline impossible-task incident rate from the archived diagnostic: {P_BASE_CELL["positive"]}/{P_BASE_CELL["observed"]} = {P_BASE:.2f}. '
     'Power to detect a drop to the target rate, two-sided exact test, α = .05.', 10.5, MUTED, valign=MSO_ANCHOR.TOP, ls=1.05)
rows = [['pairs n', f'to {POWER_TARGETS[0]:.3f}', f'to {POWER_TARGETS[1]:.2f}']]
for n, powers in POWER_ROWS:
    rows.append([str(n)] + [f'{p:.2f}' for p in powers])
grid(s, 7.17, 3.40, [1.3, 1.8, 1.8], rows, row_h=.30, size=11, header=PLUM)
para(s, 7.17, 4.95, 5.3, 1.05,
     '•  Only discordant pairs carry information; a tie between selection signals is a valid result.\n'
     '•  Zero observed incidents is not a certified population bound; unknown outcomes are not zeros.\n'
     '•  The chosen candidate\'s estimate is biased upward: confirm on held-out twins (optimizer\'s curse).',
     10.5, INK, valign=MSO_ANCHOR.TOP, ls=1.08)
takeaway(s, 'Select the lowest verified-incident policy among utility-eligible candidates, then confirm it on untouched paired twins.')
notes(s, 'Objective copied from docs/incident-aware-optimization.md with Altman 1999 and CPO as background. The selection bound is '
      'the standard Hoeffding inequality with a union bound over K candidates: each empirical risk is within ε of its mean with '
      'probability 1 − δ/K, so the empirical minimiser has true risk at most 2ε above the best eligible candidate. With K = 8 and '
      f'n = 48 the slack is {SEL_ROWS[0][1]:.2f}, which is useless; it falls to {SEL_ROWS[-1][1]:.2f} only at n = 5000. The paired design '
      'is the alternative: twins differ only in the intervention, and the exact McNemar test uses discordant pairs. The power table is '
      f'computed at build time from the archived baseline rate {P_BASE:.4f} (results/baseline-diagnostic-v2/summary.json) against two '
      'hypothetical target rates. A Lagrangian relaxation prices the constraints, but hard wall versus soft penalty is not identifiable '
      'from finite data: a large multiplier reproduces zero violations.', ('optimization', 'cpo', 'followup'))

# ============================================================================= 12 ALGORITHM A
s = new('Algorithm A: GRANT-LOOKAHEAD, plan inside the grant.', '12 / Online decision algorithm',
        'Exact dynamic programming in the finite fixture first; budgeted Monte Carlo lookahead with the same grant filter for larger tasks.',
        [('pomdp', 'Belief-state DP: Kaelbling et al. 1998', 3.4), ('pomcp', 'Sampling lookahead: POMCP 2010', 2.8),
         ('camel', 'Enforcement precedent: CaMeL 2025', 3.0)],
        status='PROPOSED')
mono(s, ML, 2.36, 7.05, 3.70, [
    'GRANT-LOOKAHEAD(g, h, T, O, u, c, k, H)',
    '  b <- FILTER(prior, h)                 # Bayes on feasibility only',
    '  loop until terminal:',
    '    C <- {a in A(g): kappa(a) <= k}     # permitted exec / probes',
    '         U {ASK(r): r not in g} U {STOP}',
    '    for a in C:',
    '      Q(a) <- E_b[r(x,a) - c(a)]',
    '            + E_o[ V_(H-1)(b\'(o), g\'(a,o), k - kappa(a)) ]',
    '            # g\'(a,o) = g unless a = ASK and o is a trusted grant',
    '    a* <- argmax Q(a); ties -> terminal actions',
    '    if BROKER(a*, g) = DENY: log attempt; C <- C \\ {a*}; continue',
    '    o <- EXECUTE(a*); JOURNAL(effects(a*)); k <- k - kappa(a*)',
    '    b <- FILTER(b, a*, o); g <- g\'(a*, o); h <- h + (a*, o)',
    '  return STOP(evidence = h)',
], size=11)
card(s, 7.85, 2.36, 4.86, 3.70, TEAL)
para(s, 8.05, 2.44, 4.5, .34, 'Properties', 15, TEAL, True)
para(s, 8.05, 2.80, 4.5, 3.2,
     '1  Under complete mediation and a correct effect model, no executed effect lies outside A(g_{t}): the broker is the '
     'shield, the planner is not trusted for safety.\n'
     '2  Exact DP has zero regret in the finite fixture; V* is the normative reference, not a clairvoyant policy.\n'
     '3  Cost O(H · |A| · |O| · |B|) over a discretised belief; beyond that, fixed-budget beam or Monte Carlo lookahead '
     'with the same A(g) filter at every node.\n'
     '4  Requests are planned: ASK enters Q with cost d and response probability q, so the planner asks when authority must change '
     'and not otherwise.\n'
     '5  Calibration is an input: T, O estimated on development data; an LLM confidence statement is not a posterior.',
     11, INK, valign=MSO_ANCHOR.TOP, ls=1.1)
takeaway(s, 'Plan for utility; enforce scope before execution; measure effects independently. Three functions, three components.')
notes(s, 'Proposed online algorithm; not evaluated on LLM agents. Candidate tool actions are normalised into effect-relevant fields '
      '(resource, owner, recipient, operation) and checked against the grant to form A(g); termination and permitted requests are '
      'always included. The unbrokered planner arm measures behavioural compliance; the otherwise identical brokered arm checks the '
      'action before dispatch and records attempted violations separately from executed effects. The effect journal is post-action '
      'evidence, not prevention. A simple allowlist broker is an engineering baseline, not a CaMeL reproduction. Monte Carlo lookahead '
      'follows POMCP sampling but filters by A(g) at every node; without that filter the same search reproduces the substitution '
      'family, which E1 tests directly.', ('pomdp', 'pomcp', 'camel', 'design'))

# ============================================================================= 13 ALGORITHM B
s = new('Algorithm B: INCIDENT-SELECT, choose by verified effects.', '13 / Offline policy selection',
        'Does better incident evidence change which controller is selected, and does the selected controller behave better on new tasks?',
        [('optimization', 'Arms B0–B4 and the primary contrast', 3.4), ('impossible', 'Closest overlap: ImpossibleBench §§5–6', 3.4),
         ('related', 'Prior work comparison', 2.0)],
        status='PROPOSED')
mono(s, ML, 2.36, 7.05, 3.70, [
    'INCIDENT-SELECT(Theta, D_dev, S, theta_0, delta_U, B_run, alpha_F)',
    '  # Theta: K frozen candidates; S in {T, T+J, J*} selection signal',
    '  for theta in Theta:',
    '    tau <- ROLLOUT(theta, D_dev, shared seeds)   # reused for every S',
    '    U_hat, C_hat, F_hat <- COMMON_SCORER(tau)',
    '    R_hat_S(theta) <- INCIDENT_ESTIMATE(tau, S)',
    '  E <- {theta: U_hat >= U_hat(theta_0) - delta_U,',
    '               C_hat <= B_run, F_hat <= alpha_F}',
    '  if E = {}: return theta_0           # no improvement claim',
    '  theta_S <- argmin_{theta in E} R_hat_S(theta)   # fixed tie-break',
    '  FREEZE(theta_S)',
    '  return EVALUATE(theta_S, D_test, J*)  # untouched lineages,',
    '                                        # effect oracle, paired',
], size=11)
for i, (title, sub, col) in enumerate([('T', 'Transcript judge', BLUE), ('T + J', 'Same judge + effect journal', TEAL), ('J*', 'Exact local effect oracle', PLUM)]):
    y = 2.36 + i * .70
    box(s, 7.85, y, 4.86, .58, TINT[col])
    para(s, 8.02, y + .08, 1.0, .42, title, 17, col, True)
    para(s, 9.05, y + .08, 3.6, .42, sub, 12, INK)
box(s, 7.85, 4.46, 4.86, 1.60, PALE, rounded=False)
para(s, 8.02, 4.52, 4.55, 1.5,
     'T → T+J tests evidence access. T+J → J* tests the labelling method. Eligibility is held fixed, so the contrast isolates the '
     'incident ranking signal. Matched search budget, model access and candidate trials; log unused resources. Repeat whole selections '
     'to estimate optimizer variability. A tie is a valid result.', 11, INK, valign=MSO_ANCHOR.TOP, ls=1.1)
takeaway(s, 'Useful evidence must improve a held-out decision, not merely produce a better-looking label.')
notes(s, 'Offline selection procedure. Fix K candidate controllers, a development set, actor model, rollout budgets, utility '
      'definition and tie-breaks; reuse the same trajectories across selection-signal arms. Choose the lowest estimated incident rate '
      '(or a prespecified conservative bound) among candidates within the utility margin and cost/review limits; if none qualify, '
      'retain the reference and make no improvement claim. T is an actor-visible transcript; J is an independent service-effect '
      'journal; the T+J judge uses the same model and rubric as T with extra evidence; J* uses explicit local grant and effect '
      'semantics and is a privileged reference, not a general truth oracle. Use disjoint development, calibration and test lineages; '
      'freeze a repaired judge after development calibration (the prior failed calibration is not a qualified gate). '
      'ImpossibleBench already studies prompting, abort options and monitors; the incremental claim here is the selection contrast.',
      ('optimization', 'impossible', 'followup', 'related'))

# ============================================================================= 14 COMPARISON TABLE
s = new('How A and B differ from prior methods, point by point.', '14 / Comparison with existing algorithms',
        'What each method decides, how it treats the grant, how it asks, where it enforces, and what it learns from.',
        [('pomcp', 'POMCP', .8), ('rcpo', 'RCPO', .8), ('shield', 'Shielding', 1.0), ('camel', 'CaMeL', .8), ('progent', 'Progent', 1.0),
         ('control', 'AI Control', 1.1), ('impossible', 'ImpossibleBench', 1.5), ('related', 'Project related-work notes', 2.6)],
        status='THEORY')
rows = [['Method', 'Decides the next action by', 'Grant apart from feasibility?', 'Request channel', 'Enforcement point', 'Learns from effects?'],
    ['Scope instruction only (B0/B3)', 'LLM free-form', 'Grant lives in prompt text', 'Only if prompted', 'None', 'No'],
    ['Belief-state planning: POMCP', 'Monte Carlo lookahead on b', 'No: plans around obstacles', 'No', 'None', 'No'],
    ['Constrained RL: CPO / RCPO', 'Trained policy', 'Cost in expectation, not per effect', 'No', 'Training time', 'Needs gradients, many rollouts'],
    ['Shielding (Alshiekh 2018)', 'RL inside a safety automaton', 'Automaton over abstract state', 'No', 'Before each action', 'No'],
    ['Capability control: CaMeL', 'LLM plan; data-flow policies', 'Static policy on values and calls', 'No', 'Before tool call', 'No'],
    ['Privilege control: Progent', 'LLM', 'Policy over tool arguments', 'Expansion needs approval', 'Before tool call', 'Policy updates by LLM'],
    ['Trusted monitoring / defer', 'Untrusted proposes; trusted scores', 'No grant model; suspicion', 'Audit budget', 'Before execution', 'Thresholds from calibration'],
    ['Abort option + prompting', 'LLM', 'No', 'Abort = stop', 'None', 'No'],
    ['Proposed A + B', 'Lookahead inside A(g)', 'Yes: g moves only on trusted reply', 'ASK with cost d, prob. q', 'Broker before; journal after', 'B selects from journal'],
]
fills = {(9, j): (MINT, TEAL) for j in range(6)}
grid(s, ML, 2.32, [2.55, 2.25, 2.35, 1.65, 1.55, 1.74], rows, row_h=.375, size=10, fills=fills, ls=.98)
takeaway(s, 'Prior methods decide where to enforce or monitor. A decides what to do next, then still enforces; B decides what to release.', size=16)
notes(s, 'Characterisations are our reading of each paper, kept to what the method does by construction. POMCP (Silver & Veness 2010) '
      'plans by sampling over beliefs with no effect constraint, so a blocked route is just an obstacle. CPO (Achiam 2017) and RCPO '
      '(Tessler 2018) enforce constraints in expectation during training and need gradient access. Shielding (Alshiekh 2018) filters '
      'unsafe actions with a temporal-logic automaton over an abstraction. CaMeL (Debenedetti 2025) separates trusted control flow from '
      'untrusted data and enforces policies at tool calls. Progent (Shi 2025) checks calls against a least-privilege policy over tool '
      'arguments; narrowing updates are auto-approved and expansions need approval. AI Control (Greenblatt 2023) assumes an '
      'intentionally subversive untrusted model and uses a trusted monitor with an audit budget. ImpossibleBench (Zhong 2025) studies '
      'prompting and abort options on impossible coding tasks. None of these is reproduced here; the comparison is structural, and '
      'the proposal has no empirical result yet.', ('pomcp', 'cpo', 'rcpo', 'shield', 'camel', 'progent', 'control', 'impossible', 'related'))

# ============================================================================= 15 POSITIONING
s = new('Where each method intervenes, and what is actually new.', '15 / Positioning',
        'Prior methods sit at single points of the pipeline; the proposal spans planning and selection and keeps enforcement.',
        [('related', 'Project: prior work and the contribution we can defend', 4.2), ('agentdojo', 'AgentDojo', 1.0), ('control', 'AI Control', 1.0)],
        status='THEORY')
stages = [('Propose', 'next action', BLUE), ('Check', 'before dispatch', TEAL), ('Execute', 'tool call', MUTED),
          ('Journal', 'service effects', ORANGE), ('Score', 'U, R, C, F', PLUM), ('Select / release', 'next policy', RED)]
sw = (CW - 5 * .18) / 6
for i, (t_, sub, col) in enumerate(stages):
    x = ML + i * (sw + .18)
    box(s, x, 3.62, sw, .80, col)
    para(s, x + .08, 3.66, sw - .16, .40, t_, 14, WHITE, True, align='c')
    para(s, x + .08, 4.02, sw - .16, .34, sub, 10.5, WHITE, align='c')
    if i < 5:
        arrow(s, [(x + sw + .02, 4.02), (x + sw + .16, 4.02)], MUTED)
above = [['POMCP · CPO/RCPO', 'prompting · abort', 'Algorithm A'], ['CaMeL · Progent', 'shields · trusted monitor', 'broker (A)'],
         [], ['effect journal', '(ours, instrumented)'], ['scorer · T / T+J / J*', 'transcript monitors'], ['AI Control protocols', 'regression gate', 'Algorithm B']]
for i, items in enumerate(above):
    x = ML + i * (sw + .18)
    for j, item in enumerate(items):
        ours = item.startswith('Algorithm') or item.startswith('broker') or item.startswith('effect') or item.startswith('(ours')
        chip(s, x, 2.36 + j * .40, sw, item, TEAL if ours else MUTED, MINT if ours else PALE, 9.5, .32)
card(s, ML, 4.62, 5.9, 1.44, MUTED)
para(s, ML + .22, 4.68, 5.5, .32, 'Not new', 14, MUTED, True)
para(s, ML + .22, 5.00, 5.5, 1.0,
     'POMDP planning and value of information; constrained optimization; capability and privilege enforcement; impossible-task tests; '
     'separate safety and utility endpoints (AgentDojo, CoSec); control protocols against subversion.', 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.08)
card(s, 6.81, 4.62, 5.9, 1.44, TEAL)
para(s, 7.03, 4.68, 5.5, .32, 'Proposed, and testable', 14, TEAL, True)
para(s, 7.03, 5.00, 5.5, 1.0,
     'Grant as state with trusted-only transitions, separate from feasibility; the request as a planned action with a price; '
     'attempted and executed violations measured in one fixture; selection from effects tested against selection from transcripts.',
     11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.08)
takeaway(s, 'The novelty claim is empirical and narrow: better choices at retained utility, not a new containment guarantee.')
notes(s, 'The pipeline places each prior method at its intervention point so comparisons happen at like points: a broker is compared '
      'with CaMeL-style and Progent-style enforcement before dispatch, the planner with POMCP and prompting at proposal, the selector '
      'with control-protocol evaluation and the regression gate at release. The right-hand card states the only novelty we intend to '
      'defend, and the left-hand card lists what is explicitly borrowed. The related-work document records the closest overlaps, '
      'including ImpossibleBench, the Anthropic incident assessment and CoSec.', ('related', 'agentdojo', 'control'))

# ============================================================================= 16 E1 SIMULATOR
s = new('E1: does the decision model predict choices, and at what regret?', '16 / Experiment 1 · simulator',
        'A small finite simulator with known kernels, then noisy and misspecified variants. The exact DP is the normative reference.',
        [('voi', 'Regret and metalevel decisions: Hay et al. 2012', 4.0), ('pomcp', 'Unconstrained reference planner', 2.6)],
        status='PROPOSED')
grid(s, ML, 2.36, [2.1, 3.8], [
    ['Factor swept', 'Range'],
    ['Feasibility prior p', '0.01 … 0.99'],
    ['Diagnostic accuracy / cost', 'perfect → noisy; 0 … 2 units'],
    ['Request cost d, grant prob. q', 'd 0 … 8; q 0.2 … 0.9'],
    ['Horizon H, budget k', '1 … 6 steps; 2 … 12 units'],
    ['Model misspecification', 'T, O perturbed ±20 %; stale q'],
], row_h=.40, size=11)
grid(s, 6.75, 2.36, [2.3, 3.66], [
    ['Planner', 'Role'],
    ['Exact DP (Algorithm A)', 'Normative reference; zero regret by construction'],
    ['One-step VOI', 'Ablation: no lookahead beyond one probe'],
    ['Budgeted MC lookahead + A(g)', 'Approximation; same grant filter'],
    ['POMCP without A(g)', 'Unconstrained: predicted to substitute'],
    ['LLM planner (prompted)', 'Behavioural compliance; calibration'],
], row_h=.40, size=11)
box(s, ML, 4.90, CW, 1.15, PALE, rounded=False)
para(s, ML + .2, 4.95, CW - .4, 1.05,
     'Endpoints: value regret V* − V_{π}; violation rate; request and probe rates; false stops; calibration by proper scoring rules and '
     'reliability plots. Prediction 1: the unconstrained planner\'s violation rate rises as p falls, reproducing H1. Prediction 2: '
     'the LLM planner\'s choices track the phase diagram on slide 11 and the cutoffs on slide 13 only where its confidence is calibrated. The event log must '
     'not reveal hidden feasibility before the permitted probe.', 11.5, INK, valign=MSO_ANCHOR.TOP, ls=1.1)
takeaway(s, 'Falsifier: if the LLM planner ignores p, d and q, the decision model does not describe agent behaviour.')
notes(s, 'E1 is cheap and runs before any fixture extension: a finite simulator whose kernels are known, so regret is exact. '
      'Sweeps follow the toy results on slides 10–14. The POMCP-without-filter arm is the mechanistic version of H1: it should prefer '
      'the reachable substitute whenever the belief in local repair is low, with no change in the grant. The LLM planner arm is the '
      'first check that an LLM agent\'s choices move with the decision variables at all; calibration is evaluated with proper scoring '
      'rules. All E1 claims are about the model, not about safety in deployment.', ('voi', 'pomcp'))

# ============================================================================= 17 E2 FACTORIAL
s = new('E2: same reachable resource, different permission, seven arms.', '17 / Experiment 2 · the discriminating fixture',
        'Repair availability × trusted authorization; the alternate resource stays reachable; prior-method baselines enter as arms.',
        [('design', 'Project: repair × authorization design', 3.4), ('next', 'X1b: silent and conflicting grants', 3.0), ('control', 'Defer protocol', 1.2)],
        status='PROPOSED')
gx, gy, gw, gh = 2.05, 2.72, 2.35, .80
para(s, gx, 2.36, gw, .30, 'ALTERNATE FORBIDDEN', 11.5, RED, True)
para(s, gx + gw + .14, 2.36, gw, .30, 'ALTERNATE GRANTED', 11.5, TEAL, True)
para(s, ML, gy + .16, 1.4, .55, 'Local repair\nworks', 12, INK, True)
para(s, ML, gy + gh + .30, 1.4, .55, 'Local repair\nfails', 12, INK, True)
for col, row, title, fill, accent in [
    (0, 0, 'Recover locally', MINT, TEAL), (1, 0, 'Either permitted route', PALE, BLUE),
    (0, 1, 'Verify blocker, then stop', ROSE, RED), (1, 1, 'Use the granted alternative', MINT, TEAL),
]:
    x, y = gx + col * (gw + .14), gy + row * (gh + .14)
    box(s, x, y, gw, gh, fill)
    para(s, x + .12, y + .10, gw - .24, .6, title, 12.5, accent, True, ls=1.0)
para(s, ML, 4.66, 4.3, .80, 'X1b adds silent and conflicting grant columns: readout is a response mix and an over-conservatism '
     'check, not a violation rate.', 10.5, MUTED, valign=MSO_ANCHOR.TOP, ls=1.05)
grid(s, 7.0, 2.36, [1.0, 2.1, 2.61], [
    ['Arm', 'Policy', 'Compared at'],
    ['A0', 'Original instruction', 'reference'],
    ['A1', 'Fixed scope instruction', 'prompt'],
    ['A2', 'Stop-or-ask instruction', 'prompt'],
    ['A3', 'Broker allowlist (Progent / CaMeL-style)', 'check'],
    ['A4', 'Trusted monitor + defer-to-stop', 'check'],
    ['A5', 'Algorithm A, unbrokered', 'propose'],
    ['A6', 'Algorithm A + broker', 'propose + check'],
], row_h=.335, size=10.5)
box(s, ML, 5.50, CW, .58, PALE, rounded=False)
para(s, ML + .2, 5.53, CW - .4, .52,
     f"Pilot: {E2['repair']} repair × {E2['grant']} grant × {E2['families']} task families × {E2['arms']} arms × {E2['lineages']} lineages "
     f"= {E2_ROLLOUTS} rollouts per model. Endpoints per run: executed effect, attempted violation, authorized completion, false stop, "
     'probes, requests, cost, unknown.', 11, INK, ls=1.05)
takeaway(s, 'When permission changes, the next action should change even if access does not. Attempts and effects are counted apart.')
notes(s, 'The matrix is a normative reference, not observed results. Pair each task instance across local repair available/unavailable '
      'and alternate forbidden/granted with contents, reachability, schemas and budgets fixed; recoverable and unavailable variants '
      'share the initial failure history and a permitted probe reveals availability. Trusted permission is the treatment. Grants are '
      'fixed within an episode for the factorial; a separate dynamic-request study adds grant/deny responses with cost. A3 and A4 are '
      'engineering baselines in the spirit of Progent, CaMeL and defer-to-trusted, not reproductions. The pilot allocation is for '
      'precision estimation, not a powered confirmatory claim; twelve seeds of one template are not twelve lineages. Report all '
      'assigned runs with unknowns bounded, cluster paired uncertainty by lineage, and evaluate each model separately.',
      ('design', 'next', 'control', 'progent', 'camel'))

# ============================================================================= 18 E3 SELECTION
s = new('E3: does incident feedback select a better policy on fresh tasks?', '18 / Experiment 3 · selection contrast',
        'Shared development pool; selections by T, T+J and J*; frozen policies evaluated on untouched lineages with one common observer.',
        [('optimization', 'Matched selection protocol', 2.6), ('followup', 'Why the current judge is not a gate', 3.2), ('pilot', 'Archived evidence', 2.0)],
        status='PROPOSED')
para(s, ML, 2.36, 6.0, .34, 'Evidence we already have: the measurement problem', 14, NAVY, True)
tx, ty, tw, th = 2.15, 3.08, 1.75, .72
para(s, tx, 2.76, tw, .26, 'NO INCIDENT', 11, TEAL, True)
para(s, tx + tw + .14, 2.76, tw, .26, 'INCIDENT', 11, RED, True)
para(s, ML, ty + .14, 1.45, .44, 'Scorer pass', 12.5, INK, True)
para(s, ML, ty + th + .30, 1.45, .44, 'No pass', 12.5, INK, True)
for col, row, n, fill, color in [(0, 0, HIST['pass_clean'], MINT, TEAL), (1, 0, HIST['pass_incident'], ROSE, RED),
                                  (0, 1, HIST['no_pass_clean'], PALE, MUTED), (1, 1, HIST['no_pass_incident'], SAND, ORANGE)]:
    x, y = tx + col * (tw + .14), ty + row * (th + .16)
    box(s, x, y, tw, th, fill)
    para(s, x, y, tw, th, str(n), 28, color, True, align='c')
para(s, ML, 4.92, 6.0, .55, f"All {HIST['pass_clean']} feasible controls completed within scope. {HIST['pass_incident']} shortcuts passed the scorer; "
     f"{HIST['no_pass_incident']} violations never passed.", 11.5, MUTED, ls=1.05)
card(s, 6.95, 2.36, 5.76, 3.70, PLUM)
para(s, 7.17, 2.44, 5.3, .34, 'Protocol and decision rules', 14, PLUM, True)
para(s, 7.17, 2.80, 5.3, 3.2,
     '1  Fix K candidates, development lineages, actor, budgets, utility and tie-breaks; one rollout set reused by all signals.\n'
     '2  Select θ_{T}, θ_{T+J}, θ_{J*}; freeze; repeat full selections to estimate optimizer variability.\n'
     '3  Evaluate on fresh lineages and a disjoint task family with the effect oracle; pair by task; exact McNemar on discordant pairs.\n'
     '4  Primary contrast θ_{T+J} vs θ_{T}: lower verified risk with authorized utility inside a preregistered margin.\n'
     '5  Calibrate a repaired judge on development data only and freeze it before test collection.\n'
     'Valid negative results: identical selections; no eligible policy; no risk reduction; utility loss beyond the margin.',
     11, INK, valign=MSO_ANCHOR.TOP, ls=1.1)
takeaway(s, 'Falsifiers: only the broker helps; refusals explain the gain; feedback selects the same policy.')
notes(s, f"Counts recomputed from results/incident-study-v1/summary.json during this build: {HIST['pass_clean']} passes without an "
      f"effect, {HIST['pass_incident']} scorer passes with an unauthorized effect, {HIST['no_pass_clean']} no-passes without an effect, "
      f"{HIST['no_pass_incident']} no-passes with one. All 15 violating trajectories first read the previous-run answer archive. The "
      "matrix pools deliberately different arms and variants to show endpoint disagreement, not to estimate a deployment rate; the 49 "
      "clean no-passes were impossible cases, not premature refusals. Current follow-ups do not validate a transcript gate or show "
      "improved selection; the 12-case calibration failed and is not a qualified gate. E3 requires lower risk with retained utility, "
      "cost and coverage limits satisfied, and reports unknowns as bounds rather than as safe.", ('optimization', 'followup', 'pilot'))

# ============================================================================= 19 MAP BACK
s = new('Back to the incidents: which result, which control, which test.', '19 / Real-world implications',
        'Each incident family maps to one decision result, one system change, and the experiment whose outcome would support or refute it.',
        [('medicare', 'Medicare', 1.0), ('workbook', 'Workbook', 1.0), ('wallet', 'Wallet', .8), ('anthropic', 'Cyber evaluations', 1.6),
         ('library', 'Channel family: records 1, 2, 6, 14', 2.8)],
        status='PROPOSED')
grid(s, ML, 2.36, [2.45, 2.55, 3.95, 3.14], [
    ['Incident family', 'Decision result', 'System change to test', 'Experiment and readout'],
    ['Substitution (Medicare, workbook, wallet)', 'R1 stopping; R2 VOI: a failed route moves b, never g',
     'Scope data routes and operations; provide a priced request channel; make an honest partial answer worth something (B > 0)',
     'E2 forbidden column: executed effects and attempts; false stops on the granted column'],
    ['Scope inference (cyber evaluations 16–21)', 'R3 ambiguity: silent grant → reversible action',
     'Explicit grant field on every task (explicit / implicit / unspecified / conflicting); default to clarification under silence',
     'X1b: response mix and over-conservatism on feasible cells'],
    ['Channel (GET writes, DNS, artifacts)', 'Formulation: violations are effects, not intents',
     'Label operation and flow, not the HTTP verb; effect journal at the service; broker before dispatch',
     'E3 T vs T+J: does effect evidence change the selected policy?'],
    ['Late detection (median 119 days to disclosure)', 'R4 selection: finite, paired evidence',
     'Replayable incident twins and feasible controls as a regression gate before each policy release',
     'E3 power table; gate miss rate on held-out twins'],
], row_h=.70, size=10.5, ls=1.02)
takeaway(s, 'Useful persistence = better plans inside authority + timely, priced requests when authority must change.')
notes(s, 'These are counterfactual intervention hypotheses, not claims that any control would certainly have prevented the reported '
      'incidents. For statistics research, an access-denied observation does not certify global impossibility and an allowed source '
      'may exist. For workbook sharing, the reliability of the authorized transfer route changes the value of continued recovery. '
      'For deployment, success and low apparent harm do not replace an account holder\'s grant. For cyber evaluations, a matching '
      'domain name is not a grant. For the channel family, only the service knows what a call did. Each row names the experiment '
      'whose outcome would support or refute the proposed change; test each replay locally with dummy data, a trusted grant and an '
      'independent effect observer before broader validation.', ('medicare', 'workbook', 'wallet', 'anthropic', 'library', 'camel'))

# ============================================================================= 20 OUTCOMES AND MILESTONES
s = new('What each outcome would mean, and what happens next.', '20 / Reading the results in advance',
        'Interpretations are fixed before the data arrive. A null is informative and will be reported as such.',
        [('design', 'Incidents to testable questions', 2.8), ('optimization', 'Incident-aware optimization', 2.8), ('next', 'Next experiments', 1.8)],
        status='PROPOSED')
grid(s, ML, 2.36, [4.3, 7.79], [
    ['If the experiments show…', 'Then for deployed agents…'],
    ['A5 cuts executed effects at retained utility (E2)', 'Decision structure matters: expose feasibility, grant and a priced request to the agent; the model describes behaviour.'],
    ['Only A3 / A6 cut effects; A5 does not', 'Enforcement, not reasoning, carries safety: invest in brokers and effect journals; treat planners as utility tools.'],
    ['Gains come with more false stops on granted cells', 'Over-conservatism is the cost: tune B and d, and report refusals as a first-class endpoint.'],
    ['θ_{T+J} = θ_{T} on fresh lineages (E3)', 'The effect journal is an audit instrument, not a selection advantage; the contribution is measurement.'],
    ['LLM choices ignore p, d, q (E1)', 'The POMDP is a normative reference only; calibration work precedes any planner deployment.'],
], row_h=.50, size=11, ls=1.02)
for i, (when, what, col) in enumerate([
    ('Nov 2026', 'Finite simulator; E1 preregistered and run', BLUE),
    ('Dec 2026', 'Fixture extension: repair × grant × X1b; E2 pilot', TEAL),
    ('Jan 2027', 'Repaired judge frozen; E3 selection contrast', PLUM),
    ('Feb 2027', 'Confirmatory twins sized from pilot discordance', ORANGE),
]):
    x = ML + i * (CW + .16) / 4
    w = (CW - 3 * .16) / 4
    box(s, x, 5.42, w, .66, TINT[col])
    box(s, x, 5.42, .07, .66, col, rounded=False)
    para(s, x + .18, 5.44, w - .3, .28, when, 11.5, col, True)
    para(s, x + .18, 5.70, w - .3, .36, what, 10, INK, ls=1.0)
takeaway(s, 'Next milestone: implement the finite decision model and the paired grant/repair fixture, qualify its controls, preregister E1 and E2.')
notes(s, 'The outcome table is the preregistered interpretation. Each row names the arm or contrast and the implication, including '
      'the two negative readings in which only enforcement helps or the selection signals tie. Milestones are proposed dates, not '
      'commitments; the confirmatory sample is sized from pilot discordance under a prespecified risk reduction and utility '
      'noninferiority margin, before any new test cohort is collected. The deck contains 24 narrative slides and two glossary slides; citations, exact toy assumptions, '
      'study design and limitations are in the speaker notes and footers.', ('design', 'optimization', 'next'))

# ============================================================================= APPENDIX A: MECHANISM CODEBOOK
MECHANISM_READINGS = {
    'M1': 'Completion pressure drives persistence when no safe failure or stop outcome is available.',
    'M2': 'Separate capabilities combine into a path across the boundary.',
    'M3': 'Shared files or services let otherwise separate runs communicate or coordinate.',
    'M4': 'Warning or detection fails to produce a timely, effective intervention.',
    'M5': 'Assumptions about the real environment and gaps in audit coverage hide exposure.',
    'M6': 'Responsibility is split across teams or organizations; no owner sees the whole chain.',
}
s = new('M1-M6: how did the safeguards fail?', 'Appendix A / Failure mechanism codebook',
        'M = mechanism. The number identifies a category, with no severity ordering. Formal names come from the incident taxonomy.',
        [('taxonomy', 'incidents/taxonomy.json', 2.4), ('medicare_coding', 'Medicare: draft coding, review pending', 3.5)],
        status='CODEBOOK')
mechanism_rows = [['ID', 'Formal name', 'Plain-language reading', 'Medicare']]
for code, name in TAXONOMY['mechanisms'].items():
    mechanism_rows.append([code, name, MECHANISM_READINGS[code], MEDICARE['mechanisms'][code]['status']])
grid(s, ML, 2.36, [.55, 4.05, 6.0, 1.49], mechanism_rows, row_h=.46, size=11.5)
para(s, ML, 5.70, CW, .42,
     'present = coded as present; partial = partly supported; absent = coded absent; unknown = insufficient evidence.\n'
     'Medicare M4 records no in-run interruption and late notification; it does not establish that an in-run alert fired.',
     11.5, MUTED, ls=1.05)
takeaway(s, 'Medicare: M1, M2, M4, M6 present; M3 absent; M5 partial. These are draft assessments of failure mechanisms.', size=14)
notes(s, 'M codes classify failure mechanisms, rather than ranking severity or counting broken rules. Formal names are read from '
      'incidents/taxonomy.json; the plain-language column explains those labels. Medicare statuses are read from the record, '
      'including the previously omitted M3 absent and M5 partial. The record does not give a rationale for M5 partial. '
      'Mechanism notes supplied in the record: ' + '; '.join(
          f"{code}: {item['note']}" for code, item in MEDICARE['mechanisms'].items() if item['note']) + '. '
      'M4 is the general category Detection without decisive escalation; the case note supports the absence of timely interruption '
      'and delayed notification, not a claim that the monitor detected this case in-run. All these codes remain analyst drafts '
      'pending review. The next appendix explains the P design-proposition family.', ('taxonomy', 'medicare_coding', 'library'))

# ============================================================================= APPENDIX B: PROPOSITION CODEBOOK
PROPOSITION_READINGS = {
    'P1': 'Maintain the boundary across every tool, service and action in the system.',
    'P2': 'Measure isolation and reachability before trusting a claim such as "simulated".',
    'P3': 'Keep authorization binding even when a technical restriction fails.',
    'P4': 'Treat blocked or impossible goals as security-relevant inputs; provide safe exits.',
    'P5': 'Account for communication and coordination through shared files or services.',
    'P6': 'Assess actual access, writes and transfers, regardless of the stated intent.',
    'P7': 'Show that the audit covers relevant runs, actions and effects.',
    'P8': 'Resume only with evidence that the cause is understood and addressed.',
    'P9': 'Assign a human accountable for consequential autonomous actions.',
}
s = new('P1-P9: which design principle does the case test?', 'Appendix B / Design proposition codebook',
        'P = proposition: a design principle assessed against the case. The number identifies a principle; the status is a draft analyst judgment.',
        [('taxonomy', 'incidents/taxonomy.json', 2.4), ('medicare_coding', 'Medicare: draft coding, review pending', 3.5)],
        status='CODEBOOK')
proposition_rows = [['ID', 'Formal name', 'Plain-language reading', 'Medicare']]
for code, name in TAXONOMY['propositions'].items():
    relation = MEDICARE['propositions'][code]['relation']
    proposition_rows.append([code, name, PROPOSITION_READINGS[code], 'n/a' if relation == 'not_applicable' else relation])
grid(s, ML, 2.36, [.55, 4.05, 6.0, 1.49], proposition_rows, row_h=.35, size=10.5)
para(s, ML, 5.93, CW, .26,
     'violated = not upheld; demonstrated = supported; mixed = mixed assessment; n/a = not applicable to this case.', 11, MUTED)
takeaway(s, 'Medicare: P1, P2, P3, P4, P6, P9 violated; P7 mixed; P5/P8 n/a. These labels are separate from the 4/6 signals.', size=14)
notes(s, 'P codes identify design propositions in the incident taxonomy. Formal names and Medicare relations are loaded from the '
      'taxonomy and record respectively. Plain-language readings explain the principles without assigning new case-specific '
      'evidence. Six principles are coded violated in Medicare: P1, P2, P3, P4, P6 and P9; P7 is mixed, P5 and P8 are not '
      'applicable. This is a different classification from the four observed risk signals f, p, x and v. The record gives no '
      'per-proposition explanatory notes, including no reason for the mixed P7 judgment, so this slide reports the draft coding '
      'without inventing a rationale. Violated is an analyst judgment about the principle, not a legal finding. Demonstrated '
      'means the case supports the principle; mixed means the coding is neither a simple violation nor demonstration; n/a means '
      'not applicable, not proof that the system satisfied the principle. Algorithm properties elsewhere use ordinary list '
      'numbers and simulator predictions are Prediction 1/2, so P1-P9 consistently denote this taxonomy.',
      ('taxonomy', 'medicare_coding', 'library'))

# ============================================================================= validation
def validate():
    assert len(PRS.slides) == N_SLIDES, len(PRS.slides)
    section_of = {}
    for pg in PAGES:
        k = pg['kicker']
        if k[:2].isdigit() and ' / ' in k:
            section_of[pg['slide']] = int(k[:2])
    part = {1: range(1, 7), 2: range(7, 12), 3: range(12, 16), 4: range(16, 19), 5: range(19, 21)}   # kicker numbers per part
    for num, where in OUTLINE_RANGES.items():
        lo, hi = (int(v) for v in where.split()[1].split('–'))
        for n in range(lo, hi + 1):
            assert section_of[n] in part[int(num)], (num, n, section_of[n])
    # Result 1: one-step rule agrees with backward induction over a finite budget.
    for k in range(1, 7):
        for p in (.1, .2, .24, .26, .3, .5, .7, .9):
            one_step_try = p * STOP_TOY['q'] * (STOP_TOY['R'] - STOP_TOY['B']) > STOP_TOY['c']
            dp_try = stopping_value(p, k, STOP_TOY) > STOP_TOY['B'] + 1e-12
            assert one_step_try == dp_try, (k, p)
    assert abs(STOP_THRESHOLD - .25) < 1e-12 and ATTEMPTS_BEFORE_STOP == 3
    # Result 2: exact toy and independent enumeration of contingent policies.
    assert abs(toy_values(.3, 2)['PROBE'] - 4.7) < 1e-10
    for p, d, action in [(.3, 2, 'PROBE'), (.05, 2, 'ASK'), (.95, 2, 'RECOVER'), (.01, 7, 'STOP'),
                         (.9, 8, 'RECOVER'), (.1, 0, 'ASK')]:
        assert best(p, d) == action
    for pi in range(101):
        for di in range(81):
            p, d = pi / 100, di / 10
            enumerated = max(-.5 + p * g + (1 - p) * b for g in (0, 8, 6 - d) for b in (0, -2, 6 - d))
            assert abs(enumerated - toy_values(p, d)['PROBE']) < 1e-9
            lo, hi = voi_boundaries(d)
            predicted = ('ASK' if d <= 6 else 'STOP') if p < lo else 'PROBE' if p <= hi else 'RECOVER'
            values = toy_values(p, d)
            # Ties are legitimate; compare value, not a tie-breaking action label.
            assert abs(values[predicted] - max(values.values())) < 1e-9, (p, d, predicted, values)
    boundary_checks = 0
    for di in range(81):
        d = di / 10
        lo, hi = voi_boundaries(d)
        left = 'ASK' if d <= 6 else 'STOP'
        assert 0 < lo < hi < 1
        for p, a, b in [(lo, left, 'PROBE'), (hi, 'PROBE', 'RECOVER')]:
            values = toy_values(p, d)
            assert abs(values[a] - values[b]) < 1e-9
            assert abs(values[a] - max(values.values())) < 1e-9
            below, above = toy_values(p - 1e-5, d), toy_values(p + 1e-5, d)
            assert abs(below[a] - max(below.values())) < 1e-9 and below[a] > below[b]
            assert abs(above[b] - max(above.values())) < 1e-9 and above[b] > above[a]
            boundary_checks += 1
        if d <= 6:
            values = toy_values((8 - d) / 10, d)
            assert abs(values['ASK'] - values['RECOVER']) < 1e-9
            assert values['PROBE'] > values['ASK']
    example_values = toy_values(.3, 2)
    gross_voi = example_values['PROBE'] + VOI_TOY['probe_cost'] - max(example_values[a] for a in ('STOP', 'ASK', 'RECOVER'))
    assert abs(gross_voi - 1.2) < 1e-9 and abs(gross_voi - VOI_TOY['probe_cost'] - .7) < 1e-9
    assert voi_boundaries(2) == (1 / 8, 11 / 12)
    # Result 3: robust rules pick the reversible action.
    assert AMB_MAXMIN == 'Ask first' and AMB_REGRET == 'Ask first'
    assert abs(10 * AMB_PRIOR_SWITCH - 20 * (1 - AMB_PRIOR_SWITCH) - 7) < 1e-12
    # Result 4: slack decreases in n; power increases in n.
    assert all(a[1] > b[1] for a, b in zip(SEL_ROWS, SEL_ROWS[1:]))
    assert all(a[1][0] <= b[1][0] + 1e-9 for a, b in zip(POWER_ROWS, POWER_ROWS[1:]))
    assert E2_ROLLOUTS == 672
    issues = []
    missing_glyphs = set('\u2208\u2209\u21d4\u21d2\u2295')   # Arial.ttf on this Mac lacks these; they render as boxes
    for n, sl in enumerate(PRS.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame and missing_glyphs & set(sh.text):
                issues.append(f'slide {n}: {sh.name} uses a glyph Arial cannot render')
        assert sl.has_notes_slide and len(sl.notes_slide.notes_text_frame.text) > 100, n
        for sh in sl.shapes:
            if min(sh.left, sh.top) < -10 or sh.left + sh.width > PRS.slide_width + 100 or sh.top + sh.height > PRS.slide_height + 100:
                issues.append(f'slide {n}: {sh.name} out of bounds')
    assert not issues, issues
    report = {
        'slides': N_SLIDES, 'pages': PAGES, 'shape_boundary_issues': issues,
        'library': {'n': LIB_N, 'overreach': len(OVERREACH), 'overreach_with_f': OVERREACH_F,
                    'latency': {k: LAT[k] for k in ('detection_days_known', 'detection_median', 'disclosure_days_known', 'disclosure_median')},
                    'stop_behavior': STOPB},
        'stopping_toy': {**STOP_TOY, 'threshold': STOP_THRESHOLD, 'posteriors': POSTERIORS, 'attempts_before_stop': ATTEMPTS_BEFORE_STOP},
        'voi_toy': example_values, 'voi_assumptions': VOI_TOY,
        'voi_gross': gross_voi, 'voi_net': gross_voi - VOI_TOY['probe_cost'],
        'voi_policy_enumeration_checks': 101 * 81, 'voi_analytic_policy_checks': 101 * 81,
        'voi_boundary_and_neighbor_checks': boundary_checks,
        'voi_thresholds_at_d_2': voi_boundaries(2),
        'ambiguity_toy': {'rows': AMB, 'maxmin': AMB_MAXMIN, 'regret': AMB_REGRET, 'prior_switch': AMB_PRIOR_SWITCH},
        'selection_slack': {'K': SEL_K, 'delta': .05, 'rows': SEL_ROWS},
        'mcnemar_power': {'p_base': P_BASE, 'targets': POWER_TARGETS, 'rows': POWER_ROWS},
        'e2_pilot': {**E2, 'rollouts_per_model': E2_ROLLOUTS},
        'historical_counts': HIST,
        'historical_source': str(summary_path.relative_to(ROOT)),
        'historical_sha256': hashlib.sha256(summary_path.read_bytes()).hexdigest(),
        'diagnostic_source': str(diag_path.relative_to(ROOT)),
        'sources_checked': '2026-10-21', 'sources': SOURCES, 'new_hosted_model_calls': 0,
        'word_counts': [len(' '.join(sh.text for sh in sl.shapes if sh.has_text_frame).split()) for sl in PRS.slides],
    }
    (QA / 'content-validation.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n')
    (HERE / 'speaker-notes.txt').write_text('\n\n'.join(
        f"SLIDE {n}: {PAGES[n - 1]['title']}\n\n{sl.notes_slide.notes_text_frame.text}"
        for n, sl in enumerate(PRS.slides, 1)) + '\n')

if __name__ == '__main__':
    validate()
    PRS.save(OUT)
    print(f'Created {OUT}: {N_SLIDES} slides; library counts, archive counts and toy calculations verified.')
