"""Render this deck's native text, tables and simple shapes as vector PDF.

This intentionally supports only the elements used in build_deck.py and fails
on unsupported objects or text overflow. It is not a general PPTX converter.
Requires python-pptx and reportlab, plus the macOS Arial and Menlo fonts.
"""
from __future__ import annotations

import json
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, Preformatted

HERE = Path(__file__).resolve().parent
EMU_PER_PT = 12700


def pt(value):
    return value / EMU_PER_PT


def color(value, fallback="#111111"):
    try:
        return HexColor("#" + str(value.rgb))
    except (AttributeError, TypeError):
        return HexColor(fallback)


def register_fonts():
    base = Path("/System/Library/Fonts/Supplemental")
    for name, filename in [
        ("Arial", "Arial.ttf"), ("Arial-Bold", "Arial Bold.ttf"),
        ("Arial-Italic", "Arial Italic.ttf"), ("Arial-BoldItalic", "Arial Bold Italic.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(base / filename)))
    pdfmetrics.registerFont(TTFont("Menlo", "/System/Library/Fonts/Menlo.ttc"))


def font_name(font):
    if font.name == "Menlo":
        return "Menlo"
    if font.bold and font.italic:
        return "Arial-BoldItalic"
    return "Arial-Bold" if font.bold else "Arial-Italic" if font.italic else "Arial"


def render_text(canvas, frame, x, y, width, height, label, issues):
    """Coordinates are points measured from the top left of the slide."""
    left, right = pt(frame.margin_left), pt(frame.margin_right)
    top, bottom = pt(frame.margin_top), pt(frame.margin_bottom)
    available = width - left - right
    blocks = []
    for paragraph in frame.paragraphs:
        if not paragraph.runs:
            continue
        sizes = [pt(run.font.size) if run.font.size else 14 for run in paragraph.runs]
        size = max(sizes)
        spacing = paragraph.line_spacing or 1.1
        leading = size * spacing if isinstance(spacing, float) else pt(spacing)
        if not leading:
            leading = size * 1.1
        align = {PP_ALIGN.LEFT: 0, PP_ALIGN.CENTER: 1, PP_ALIGN.RIGHT: 2}.get(paragraph.alignment, 0)
        style = ParagraphStyle(label, fontName=font_name(paragraph.runs[0].font),
                               fontSize=size, leading=leading, alignment=align)
        if all(run.font.name == "Menlo" for run in paragraph.runs):
            line = paragraph.text
            measured = pdfmetrics.stringWidth(line, "Menlo", size)
            if measured > available + .5:
                issues.append(f"{label}: code width {measured:.1f} > {available:.1f}: {line[:65]}")
            block = Preformatted(line or " ", style)
        else:
            spans = []
            for run, run_size in zip(paragraph.runs, sizes, strict=True):
                c = color(run.font.color).hexval().replace("0x", "#")
                spans.append(f'<font name="{font_name(run.font)}" size="{run_size}" color="{c}">'
                             + escape(run.text).replace("\v", "<br/>") + "</font>")
            block = Paragraph("".join(spans), style)
        _, block_height = block.wrap(available, 10000)
        before = pt(paragraph.space_before) if paragraph.space_before else 0
        after = pt(paragraph.space_after) if paragraph.space_after else 0
        blocks.append((block, block_height, before, after))
    used = sum(h + before + after for _, h, before, after in blocks)
    room = height - top - bottom
    if used > room + .5:
        issues.append(f"{label}: text height {used:.1f} > {room:.1f}: {frame.text[:70]}")
    offset = max(0, room-used)
    anchor = frame.vertical_anchor
    cursor = y + top + (offset/2 if anchor == MSO_ANCHOR.MIDDLE else offset if anchor == MSO_ANCHOR.BOTTOM else 0)
    for block, h, before, after in blocks:
        cursor += before
        block.drawOn(canvas, x+left, canvas._pagesize[1]-cursor-h)
        cursor += h + after


def fill_shape(canvas, shape, x, y, w, h):
    solid = shape.fill.type == MSO_FILL_TYPE.SOLID
    stroke = shape.line.fill.type == MSO_FILL_TYPE.SOLID
    if solid:
        canvas.setFillColor(color(shape.fill.fore_color, "#ffffff"))
    if stroke:
        canvas.setStrokeColor(color(shape.line.color))
        canvas.setLineWidth(pt(shape.line.width) if shape.line.width else .75)
    bottom = canvas._pagesize[1]-y-h
    if shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE:
        if shape.rotation != 90:
            raise ValueError("Only the deck's right-facing triangle is supported")
        path = canvas.beginPath()
        path.moveTo(x, bottom)
        path.lineTo(x+w, bottom+h/2)
        path.lineTo(x, bottom+h)
        path.close()
        canvas.drawPath(path, stroke=int(stroke), fill=int(solid))
    elif shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
        canvas.roundRect(x, bottom, w, h, min(w,h)*.05, stroke=int(stroke), fill=int(solid))
    elif shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.RECTANGLE:
        canvas.rect(x, bottom, w, h, stroke=int(stroke), fill=int(solid))
    else:
        raise ValueError(f"Unsupported shape: {shape.auto_shape_type}")


def main():
    register_fonts()
    prs = Presentation(HERE / "incident-library.pptx")
    target = HERE / "incident-library.pdf"
    canvas = Canvas(str(target), pagesize=(pt(prs.slide_width), pt(prs.slide_height)))
    canvas.setTitle("One incident to executable policy - INC-2026-015")
    canvas.setAuthor("containment-extension")
    issues = []
    for number, slide in enumerate(prs.slides, 1):
        for index, shape in enumerate(slide.shapes):
            label = f"slide {number}, shape {index}"
            x, y, w, h = map(pt, (shape.left, shape.top, shape.width, shape.height))
            if shape.has_table:
                table = shape.table
                row_y = y
                for i, row in enumerate(table.rows):
                    col_x = x
                    for j, column in enumerate(table.columns):
                        cell = table.cell(i,j)
                        cw, rh = pt(column.width), pt(row.height)
                        canvas.setFillColor(color(cell.fill.fore_color, "#ffffff"))
                        canvas.rect(col_x, pt(prs.slide_height)-row_y-rh, cw, rh, fill=1, stroke=0)
                        render_text(canvas, cell.text_frame, col_x, row_y, cw, rh,
                                    f"{label}, cell {i},{j}", issues)
                        col_x += cw
                    row_y += pt(row.height)
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill_shape(canvas, shape, x,y,w,h)
            elif shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                raise ValueError(f"Unsupported object: {shape.shape_type}")
            if shape.has_text_frame and shape.text:
                render_text(canvas, shape.text_frame, x,y,w,h, label, issues)
        canvas.showPage()
    canvas.save()
    report = {"slides": len(prs.slides), "renderer": "reportlab; native vector shapes and text",
              "overflow_issues": issues}
    (HERE / "qa/layout-validation.json").write_text(json.dumps(report, indent=2) + "\n")
    if issues:
        raise SystemExit("Layout requires adjustment:\n" + "\n".join(issues))
    print(f"Rendered {target}: {len(prs.slides)} slides; no text overflow.")


if __name__ == "__main__":
    main()
