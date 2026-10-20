"""Render trust4rsi-paper-plan.pptx to vector PDF with the incident-library deck's renderer.

Reuses ../incident-library-deck/render_pdf.py (text, tables, simple shapes only) and
writes an overflow check to qa/layout-validation.json. Requires python-pptx and reportlab.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "incident-library-deck"))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from reportlab.pdfgen.canvas import Canvas  # noqa: E402

import render_pdf as base  # noqa: E402


def main():
    base.register_fonts()
    source = HERE / "trust4rsi-paper-plan.pptx"
    target = HERE / "trust4rsi-paper-plan.pdf"
    prs = Presentation(source)
    canvas = Canvas(str(target), pagesize=(base.pt(prs.slide_width), base.pt(prs.slide_height)))
    canvas.setTitle("Paper plan: containment-extension for Trust4RSI-Agent (AAAI 2027)")
    canvas.setAuthor("containment-extension")
    issues = []
    for number, slide in enumerate(prs.slides, 1):
        for index, shape in enumerate(slide.shapes):
            label = f"slide {number}, shape {index}"
            x, y, w, h = map(base.pt, (shape.left, shape.top, shape.width, shape.height))
            if shape.has_table:
                tbl = shape.table
                row_y = y
                for i, row in enumerate(tbl.rows):
                    col_x = x
                    for j, column in enumerate(tbl.columns):
                        cell = tbl.cell(i, j)
                        cw, rh = base.pt(column.width), base.pt(row.height)
                        canvas.setFillColor(base.color(cell.fill.fore_color, "#ffffff"))
                        canvas.rect(col_x, base.pt(prs.slide_height) - row_y - rh, cw, rh, fill=1, stroke=0)
                        base.render_text(canvas, cell.text_frame, col_x, row_y, cw, rh, f"{label}, cell {i},{j}", issues)
                        col_x += cw
                    row_y += base.pt(row.height)
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                base.fill_shape(canvas, shape, x, y, w, h)
            elif shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                raise ValueError(f"Unsupported object: {shape.shape_type}")
            if shape.has_text_frame and shape.text:
                base.render_text(canvas, shape.text_frame, x, y, w, h, label, issues)
        canvas.showPage()
    canvas.save()
    report = {"slides": len(prs.slides), "renderer": "reportlab; native vector shapes and text",
              "overflow_issues": issues}
    (HERE / "qa").mkdir(exist_ok=True)
    (HERE / "qa/layout-validation.json").write_text(json.dumps(report, indent=2) + "\n")
    if issues:
        raise SystemExit("Layout requires adjustment:\n" + "\n".join(issues))
    print(f"Rendered {target}: {len(prs.slides)} slides; no text overflow.")


if __name__ == "__main__":
    main()
