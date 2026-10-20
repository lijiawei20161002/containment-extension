"""Editable vector component icons and orthogonal arrows for measured trace diagrams."""
from deck_lib import BLUE, GRAY, INK, LINE, WHITE, rect, textbox
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def oval(s, x, y, w, h, fill, line=None):
    shape = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def segment(s, x1, y1, x2, y2, color=GRAY, width=.025):
    weight = Inches(width)
    if x1 == x2:
        return rect(s, x1-weight/2, min(y1, y2), weight, max(abs(y2-y1), weight),
                    color, None, rounded=False)
    if y1 == y2:
        return rect(s, min(x1, x2), y1-weight/2, max(abs(x2-x1), weight), weight,
                    color, None, rounded=False)
    raise ValueError("Use orthogonal routes")


def arrow(s, points, color=GRAY):
    for a, b in zip(points, points[1:], strict=False):
        segment(s, *a, *b, color)
    x1, y1 = points[-2]
    x2, y2 = points[-1]
    size = Inches(.14)
    head = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                              x2-size/2, y2-size/2, size, size)
    head.rotation = 90 if x2 > x1 else 270 if x2 < x1 else 180 if y2 > y1 else 0
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()


def icon(s, kind, x, y, scale=.8, color=BLUE):
    """A small physical metaphor. Geometry is a diagram; outcomes come from logs."""
    u = Inches(scale)
    if kind == "agent":
        rect(s, x+u*.12, y+u*.16, u*.76, u*.56, WHITE, color, width=1.8)
        oval(s, x+u*.29, y+u*.34, u*.1, u*.1, color)
        oval(s, x+u*.61, y+u*.34, u*.1, u*.1, color)
        segment(s, x+u*.5, y, x+u*.5, y+u*.16, color)
        oval(s, x+u*.45, y, u*.1, u*.1, color)
        rect(s, x+u*.29, y+u*.79, u*.42, u*.08, color, None)
    elif kind == "database":
        rect(s, x+u*.13, y+u*.17, u*.74, u*.60, WHITE, color, width=1.5,
             rounded=False)
        oval(s, x+u*.13, y+u*.65, u*.74, u*.24, WHITE, color)
        oval(s, x+u*.13, y+u*.36, u*.74, u*.24, WHITE, color)
        oval(s, x+u*.13, y+u*.08, u*.74, u*.24, WHITE, color)
    elif kind == "server":
        for i in range(3):
            top = y+u*(.1+i*.26)
            rect(s, x+u*.09, top, u*.82, u*.20, WHITE, color, width=1.5)
            oval(s, x+u*.72, top+u*.07, u*.06, u*.06, color)
            rect(s, x+u*.18, top+u*.08, u*.32, u*.04, color, None)
    elif kind == "folder":
        rect(s, x+u*.1, y+u*.15, u*.38, u*.22, WHITE, color, width=1.5)
        rect(s, x+u*.1, y+u*.32, u*.8, u*.48, WHITE, color, width=1.5)
        rect(s, x+u*.24, y+u*.47, u*.50, u*.05, color, None)
        rect(s, x+u*.24, y+u*.61, u*.32, u*.05, color, None)
    elif kind in {"score", "observer"}:
        rect(s, x+u*.12, y+u*.1, u*.76, u*.62, WHITE, color, width=1.5)
        textbox(s, x+u*.17, y+u*.25, u*.66, u*.33,
                "PASS" if kind == "score" else "LOG", 12*scale, color,
                bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        segment(s, x+u*.5, y+u*.72, x+u*.5, y+u*.85, color)
        segment(s, x+u*.3, y+u*.85, x+u*.7, y+u*.85, color)
    else:
        raise ValueError(kind)


def component(s, x, y, w, h, title, subtitle, kind, color, pale, number=None):
    rect(s, x, y, w, h, pale, color, width=1)
    icon(s, kind, x+(w-Inches(.8))/2, y+Inches(.16), .8, color)
    textbox(s, x+Inches(.08), y+Inches(1.02), w-Inches(.16), Inches(.46),
            title, 16, color, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, x+Inches(.08), y+Inches(1.52), w-Inches(.16), h-Inches(1.6),
            subtitle, 13, INK, align=PP_ALIGN.CENTER)
    if number is not None:
        oval(s, x-Inches(.1), y-Inches(.1), Inches(.36), Inches(.36), color)
        textbox(s, x-Inches(.1), y-Inches(.1), Inches(.36), Inches(.36), str(number),
                13, WHITE, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def ratebar(s, x, y, w, stat, color=BLUE):
    """A measured proportion with its stored Wilson interval, using native shapes."""
    segment(s, x, y, x+w, y, LINE, .035)
    rate, interval = stat["rate"], stat["wilson95"]
    if rate is None:
        textbox(s, x, y-Inches(.2), w, Inches(.4), "Unknown", 15, GRAY)
        return
    lo, hi = interval
    segment(s, x+w*lo, y, x+w*hi, y, color, .055)
    for endpoint in (lo, hi):
        segment(s, x+w*endpoint, y-Inches(.10), x+w*endpoint, y+Inches(.10), color)
    oval(s, x+w*rate-Inches(.075), y-Inches(.075), Inches(.15), Inches(.15), color)
