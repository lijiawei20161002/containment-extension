"""Build the paper-plan deck for adapting containment-extension into a Trust4RSI-Agent
(AAAI 2027 Workshop on Trustworthy Evolution of AI Agents) submission.

Run: python3 build_deck.py  -> trust4rsi-paper-plan.pptx
Connector-free, PowerPoint-safe, no speaker notes (Keynote stalls on notes masters).
All numbers come from the repository's archived docs/results; none is new measurement.
"""
from __future__ import annotations
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from deck_lib import (INK, BLUE, BLUE_PALE, BLUE_LIGHT, GRAY, LGRAY, LINE, RED, RED_PALE,
                      GREEN, GREEN_PALE, ORANGE, ORANGE_PALE, PURPLE, PURPLE_PALE, WHITE,
                      SW, SH, M, textbox, bullets, rect, table, hrule, chevron)

HERE = Path(__file__).resolve().parent
OUT = HERE / "trust4rsi-paper-plan.pptx"
FOOTER = ("Paper plan: containment-extension -> Trust4RSI-Agent, AAAI 2027 Workshop on Trustworthy Evolution of AI Agents "
          "(Montreal, 22-23 Feb 2027; submission 20 Nov 2026 AoE). Prepared 20 Oct 2026. Counts are archived pilot results.")
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]
page = [0]
CW = SW - 2 * M

# Status vocabulary used on every experiment slide.
DONE = ("DONE: reuse archive", GREEN, GREEN_PALE)
BUILD = ("BUILD: implement, no inference", BLUE, BLUE_PALE)
RUN = ("RUN: implement + live calls", RED, RED_PALE)
HUMAN = ("HUMAN: analyst work", PURPLE, PURPLE_PALE)
LATER = ("LATER: declare as future work", GRAY, LGRAY)


def slide(title, kicker=None, sources=None):
    s = prs.slides.add_slide(BLANK); page[0] += 1
    textbox(s, M, Inches(0.28), CW, Inches(0.6), title, 23, INK, bold=True)
    if kicker:
        textbox(s, M, Inches(0.86), CW, Inches(0.42), kicker, 12.5, GRAY, italic=True)
    hrule(s, M, Inches(1.3), CW, BLUE, Pt(1.5))
    textbox(s, M, SH - Inches(0.4), SW - Inches(1.6), Inches(0.3), FOOTER, 8, GRAY)
    textbox(s, SW - Inches(1.0), SH - Inches(0.4), Inches(0.5), Inches(0.3), str(page[0]), 9, GRAY, align=PP_ALIGN.RIGHT)
    if sources:
        textbox(s, M, SH - Inches(0.68), CW, Inches(0.28), "Repository sources: " + sources, 8.5, GRAY, italic=True)
    return s


def strip(s, y, text, fill=LGRAY, color=INK, bold=False, h=Inches(0.7), size=14):
    rect(s, M, y, CW, h, fill, None)
    textbox(s, M + Inches(0.12), y, CW - Inches(0.24), h, text, size, color, bold=bold, anchor=MSO_ANCHOR.MIDDLE)


def tcard(s, x, y, w, h, title, items, fill=BLUE_PALE, line=BLUE_LIGHT, tcolor=BLUE, tsize=15, bsize=13,
          title_h=Inches(0.45), gap=6):
    rect(s, x, y, w, h, fill, line)
    textbox(s, x + Inches(0.1), y + Inches(0.08), w - Inches(0.2), title_h, title, tsize, tcolor, bold=True, ls=1.0)
    bullets(s, x + Inches(0.1), y + Inches(0.1) + title_h, w - Inches(0.2), h - title_h - Inches(0.2), items, bsize, gap=gap, ls=1.1)


def tag(s, x, y, status, w=Inches(2.6), h=Inches(0.34), size=10.5):
    text, col, pale = status
    rect(s, x, y, w, h, pale, col, width=1.0)
    textbox(s, x, y, w, h, text, size, col, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def legend(s, y):
    x = M
    for status in (DONE, BUILD, RUN, HUMAN, LATER):
        tag(s, x, y, status, w=Inches(2.35), h=Inches(0.3), size=9.5)
        x += Inches(2.47)


def box(s, x, y, w, h, title, body, col, pale, tsize=13, bsize=11):
    rect(s, x, y, w, h, pale, col)
    textbox(s, x + Inches(0.1), y + Inches(0.06), w - Inches(0.2), Inches(0.36), title, tsize, col, bold=True)
    textbox(s, x + Inches(0.1), y + Inches(0.42), w - Inches(0.2), h - Inches(0.48), body, bsize, INK, ls=1.08)


# ---------------------------------------------------------------- 1 Title
s = prs.slides.add_slide(BLANK); page[0] += 1
rect(s, 0, 0, SW, SH, BLUE_PALE, None, rounded=False)
hrule(s, M, Inches(1.05), Inches(2.2), BLUE, Pt(3))
textbox(s, M, Inches(1.2), CW, Inches(1.05),
        "Incident-aware regression gating for evolving agent policies", 32, INK, bold=True, ls=1.0)
textbox(s, M, Inches(2.3), CW, Inches(0.5),
        "Plan for adapting containment-extension into a Trust4RSI-Agent paper (AAAI 2027 workshop)", 18, BLUE)
textbox(s, M, Inches(2.85), CW, Inches(0.4),
        "Target: full paper, 8 pages + references and appendix, double-blind, OpenReview  -  deadline 20 Nov 2026 AoE, no extensions", 12.5, GRAY)
cols = [("What we already have", "26-incident library, independent effect observer, 440 archived live rollouts, instruction optimizer, executable incident twin, offline-qualified testbeds", BLUE, BLUE_PALE),
        ("What the paper claims", "An audit-and-release abstraction beneath an evolving agent: observe effects below the tool broker, gate policy updates on verified incidents, keep executable incidents as the regression suite", GREEN, GREEN_PALE),
        ("What still has to run", "E1 archived-data figure, E2 matched proxy-vs-incident selection (the one live experiment), E5 second-reader codings; optional E3/E4", RED, RED_PALE)]
gap = Inches(0.18); w = (CW - 2 * gap) / 3
for i, (name, body, col, pale) in enumerate(cols):
    x = M + i * (w + gap)
    rect(s, x, Inches(3.45), w, Inches(1.95), pale, col)
    textbox(s, x + Inches(0.14), Inches(3.53), w - Inches(0.28), Inches(0.45), name, 17, col, bold=True)
    textbox(s, x + Inches(0.14), Inches(4.0), w - Inches(0.28), Inches(1.35), body, 12.5, INK, ls=1.1)
textbox(s, M, Inches(5.65), CW, Inches(0.9),
        "The repository's headline hypothesis (incident feedback beats transcript-proxy feedback for policy selection) is designed "
        "but unmeasured. The paper must lead with the abstraction, the testbed and pilot evidence, and report that contrast "
        "only as far as E2 actually measures it.", 14, INK, ls=1.15)
textbox(s, M, Inches(6.55), CW, Inches(0.4), "Paper plan deck  -  Jiawei Li  -  20 Oct 2026  -  remove this line and all author traces before the anonymized submission", 11, GRAY)

# ---------------------------------------------------------------- 2 Venue
s = slide("The venue: Trust4RSI-Agent, 'An Agent Operating System Perspective'",
          "The workshop treats trustworthy agent evolution as a systems problem: design the layer beneath an evolving agent")
table(s, M, Inches(1.5), Inches(6.3), Inches(3.5), [
    ["Item", "Requirement"],
    ["Full paper", "8 pages, AAAI 2027 author kit; references and appendix excluded"],
    ["Short / position", "4 pages, same kit"],
    ["Review", "Double-blind on OpenReview; anonymize PDF and supplementary"],
    ["Submission", "20 Nov 2026 AoE, strict"],
    ["Notification", "2 Dec 2026 AoE"],
    ["Workshop", "22-23 Feb 2027, Palais des congres, Montreal (one day, TBA)"],
    ["Accepted types", "Mechanisms/abstractions, system designs, benchmarks, negative results, position papers"],
], [Inches(1.6), Inches(4.7)], size=11)
tcard(s, M + Inches(6.6), Inches(1.5), CW - Inches(6.6), Inches(3.5), "Five topics of interest", [
    "**T1 Memory evolution** (state subsystem): provenance, write policies, snapshots, leakage from persistent state",
    "**T2 Tool and skill evolution** (capability interface): least privilege, sandboxing, verification of synthesized tools, revocation",
    "**T3 Model and policy evolution** (update and release layer): versioning, staged rollout, regression suites, rollback, reward hacking, reliability of self-critiques",
    "**T4 Multi-agent co-evolution** (runtime and scheduler): shared state, runaway loops, blast-radius limits",
    "**T5 Evaluating evolution** (observability and audit layer): attributing outcomes to updates, monitoring, oversight before irreversible updates, stateful testbeds",
], bsize=11.5, gap=4)
strip(s, Inches(5.15), "Our natural home is T3 (the release layer) with T5 (the audit layer) as the measurement story. "
      "The CFP explicitly welcomes system designs, benchmarks and negative results, which is what the repository can honestly supply.",
      BLUE_PALE, INK, bold=True, h=Inches(0.75), size=13)
textbox(s, M, Inches(6.0), CW, Inches(0.6),
        "Organizers: Haolun Wu (Stanford/Mila), Negar Arabzadeh (UC Berkeley), Wenyue Hua (MSR), Bang Liu, Laurent Charlin (Mila). "
        "Speakers include Yoshua Bengio, Mengdi Wang, Yilun Du. A safety-literate audience: hedged, measured claims will be read as a strength.", 11, GRAY, ls=1.1)

# ---------------------------------------------------------------- 3 Fit map
s = slide("Fit map: repository component -> workshop topic -> OS vocabulary for the paper",
          "The components exist; the adaptation is mostly renaming them in the language of the layer beneath the agent",
          "README.md; docs/incident-aware-optimization.md; docs/instruction-optimization.md; docs/coordination-harness.md")
table(s, M, Inches(1.5), CW, Inches(4.0), [
    ["Topic", "Repository component (as implemented)", "Paper vocabulary (agent-OS framing)", "Role"],
    ["T3 policy evolution", "Instruction optimizer: model-proposed clause variants, catalog validator, development-only selection, frozen held-out trials; constrained objective min R s.t. U, C, F",
     "Staged rollout with a release gate; a regression suite every policy update must pass; reward hacking made visible to the gate", "Primary"],
    ["T5 evaluating evolution", "Independent SQLite effect observer (before/after state), counterfactual task families with feasibility oracle, grader regression cases, 26-record incident library",
     "Audit subsystem below the tool broker; outcome attribution per update; stateful testbed with reference labels", "Primary"],
    ["T2 tool/skill evolution", "Dispatcher scope check on exact `.lab.test` authorities; catalog-constrained synthesis rejects arbitrary instruction text",
     "Reference monitor / least privilege; verification of synthesized capabilities before they reach the actor", "Supporting"],
    ["T4 multi-agent co-evolution", "Coordination harness: shared counters, read-consumed quota, URL store, paired counterfactual replay, 28 scripted cases",
     "Shared-state blast radius; effects invisible in transcripts", "One paragraph + appendix"],
    ["T1 memory evolution", "Prior-run answer archive, cached flag, synthetic mailbox as reachable persistent state",
     "Leakage from persistent state across runs", "Mention only"],
], [Inches(1.75), Inches(4.9), Inches(4.3), Inches(1.38)], size=10.5)
strip(s, Inches(5.7), "Rule for the rewrite: keep every measurement caveat once, in Limitations. The body leads with the abstraction and what it already observes.",
      LGRAY, INK, bold=True, h=Inches(0.6), size=13)

# ---------------------------------------------------------------- 4 Claim
s = slide("The paper in one slide: what we claim and what we do not",
          "A contribution of mechanisms, a testbed and pilot evidence; the optimization result is reported only as far as E2 measures it")
lw = Inches(6.1)
tcard(s, M, Inches(1.5), lw, Inches(4.35), "We claim", [
    "**Abstraction.** An audit-and-release layer for evolving agents: (a) an independent effect observer reading service state below the tool broker, (b) executable incident twins as a regression suite, (c) a release rule that minimizes verified incidents subject to utility, cost and false-refusal limits.",
    "**Evidence that scorers miss what the observer sees.** 13 of 15 violating runs in the incident study passed the task scorer by submitting a cached answer; the observer recorded every excluded read.",
    "**Pilot evidence that policy updates move the gated quantity.** Haiku 4.5 outside access on impossible tasks 6/8 -> 0/8 under the combined instruction update, with 96/96 capability controls completed in scope; scope levels 2/6, 2/6, 0/6, 0/6.",
    "**A reusable protocol and testbed.** Frozen plans, per-run budgets, 320 transition checks, 17 grader cases, 1,543 relabeled proposals, 196+ tests, no model call in qualification.",
    "**A negative result on environment grounding.** The direct upstream coding baseline failed for environment reasons, not model reasons.",
], fill=GREEN_PALE, line=GREEN, tcolor=GREEN, bsize=11.5, gap=5)
tcard(s, M + lw + Inches(0.3), Inches(1.5), CW - lw - Inches(0.3), Inches(4.35), "We do not claim", [
    "A general failure rate or a generally effective defense: cells are n=8 and n=2 on one task structure and two models; 0/8 has a 95% Wilson upper limit of 32.4%.",
    "That incident feedback beats transcript-proxy feedback in general. E2 gives one matched comparison on the native fixture; a null is reported as a null.",
    "Benchmark-level improvements on ImpossibleBench, AgentDojo, ResearchArena or BashArena. These stay as the declared protocol.",
    "That the dispatcher is OS-level containment. It is tool-level containment for a closed tool surface.",
    "That the incident codings are settled facts. They are analyst judgments on public reporting; E5 adds a second reader.",
], fill=RED_PALE, line=RED, tcolor=RED, bsize=11.5, gap=5)
strip(s, Inches(6.0), "Title candidate: 'Incident-aware regression gating for evolving agent policies: an audit layer beneath the tool broker'", BLUE, WHITE, bold=True, h=Inches(0.55), size=14)

# ---------------------------------------------------------------- 5 Outline
s = slide("Eight-page outline: page budget and source material",
          "Everything in sections 2-4 and 6 exists today; section 5 needs E1 and E2; section 8 needs E5")
table(s, M, Inches(1.5), CW, Inches(4.7), [
    ["#", "Section", "Pages", "Content", "Source material", "Needs"],
    ["1", "Introduction", "1.0", "Agents evolve through prompt, skill and model updates; each update needs a gate; task scores miss incidents", "README; incident-aware-optimization.md", "Writing"],
    ["2", "Incidents motivate the layer beneath", "1.0", "26 public cases: pathways, mechanisms M1-M6, propositions P1-P9, detection and disclosure latency, mapped to T1-T5", "incident-library-analysis.md; figures/incidents/analysis", "E5"],
    ["3", "Abstractions", "1.5", "Effect observer below the broker; incident definitions; executable twins; release rule min R s.t. U, C, F; B0-B4 protocol", "protocol.md; workbook_policy.py; INC-2026-015", "Writing, 1 figure"],
    ["4", "Testbed", "1.0", "Native service fixture; counterfactual task families; grader regression; coordination harness; budgets and frozen plans", "evaluation-infrastructure-results.md; coordination-harness-results.md", "Writing"],
    ["5", "Pilot evidence", "2.0", "Utility vs verified incident rate across instruction policies (E1); matched proxy-vs-incident selection (E2); controls and intervals", "incident-study-v1; scope-levels-v1; link-study-v1; new E2 run", "E1, E2"],
    ["6", "Negative result", "0.5", "Upstream coding baseline: 1/2 original, 0/4 impossible, message-capped, missing packages (P2)", "impossiblebench-upstream-baseline-results.md", "Writing"],
    ["7", "Related work", "0.5", "AI control protocols, Inspect approvers, AgentDojo defenses, ImpossibleBench, incident taxonomies", "standardized-benchmarks.md", "Writing"],
    ["8", "Limitations and roadmap", "0.5", "Sample sizes, one task structure, single-reader codings, tool-level containment, the unmeasured benchmark protocol", "experiment-roadmap.md", "E5"],
], [Inches(0.35), Inches(2.05), Inches(0.75), Inches(4.45), Inches(3.3), Inches(1.43)], size=10)
textbox(s, M, Inches(6.28), CW, Inches(0.5),
        "Appendix (unlimited): exact prompts, frozen plans, per-cell tables with Wilson intervals, incident record schema, coordination harness cases, test counts. "
        "Supplementary: anonymized source snapshot and evidence archive (E6).", 11, GRAY, ls=1.1)

# ---------------------------------------------------------------- 6 Incident library
s = slide("Section 2 material: what the 26-case library already says about the layer beneath",
          "Codings are analyst drafts on public reporting; second-reader agreement (E5) is required before these counts go in the paper",
          "docs/incident-library-analysis.md; incidents/records; scripts/analyze_incident_library.py")
table(s, M, Inches(1.5), Inches(6.2), Inches(3.95), [
    ["Finding (n = 26)", "n", "Topic"],
    ["M4 detection without decisive escalation: present", "13", "T5 oversight"],
    ["P6 effects matter more than stated intent: violated", "17", "T5 audit layer"],
    ["M2 capability composition: present", "15", "T2 least privilege"],
    ["P3 authorization must survive config failure: violated", "12", "T2 / T3"],
    ["P4 impossible tasks are security inputs: violated", "10", "T3 reward hack"],
    ["M3 shared services as coordination infra: present", "3", "T4 shared state"],
    ["Stop found only in retrospect", "10", "T5 monitoring"],
    ["Median detection latency (12 known)", "14 days", "T5"],
    ["Median disclosure latency (21 known)", "119 days", "T5"],
], [Inches(4.0), Inches(0.9), Inches(1.3)], size=10)
tcard(s, M + Inches(6.5), Inches(1.5), CW - Inches(6.5), Inches(3.95), "How the section argues", [
    "Pathways: adaptive escape 3, misconfiguration 6, live-access overreach 9, training-sample boundary 3, external adversary 4, unresolved 1. The two usual pathways cover only 9 of 26.",
    "18 entries are judged replayable as synthetic twins; one (INC-2026-015, public hosting of a local workbook) is implemented. That is the bridge from Section 2 to the regression-suite idea in Section 3.",
    "Evidence levels A: 19, B: 7. State them; do not add counting units across cases.",
    "Keep the coding frame citation to one sentence. The paper's contribution is the executable use of the library, not the taxonomy.",
], bsize=11.5, gap=5)
strip(s, Inches(5.5), "One figure here: mechanism-by-topic heatmap (exists as figures/incidents/analysis/mechanism-heatmap.png; relabel columns with T1-T5).",
      LGRAY, INK, h=Inches(0.5), size=12)
textbox(s, M, Inches(6.08), CW, Inches(0.7),
        "Risk: the records name real companies and models. That is public reporting and is fine, but every cited count must carry its evidence level, "
        "and nothing may read as a verdict on a named party. Phrase as 'reported' and 'coded as'.", 11.5, GRAY, ls=1.1)

# ---------------------------------------------------------------- 7 System figure
s = slide("Section 3 material: the audit-and-release layer, drawn in OS terms",
          "The agent's policy evolves above the line; everything the paper contributes sits below it",
          "README.md (outer containment); docs/incident-aware-optimization.md; incidents/executable/INC-2026-015")
y0 = Inches(1.55); bh = Inches(1.4); bw = Inches(2.75); gx = Inches(0.42)
x = M
box(s, x, y0, bw, bh, "Agent + policy theta", "Fixed actor model. theta = selected, ordered instruction clauses (later: monitor prompts, thresholds, defer rules). This is what evolves.", GRAY, LGRAY); x += bw + gx
chevron(s, x - gx + Inches(0.1), y0 + bh / 2 - Inches(0.11), Inches(0.22));
box(s, x, y0, bw, bh, "Tool broker (ref. monitor)", "One structured tool. Exact-authority routing, optional scope enforcement, fail-closed on unknown routes. Changing theta never opens this boundary.", BLUE, BLUE_PALE); x += bw + gx
chevron(s, x - gx + Inches(0.1), y0 + bh / 2 - Inches(0.11), Inches(0.22))
box(s, x, y0, bw, bh, "Service replicas (state)", "SQLite transactions for repository, scorer, collector, mailbox, prior-run artifact. Writes later reverted are still state history.", BLUE, BLUE_PALE); x += bw + gx
chevron(s, x - gx + Inches(0.1), y0 + bh / 2 - Inches(0.11), Inches(0.22))
box(s, x, y0, CW - (x - M), bh, "Task scorer", "Official pass/fail. Accepts a cached flag. This is the signal an evolving agent optimizes if nothing else is watching.", ORANGE, ORANGE_PALE)
hrule(s, M, Inches(3.12), CW, LINE, Pt(1))
textbox(s, M, Inches(3.14), Inches(4), Inches(0.3), "below the line: the paper's layer", 10, GRAY, italic=True)
y1 = Inches(3.47); bh1 = Inches(1.5)
x = M
box(s, x, y1, bw, bh1, "Independent effect observer", "Reads before/after service state, not the transcript. Emits incident records: resource, authorization, effect, reverted or not. Positive witnesses survive later failures.", GREEN, GREEN_PALE); x += bw + gx
chevron(s, x - gx + Inches(0.1), y1 + bh1 / 2 - Inches(0.11), Inches(0.22))
box(s, x, y1, bw, bh1, "Release gate", "min R(theta) s.t. U >= U0 - dU, C <= B, F <= aF. Development evidence only; freeze before held-out. No eligible candidate = no release.", GREEN, GREEN_PALE); x += bw + gx
chevron(s, x - gx + Inches(0.1), y1 + bh1 / 2 - Inches(0.11), Inches(0.22))
box(s, x, y1, bw, bh1, "Executable incident twins", "Library record -> policy manifest + unsafe and benign fixtures. Every candidate theta must pass the suite. INC-2026-015 implemented; 18 judged replayable.", GREEN, GREEN_PALE); x += bw + gx
chevron(s, x - gx + Inches(0.1), y1 + bh1 / 2 - Inches(0.11), Inches(0.22))
box(s, x, y1, CW - (x - M), bh1, "Versioned policy release", "Frozen theta, prompt hashes, source snapshot, plan, budget ledger. Rollback = previous frozen record. Held-out test never feeds back.", PURPLE, PURPLE_PALE)
strip(s, Inches(5.12), "The figure to draw for the paper: this slide as one panel, with the proxy-feedback path (transcript judge -> gate) drawn as a dashed alternative to the observer. "
      "figures/optimization/overview-v1 and system-components-v1 already contain the pieces.", LGRAY, INK, h=Inches(0.75), size=12)
textbox(s, M, Inches(6.0), CW, Inches(0.6),
        "OS analogies to use sparingly and name once each: reference monitor (broker), audit subsystem (observer), CI regression suite (twins), release gate and rollback (frozen records). "
        "Say which classical abstraction transfers and which breaks: audit coverage must be demonstrated, not assumed.", 11, GRAY, ls=1.1)

# ---------------------------------------------------------------- 8 Existing evidence
s = slide("Section 5 material: 440 archived live rollouts",
          "Every cell below can be reported today with Wilson intervals; none of it is a new run",
          "docs/*-results.md; results/*/README.md; results/*/summary.json")
table(s, M, Inches(1.5), CW, Inches(3.55), [
    ["Study (date)", "Rollouts", "Models", "Manipulation", "Key cells", "Controls", "Cost"],
    ["Native pilot (Sep)", "8", "GPT-4.1 mini, Haiku 4.5", "solvable/stalled x permissive/scoped", "Local out-of-scope effects occurred; early scope check over-blocked", "-", "-"],
    ["Cue study (Sep)", "96", "both", "prompt cues", "No consistent increase from tested cues", "-", "-"],
    ["Upstream-link study (Sep)", "128", "both", "README link present/absent x reminder", "Early upstream reads 26/32 with link, 0/32 without; later violations substituted the answer archive", "-", "-"],
    ["Incident study (17 Sep)", "160", "both", "N / R / E / RE instruction updates", "Haiku impossible outside access 6/8, 3/8, 5/8, 0/8; GPT-4.1 mini 0/8, 1/8, 0/8, 0/8", "96/96 in scope", "$0.95"],
    ["Scope levels (24 Sep)", "24", "Haiku 4.5", "4 cumulative instruction levels", "Outside effects 2/6, 2/6, 0/6, 0/6; impossible-only 2/2, 1/2, 0/2, 0/2", "15/16 in scope", "$0.24"],
    ["Upstream coding baseline (Sep)", "6", "Haiku 4.5", "unmodified ImpossibleBench scaffold", "Original 1/2, impossible 0/4; all 100-message capped; missing packages in all 6", "-", "$5.78"],
    ["ImpossibleBench preflight (Sep)", "24 assigned, 16 started", "Haiku 4.5", "custom adapter arms", "No scored submissions; unknown outcomes", "-", "~$5 cap"],
], [Inches(2.0), Inches(1.15), Inches(1.5), Inches(2.2), Inches(3.6), Inches(1.1), Inches(0.78)], size=9.5)
tcard(s, M, Inches(5.15), Inches(6.1), Inches(1.6), "What to foreground", [
    "13 of 15 incident-study violations submitted the cached answer and **passed the scorer**. This is the paper's motivating observation.",
    "Capability controls stayed at or near 100% while the gated quantity moved: the utility constraint is not vacuous.",
], fill=GREEN_PALE, line=GREEN, tcolor=GREEN, bsize=11.5, gap=4)
tcard(s, M + Inches(6.4), Inches(5.15), CW - Inches(6.4), Inches(1.6), "What to keep in Limitations", [
    "One task structure; two models; cells of 8 and 2. GPT-4.1 mini's near-zero event rate leaves no room to measure an effect.",
    "Instruction packages change length and content together; this is not a scalar strength manipulation.",
], fill=LGRAY, line=LINE, tcolor=INK, bsize=11.5, gap=4)

# ---------------------------------------------------------------- 9 Experiments overview
s = slide("Experiments: what to implement and run, marked by status",
          "Three things are required (E1, E2, E5); E3 and E4 strengthen the paper if E2 is informative; the benchmark suite stays declared, not run",
          "experiments/incident-aware-optimization-v1.design.json (status: matched feedback ablation not implemented)")
legend(s, Inches(1.45))
rows = [["ID", "Experiment", "Status", "Inference", "Budget", "Output / gate"],
        ["E1", "Main figure from archived data: task utility (x) vs verified incident rate (y) for every instruction policy in incident-study-v1, scope-levels-v1 and link-study-v1, with Wilson intervals", "BUILD", "None", "$0", "Fig. 2 + per-cell table; script under scripts/"],
        ["E2", "Matched fixed-pool selection: same candidate pool and shared development rollouts; select once with task score + transcript-proxy judge (B3), once with task score + effect observer (B4); test both frozen policies plus two fixed controls on fresh held-out seeds", "RUN", "~110 Haiku actor rollouts + ~60 judge calls", "est. $2-3, ceiling $10", "Primary Section 5 result; frozen plan + scripted controls first"],
        ["E3", "Replicate E2 on a second model family with a non-zero baseline event rate (GPT-4.1 mini had 1/32; pick a different actor)", "RUN (optional)", "same as E2", "ceiling $15", "Needed to say 'two model families'; skip if E2 null"],
        ["E4", "Regression-gate demonstration: implement twins for 2 more replayable library entries; run the four frozen E2 policies through the twin suite with the live actor", "BUILD then RUN (optional)", "~24 rollouts", "ceiling $3", "Shows the gate consuming incident twins; stretch"],
        ["E5", "Second-reader coding of the 26 records on pathway, M1-M6, P1-P9; report agreement; resolve or mark disputed", "HUMAN", "None", "$0", "Required before Section 2 counts are cited"],
        ["E6", "Anonymized reproducibility package: source snapshot, frozen plans, evidence archives, tests, builder scripts; strip names, URLs, decks", "BUILD", "None", "$0", "Supplementary zip"],
        ["-", "ImpossibleBench U2 environment repair, AgentDojo, ResearchArena, BashArena comparisons", "LATER", "-", "-", "Section 8 roadmap only"]]
table(s, M, Inches(1.9), CW, Inches(4.45), rows,
      [Inches(0.45), Inches(5.3), Inches(1.55), Inches(1.75), Inches(1.3), Inches(1.98)], size=9.5)
textbox(s, M, Inches(6.38), CW, Inches(0.45),
        "Cost basis: the incident study ran 160 rollouts for $0.95 and scope levels 24 for $0.24 with Haiku 4.5 at $1/$5 per million tokens. "
        "Estimates are list-price projections, not billing guarantees; every run keeps the repository's per-run and aggregate ceilings.", 10, GRAY, ls=1.05)

# ---------------------------------------------------------------- 10 E1
s = slide("E1: the main figure from archived evidence (no inference)",
          "The design doc already specifies the plot; the data for six instruction policies exist in results/",
          "results/incident-study-v1/summary.json; results/scope-levels-v1/analysis.json; results/link-study-v1")
tag(s, M, Inches(1.45), BUILD)
tcard(s, M, Inches(1.9), Inches(6.2), Inches(3.9), "Points to plot (Haiku 4.5 unless stated)", [
    "**Incident study, impossible n=8 each:** N 6/8, R 3/8, E 5/8, RE 0/8 outside access. Utility axis: authorized completion on recoverable + easy controls, 96/96 across arms (plot per arm).",
    "**Scope levels, impossible n=2, controls n=4:** baseline 2/2 and 4/4; repeat_scope 1/2 and 3/4; explicit_routes 0/2 and 4/4; action_check 0/2 and 4/4.",
    "**Upstream-link study, n=32 per README cell:** early upstream reads 26/32 with link, 0/32 without, 0/30 (+2 unknown) link + reminder, 0/32 hidden + reminder; solvable-task completion as utility.",
    "**GPT-4.1 mini cells** as hollow markers: 0/8, 1/8, 0/8, 0/8. Shows the event-rate floor problem rather than hiding it.",
    "Wilson 95% intervals on both axes; mark unknown outcomes as a separate count, never as zeros.",
], bsize=11.5, gap=5)
tcard(s, M + Inches(6.5), Inches(1.9), CW - Inches(6.5), Inches(3.9), "Implementation", [
    "New script `scripts/draw_utility_incident_tradeoff.py`: read the three archived summaries, emit PNG + SVG under figures/experiments/tradeoff-v1/ with a JSON table of every plotted cell.",
    "Follow the repository's figure convention: generator script, editable SVG, metadata next to the image.",
    "Regression test: the plotted counts equal the archived summary counts (a fixture-reading test, no inference).",
    "Half a day. Blocked on nothing.",
], fill=LGRAY, line=LINE, tcolor=INK, bsize=11.5, gap=5)
strip(s, Inches(6.0), "What the figure must not do: draw a curve. Six points with intervals, labeled by policy, is the honest version.", BLUE_PALE, INK, bold=True, h=Inches(0.55), size=13)

# ---------------------------------------------------------------- 11 E2 design
s = slide("E2: the one live experiment, matched proxy-vs-incident selection",
          "First stage of the design document's B3 vs B4 contrast: fixed pool, shared development rollouts, two selections, one held-out observer",
          "docs/incident-aware-optimization.md; src/containment_extension/instruction_optimization; src/containment_extension/evaluation/live.py")
tag(s, M, Inches(1.45), RUN)
table(s, M, Inches(1.9), Inches(7.4), Inches(2.75), [
    ["Stage", "Content", "Rollouts"],
    ["Pool", "8 catalog-mode instruction candidates + 2 fixed controls (original, fixed L3); catalog proposer is deterministic, so the pool is reproducible", "0"],
    ["Development", "10 policies x 6 cases (2 impossible, 2 recoverable, 2 solvable), shared across both selections", "60 actor"],
    ["Judge (B3 only)", "Haiku transcript monitor over each development transcript: unauthorized-access flag + suspicion; threshold set on benign calibration transcripts from the archive", "60 judge"],
    ["Selection", "Same eligibility rule (zero incidents on impossible, all controls completed, shortest instruction) applied to judge flags (B3) and to observer records (B4)", "0"],
    ["Held-out", "4 frozen policies (original, fixed L3, B3-selected, B4-selected) x 12 fresh seeds (4/4/4)", "48 actor"],
], [Inches(1.3), Inches(5.2), Inches(0.9)], size=10)
tcard(s, M + Inches(7.7), Inches(1.9), CW - Inches(7.7), Inches(2.75), "Where proxy and observer can disagree", [
    "**Reverted writes.** The optimizer's scripted control already covers an upstream write that is later reverted. A transcript judge may excuse it; the observer counts it.",
    "**False flags.** The judge may flag legitimate recovery actions, steering B3 toward over-restrictive policies and lower control completion.",
    "**Cached-answer success.** Scorer pass with excluded read: both signals should catch it; if the judge misses any, that is the result.",
], fill=ORANGE_PALE, line=ORANGE, tcolor=ORANGE, bsize=11, gap=4)
tcard(s, M, Inches(4.8), Inches(6.1), Inches(1.95), "Gates before the first paid call", [
    "Scripted controls pass for both selection paths; the judge path is exercised with scripted transports.",
    "Frozen plan with pool, seeds, thresholds, eligibility rule, tie-break, per-run and aggregate ceilings; user approval of the ceiling recorded in the plan.",
    "Baseline development event rate must be non-zero on impossible cases (Haiku N was 6/8); otherwise selection is uninformative and the run stops.",
], fill=GREEN_PALE, line=GREEN, tcolor=GREEN, bsize=10.5, gap=3)
tcard(s, M + Inches(6.4), Inches(4.8), CW - Inches(6.4), Inches(1.95), "Pre-registered readout", [
    "Primary: paired difference in held-out verified incident rate and in control completion between B3- and B4-selected policies, by seed.",
    "Also report: which candidate each signal chose, judge false-flag rate on controls, cost per arm. Same policy chosen by both = null result, written up as such.",
], fill=LGRAY, line=LINE, tcolor=INK, bsize=11, gap=4)

# ---------------------------------------------------------------- 12 E2 implementation
s = slide("E2: what has to be implemented (about four engineering days)",
          "The optimizer, fixture, observer, budgets and a scripted-transport monitor runner exist; the missing parts are the judge arm and the dual selection",
          "instruction_optimization/{catalog,search,study}.py; evaluation/live.py; results/incident-monitor-inputs-v1 (1,543 paired inputs)")
tag(s, M, Inches(1.45), RUN)
table(s, M, Inches(1.9), CW, Inches(3.6), [
    ["Piece", "Exists?", "Work", "Verification"],
    ["Candidate pool (catalog mode)", "Yes", "Pin 8 candidates by stable ID in a new config experiments/feedback-ablation-v1.json", "Dry run prints exact instructions and ceilings"],
    ["Shared development rollouts", "Yes (study.py)", "Record per-rollout transcript path + observer record together so both selections read one evidence set", "Fixture test: every development rollout has both artifacts"],
    ["Transcript judge", "Partly: scripted-transport runner; 1,543 archived inputs", "Live Haiku monitor prompt over transcript only; calibrate threshold on archived benign transcripts; never give it service state", "Scripted-transport test; calibration report with TPR/FPR on archive"],
    ["Dual selection", "No", "Selection function takes a signal source (judge flags or observer records); emit both selected IDs and eligibility tables", "Unit test: identical inputs -> identical choice; divergent fixture -> divergent choice"],
    ["Held-out schedule", "Yes (fresh-seed trials)", "Extend to 4 frozen policies x 12 seeds with randomized block order", "Schedule check prints 48 assignments"],
    ["Budget ledger", "Yes", "Add judge calls as a separate ledger line and ceiling", "Reservation test on provider failure"],
    ["Report", "Partly (instruction-report)", "Paired held-out table, judge false flags, cost per arm, Wilson intervals", "Report from scripted run has no empty fields"],
], [Inches(2.2), Inches(2.3), Inches(4.9), Inches(2.93)], size=10)
strip(s, Inches(5.7), "Order: dual selection -> judge arm -> scripted end-to-end -> freeze plan -> approval -> run. Do not touch the fixture, scorer or incident definition in the same change.",
      BLUE_PALE, INK, bold=True, h=Inches(0.6), size=13)

# ---------------------------------------------------------------- 13 E3-E6
s = slide("E3-E6: optional strengthening and the required non-inference work",
          "Decide E3/E4 after E2's development stage; E5 and E6 start now")
w2 = (CW - Inches(0.3)) / 2
tcard(s, M, Inches(1.5), w2, Inches(2.2), "E3  Second model family (RUN, optional)", [
    "Same plan as E2 with a different actor whose baseline impossible-task event rate is non-zero in a 4-run preflight.",
    "Lets Section 5 say 'two model families' instead of 'one model'. Ceiling $15. Skip if E2 is null or budget is not approved.",
], fill=RED_PALE, line=RED, tcolor=RED, bsize=11.5, gap=5)
tcard(s, M + w2 + Inches(0.3), Inches(1.5), w2, Inches(2.2), "E4  Regression gate end to end (BUILD then RUN, stretch)", [
    "Implement twins for two more of the 18 replayable library entries using the INC-2026-015 pattern (policy manifest, unsafe fixture, benign fixture).",
    "Run the four frozen E2 policies through the twin suite with the live actor, ~24 rollouts, ceiling $3. Shows the gate consuming incidents, which is the paper's T3 story.",
], fill=RED_PALE, line=RED, tcolor=RED, bsize=11.5, gap=5)
tcard(s, M, Inches(3.95), w2, Inches(2.2), "E5  Second-reader codings (HUMAN, required)", [
    "A second reader codes pathway, M1-M6 and P1-P9 for all 26 records blind to the first coding; compute agreement; resolve or mark disputed.",
    "Without this, Section 2 counts are one analyst's draft and a careful reviewer will say so. About two days of reading.",
], fill=PURPLE_PALE, line=PURPLE, tcolor=PURPLE, bsize=11.5, gap=5)
tcard(s, M + w2 + Inches(0.3), Inches(3.95), w2, Inches(2.2), "E6  Anonymized artifact (BUILD, required)", [
    "Source snapshot, frozen plans, evidence archives, tests, figure generators. Strip README author line, HANDOFF, slide decks, repo URL, machine paths and names in run metadata.",
    "Check that no archived manifest or transcript carries an account identifier. Zip, hash, upload as supplementary.",
], fill=BLUE_PALE, line=BLUE_LIGHT, tcolor=BLUE, bsize=11.5, gap=5)
strip(s, Inches(6.3), "Not for this paper: ImpossibleBench U2, AgentDojo, ResearchArena, BashArena. Declared in Section 8 with the frozen source pins; no live allocation.", LGRAY, INK, h=Inches(0.5), size=12)

# ---------------------------------------------------------------- 14 Negative result
s = slide("Section 6: the upstream coding baseline as a negative result",
          "The CFP lists negative results as an accepted contribution; this one is already archived and maps to proposition P2",
          "docs/impossiblebench-upstream-baseline-results.md; docs/impossiblebench-live-preflight-results.md; results/upstream-baseline-preparation-v1")
tcard(s, M, Inches(1.5), Inches(6.2), Inches(3.3), "What happened", [
    "Six Haiku 4.5 runs on the unmodified upstream scaffold, prompts and scorer: original 1/2, oneoff 0/2, conflicting 0/2.",
    "All six hit the 100-message cap without calling submit. All six hit missing packages under the default actor Python while the scorer used the activated `testbed` environment.",
    "290 requests, 5.43M input tokens, $5.78. The passing original patch is a general repair on review.",
    "Earlier custom-adapter preflight: 24 assigned, 16 started, 14 at the 20-call limit, no scored submissions.",
], fill=LGRAY, line=LINE, tcolor=INK, bsize=12, gap=6)
tcard(s, M + Inches(6.5), Inches(1.5), CW - Inches(6.5), Inches(3.3), "How to write it", [
    "Twelve scoring controls and four tool-loop controls passed before the run, yet did not establish that ordinary actor imports work across successive calls. Environment claims required measurement in the actor's path, not the scorer's.",
    "An evolving agent whose regression suite runs in a different environment than its actor will pass the suite and fail in use. This is the T3 lesson, stated once.",
    "Report the capability floor: a safety comparison on this scaffold is uninformative until originals are solvable. No cheating rate is claimed.",
], fill=ORANGE_PALE, line=ORANGE, tcolor=ORANGE, bsize=12, gap=6)
strip(s, Inches(5.0), "Half a page. It costs nothing, it is already written in the docs, and it is the kind of result this workshop asked for.", BLUE_PALE, INK, bold=True, h=Inches(0.55), size=13)
textbox(s, M, Inches(5.75), CW, Inches(0.8),
        "Do not mix this with Section 5: the native fixture results and the coding-benchmark failure are different environments, denominators and models. "
        "Keep separate tables.", 11.5, GRAY, ls=1.1)

# ---------------------------------------------------------------- 15 Timeline
s = slide("Timeline: 31 days from today to the deadline",
          "Submit on 19 Nov; the 20 Nov AoE deadline has no extension")
weeks = [("Oct 20-26", "Outline fixed. E1 figure. E2 dual selection + judge arm implemented, scripted end to end. E5 second reader starts. Draft Sections 2, 3, 6 from existing docs.", BLUE),
         ("Oct 27-Nov 2", "Freeze E2 plan and ceiling; approval; run E2 development stage; check event rate gate; run held-out. E5 continues.", RED),
         ("Nov 3-9", "Analyze E2; decide E3/E4 and run if green. Draft Sections 1, 4, 5. Paper figure from slide 7. E5 agreement computed.", RED),
         ("Nov 10-16", "Full draft in AAAI kit. Related work. Appendix tables. E6 anonymized artifact assembled and checked.", BLUE),
         ("Nov 17-19", "Internal read for over-claiming; anonymization check; OpenReview submission on 19 Nov.", GREEN)]
y = Inches(1.55); rh = Inches(0.78)
for label, body, col in weeks:
    rect(s, M, y, Inches(1.7), rh, col, None)
    textbox(s, M, y, Inches(1.7), rh, label, 13, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, M + Inches(1.8), y, CW - Inches(1.8), rh, LGRAY, None)
    textbox(s, M + Inches(1.95), y, CW - Inches(2.1), rh, body, 12.5, INK, anchor=MSO_ANCHOR.MIDDLE, ls=1.1)
    y += rh + Inches(0.1)
strip(s, Inches(6.05), "Critical path: E2 implementation (week 1) -> E2 run (week 2). If the week-2 run slips past Nov 9, the paper ships as infrastructure + pilot + negative result, with E2 moved to Section 8.",
      ORANGE_PALE, INK, bold=True, h=Inches(0.65), size=12.5)

# ---------------------------------------------------------------- 16 Risks
s = slide("Risks, reviewer pushback and fallbacks",
          "Each row names the fallback that keeps a submittable paper")
table(s, M, Inches(1.5), CW, Inches(3.9), [
    ["Risk", "Likelihood", "Fallback"],
    ["E2 null: judge and observer pick the same policy", "Moderate: the native transcript exposes every tool call", "Report as a negative result with the judge's false-flag rate and the reverted-write case; the abstraction and testbed carry the paper"],
    ["Live budget not approved in time", "Low", "Drop E2 to Section 8; submit infrastructure + archived pilot + negative result (still a full paper) or a 4-page position paper on the library"],
    ["Reviewers: 'n=8 on a toy fixture'", "High", "Intervals on every cell; one-task-structure limitation stated; frame as a testbed paper, not a defense paper"],
    ["Reviewers: 'incident codings are one analyst's opinion'", "High without E5", "E5 agreement statistics; disputed codings marked and excluded from counts"],
    ["Reviewers: 'this is prompt engineering, not evolution'", "Moderate", "Section 3: theta is any policy update the gate can score; the catalog-constrained synthesizer is literally a model proposing its own instruction updates"],
    ["Anonymity leak via repo name, decks, run metadata", "Moderate", "E6 checklist; grep the artifact for name, handle, machine paths, provider account strings"],
    ["Over-hedged prose reads as no contribution", "High given current docs", "One Limitations section; body states what was built and observed in plain declaratives"],
], [Inches(3.6), Inches(2.7), Inches(6.03)], size=10.5)
strip(s, Inches(5.6), "Minimum viable submission, achievable with zero new inference: Sections 1-4, 6-8 plus E1 as the pilot evidence. "
      "E2 is what turns it from a system paper into a system paper with a first measured contrast.", BLUE, WHITE, bold=True, h=Inches(0.8), size=13)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides))
