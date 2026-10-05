"""Restyle v2 without rewriting its content.

Run: python3 scripts/build_incident_optimization_crayon_deck.py
The source builder is recorded, never run with its save operation. Text, notes
and original image bytes are checked against v2. All new text and decorations
are editable PowerPoint objects; figures keep their original image data.
"""

from __future__ import annotations

import collections
import hashlib
import inspect
import json
import math
import random
import re
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "scripts/build_incident_optimization_deck.py"
ORIGINAL = ROOT / "slides/incident-aware-optimization-v2.pptx"
OUT = ROOT / "slides/incident-aware-optimization-crayon-v3.pptx"
AUDIT = ROOT / "slides/incident-aware-optimization-crayon-v3.validation.json"

PAPER = "FBF8F0"
INK = "293E46"
MUTED = "68736E"
BLUE = ("4E7892", "EAF1F4")
SAGE = ("557B65", "EEF3E8")
CORAL = ("AE6B55", "FAECE4")
GOLD = ("9D7B35", "FBF1D5")
LILAC = ("7C7196", "F0EDF6")
PALE = "DDD8CA"
COLORS = [BLUE, SAGE, GOLD, CORAL, LILAC]
FONT = "Trebuchet MS"
HAND = "Chalkboard SE"
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
W, H = 16.0, 9.0
X, CW = 0.7, 14.6
FOOTER = "Incident-aware optimization · design proposal · containment-extension · October 2026 · no benchmark improvement measured"


@dataclass
class SourceSlide:
    title: str = ""
    subtitle: str | None = None
    section: str | None = None
    events: list = field(default_factory=list)
    note: str = ""

    def of(self, kind):
        return [v for k, v in self.events if k == kind]


def record_source():
    code = SOURCE.read_text()
    prefix, body = code.split("# ===================================================================== 1 title", 1)
    body = body.split("OUT.parent.mkdir", 1)[0]
    ns = {"__file__": str(SOURCE), "__name__": "recorded_deck"}
    exec(compile(prefix, str(SOURCE), "exec"), ns)
    slides = []

    def new_slide(title="", subtitle=None, section=None):
        slide = SourceSlide(title, subtitle, section)
        slides.append(slide)
        return slide

    def recorder(kind, signature):
        def capture(*args, **kwargs):
            bound = signature.bind(*args, **kwargs)
            bound.apply_defaults()
            values = dict(bound.arguments)
            slide = values.pop("slide")
            slide.events.append((kind, values))
        return capture

    for name in ["textbox", "bullets", "big_bullets", "rect", "card", "arrow_right", "table", "picture", "band"]:
        ns[name] = recorder(name, inspect.signature(ns[name]))
    ns["notes"] = lambda slide, text: setattr(slide, "note", text)
    ns["new_slide"] = new_slide
    ns["prs"] = SimpleNamespace(slides=SimpleNamespace(add_slide=lambda _: new_slide()))
    exec(compile(body, str(SOURCE), "exec"), ns)
    return slides


def clean(text):
    text = str(text).replace("**", "")
    return re.sub(r"\s+", " ", text).strip()


def tokens(text):
    # Whitespace and bullet markers belong to layout, not the source wording.
    return re.findall(r"\w+|[^\w\s•]", clean(text), flags=re.UNICODE)


def source_strings(source):
    strings = []
    for kind, ev in source.events:
        if kind in ("textbox", "band"):
            strings.append(ev["text"])
        elif kind in ("bullets", "big_bullets"):
            strings.extend(i if isinstance(i, str) else i[0] for i in ev["items"])
        elif kind == "card":
            strings.append(ev["heading"])
            strings.extend(ev["body"] if isinstance(ev["body"], list) else [ev["body"]])
            if ev["badge"] is not None:
                strings.append(str(ev["badge"]))
        elif kind == "table":
            strings.extend(str(cell) for row in ev["rows"] for cell in row)
    return strings


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


FONT_CACHE = {}


def font_file(family, bold):
    if family == HAND:
        # The collection contains Light, Regular and Bold.
        return FONT_DIR / "ChalkboardSE.ttc", 2 if bold else 1
    if family == "Arial Unicode MS":
        return FONT_DIR / "Arial Unicode.ttf", 0
    return FONT_DIR / ("Trebuchet MS Bold.ttf" if bold else "Trebuchet MS.ttf"), 0


def measured_font(size, family=FONT, bold=False):
    key = (size, family, bold)
    if key not in FONT_CACHE:
        path, index = font_file(family, bold)
        FONT_CACHE[key] = ImageFont.truetype(str(path), round(size * 4), index=index)
    return FONT_CACHE[key]


def wrap(text, width, size, family=FONT, bold=False):
    f = measured_font(size, family, bold)
    max_width = max(1, (width * 72 - 4) * 4 / 1.035)
    lines = []
    for paragraph in str(text).replace("**", "").split("\n"):
        current = ""
        for word in paragraph.split():
            test = (current + " " + word).strip()
            if current and f.getlength(test) > max_width:
                lines.append(current)
                current = word
            else:
                current = test
        lines.append(current)
    return lines


def text_height(lines, size, family=FONT):
    return len(lines) * size * (1.36 if family == HAND else 1.27) / 72 + 0.075


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.source_index = 0
        self.pages = []
        self.ledger = collections.defaultdict(list)
        self.images = collections.defaultdict(list)
        self.boxes = []
        self.sizes = []
        self.rng = random.Random(20261005)

    def path(self, points, color, width=1.2, closed=False, opacity=1.0):
        # PowerPoint/Keynote importers require non-degenerate custom geometry.
        points = list(points)
        if max(x for x, _ in points) - min(x for x, _ in points) < .0001:
            points[-1] = (points[-1][0] + .0002, points[-1][1])
        if max(y for _, y in points) - min(y for _, y in points) < .0001:
            points[-1] = (points[-1][0], points[-1][1] + .0002)
        scale = 10000
        builder = self.slide.shapes.build_freeform(
            round(points[0][0] * scale), round(points[0][1] * scale), scale=Inches(1) / scale
        )
        builder.add_line_segments([(round(x * scale), round(y * scale)) for x, y in points[1:]], close=closed)
        shape = builder.convert_to_shape()
        shape.fill.background()
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
        shape.line._get_or_add_ln().set("cap", "rnd")
        if opacity < 1:
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(round(opacity * 100000)))
            shape.line._get_or_add_ln().find(".//" + "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr").append(alpha)
        return shape

    def shape(self, x, y, w, h, fill, edge=None, oval=False, sketch=True):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.ROUNDED_RECTANGLE,
                                            Inches(x), Inches(y), Inches(w), Inches(h))
        if not oval:
            shape.adjustments[0] = 0.18
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if edge:
            if sketch:
                self.outline(x, y, w, h, edge, oval)
            else:
                shape.line.color.rgb = rgb(edge)
                shape.line.width = Pt(0.6)
        return shape

    def outline(self, x, y, w, h, color, oval=False):
        # Two imperfect pencil/crayon passes, plus sparse grain on the edge.
        exponent = 2.0 if oval else 5.0
        base = []
        for i in range(121):
            theta = 2 * math.pi * i / 120
            c, s = math.cos(theta), math.sin(theta)
            bx = math.copysign(abs(c) ** (2 / exponent), c)
            by = math.copysign(abs(s) ** (2 / exponent), s)
            base.append((x + w / 2 + (w / 2 - .012) * bx,
                         y + h / 2 + (h / 2 - .012) * by))
        for pass_index in range(2):
            points = [(px + self.rng.uniform(-.013, .013), py + self.rng.uniform(-.012, .012)) for px, py in base]
            self.path(points, color, 1.55 if pass_index == 0 else .6, closed=True,
                      opacity=.72 if pass_index == 0 else .42)
        for start in range(4, 118, 15):
            self.path(base[start:start + 5], color, 2.2, opacity=.18)

    def stroke(self, x, y, w, color, thickness=3.0):
        points = [(x + w * i / 24, y + .018 * math.sin(i * .8) + self.rng.uniform(-.008, .008)) for i in range(25)]
        self.path(points, color, thickness, opacity=.44)
        self.path([(px, py + .018) for px, py in points], color, .8, opacity=.58)

    def sparkle(self, x, y, color=GOLD[0], radius=.12):
        for dx, dy in [(radius, 0), (0, radius), (.065, .065), (.065, -.065)]:
            self.path([(x-dx, y-dy), (x+dx, y+dy)], color, 1.3, opacity=.7)

    def text(self, text, x, y, w, h, size=21, color=INK, bold=False, family=FONT,
             align="center", min_size=14, role="body", valign="middle"):
        requested = size
        while True:
            lines = wrap(text, w, size, family, bold)
            need = text_height(lines, size, family)
            longest = max((measured_font(size, family, bold).getlength(line) / 4 / 72 for line in lines), default=0)
            if need <= h and longest <= w - .025:
                break
            size -= .25
            if size < min_size:
                raise ValueError(f"Slide {len(self.pages)}: text will not fit: {text!r} in {w} x {h}; minimum {min_size}")
        actual_y = y + (h - need) / 2 if valign == "middle" else y
        if family == HAND and role == "body" and len(lines) > 1:
            # Chalkboard's visible glyphs sit high within its line metrics.
            # Give two-line headings optical centering in PowerPoint.
            actual_y += .06
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(actual_y), Inches(w), Inches(need))
        tb.name = f"{role}: {clean(text)[:64]}"
        tf = tb.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.text = "\v".join(lines)
        p.alignment = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = Pt(size * (1.36 if family == HAND else 1.27))
        p.space_before = p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = family
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
        if role == "body":
            self.ledger[self.source_index].append(str(text))
        self.boxes.append({"page": len(self.pages), "role": role, "text": clean(text),
                           "x": x, "y": actual_y, "w": w, "h": need})
        self.sizes.append({"page": len(self.pages), "role": role, "size": size, "requested": requested})
        return tb

    def bullets(self, items, x, y, w, h, size=21, color=INK, gap=.15, min_size=16):
        texts = [i if isinstance(i, str) else i[0] for i in items]
        while True:
            heights = [text_height(wrap(t, w-.3, size), size) for t in texts]
            total = sum(heights) + max(0, len(texts)-1) * gap
            if total <= h:
                break
            size -= .25
            if size < min_size:
                raise ValueError(f"Slide {len(self.pages)} bullets overflow: {texts}")
        cy = y + (h-total)/2
        for t, height in zip(texts, heights):
            self.shape(x+.015, cy+.11, .068, .068, BLUE[0], oval=True, sketch=False)
            self.text(t, x+.27, cy, w-.27, height, size, color, align="left", min_size=min_size, valign="top")
            cy += height + gap

    def card(self, ev, x, y, w, h, palette=BLUE, size=20, heading_size=21,
             centered=False, badge=None):
        edge, pale = palette
        self.shape(x, y, w, h, pale, edge)
        heading = ev["heading"]
        badge = ev.get("badge") if badge is None else badge
        compact = h < 3.3
        head_h = .76 if len(heading) < 33 else 1.03
        if compact:
            head_h = .52 if len(heading) < 29 else .73
            heading_size = min(heading_size, 18.5)
        header_x = x + .23
        header_w = w - .46
        if badge is not None:
            # The index has its own margin; it never shares a text area.
            self.text(str(badge), x+.18, y+.2, .33, .31, 13, edge, bold=True, min_size=11)
            header_x += .36
            header_w -= .43
        if compact:
            head_h = max(.52, text_height(wrap(heading, header_w-.32, heading_size, HAND, True),
                                          heading_size, HAND) + .06)
        head_y = y + (.10 if compact else .16)
        self.text(heading, header_x+.16, head_y+.03, header_w-.32, head_h-.06,
                  heading_size, edge, bold=True, family=HAND, min_size=15)
        body_y = y + head_h + (.22 if compact else .32)
        body_h = y + h - (.12 if compact else .22) - body_y
        body = ev["body"]
        if isinstance(body, list):
            self.bullets(body, x+.28, body_y, w-.56, body_h, size, gap=.08 if compact else .13)
        else:
            self.text(body, x+.28, body_y, w-.56, body_h, size,
                      align="center" if centered else "left", min_size=16)

    def callout(self, text, x, y, w, h, palette=GOLD, size=23, oval=False, family=FONT, bold=False):
        self.shape(x, y, w, h, palette[1], palette[0], oval=oval)
        pad = .12*w if oval else .3
        self.text(text, x+pad, y+.15, w-2*pad, h-.30, size, INK, bold=bold, family=family, min_size=17)

    def table(self, ev, x, y, w, h, widths=None, size=20, rows=None, repeat_header=False):
        rows = rows or ev["rows"]
        weights = widths or ev["col_widths"] or [1] * len(rows[0])
        widths = [w * z / sum(weights) for z in weights]
        while True:
            heights = []
            for ri, row in enumerate(rows):
                sizes = [size if ri else min(size, 19)]*len(row)
                height = max(text_height(wrap(str(t), widths[ci]-.3, sizes[ci], FONT, ri == 0 or ci == 0), sizes[ci])
                             for ci, t in enumerate(row)) + (.19 if ri else .17)
                heights.append(max(.43 if ri else .51, height))
            if sum(heights) + .065*(len(rows)-1) <= h:
                break
            size -= .25
            if size < 15:
                raise ValueError(f"Slide {len(self.pages)} table too dense at {size}: {h}")
        extra = (h - sum(heights) - .065*(len(rows)-1)) / len(rows)
        heights = [v + extra for v in heights]
        cy = y
        for ri, (row, row_h) in enumerate(zip(rows, heights)):
            self.shape(x, cy, w, row_h, BLUE[1] if ri == 0 else ("F0F2E9" if ri%2 else "FFFDF7"))
            cx = x
            for ci, value in enumerate(row):
                role = "repeated-table-header" if ri == 0 and repeat_header else "body"
                align = "center" if ri == 0 or (ci == 0 and widths[ci] < 1.7) else "left"
                self.text(str(value), cx+.15, cy+.07, widths[ci]-.3, row_h-.14,
                          min(size, 19) if ri == 0 else size,
                          BLUE[0] if ri == 0 or ci == 0 else INK, bold=ri == 0 or ci == 0,
                          align=align, min_size=15, role=role)
                cx += widths[ci]
            cy += row_h + .065

    def picture(self, ev, x, y, w, h):
        path = Path(ev["path"])
        with Image.open(path) as image:
            ratio = image.width / image.height
        self.shape(x-.08, y-.08, w+.16, h+.16, "FFFFFF", PALE)
        ph = min(h, w/ratio)
        pw = ph*ratio
        self.slide.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), Inches(pw), Inches(ph))
        self.images[self.source_index].append(hashlib.sha256(path.read_bytes()).hexdigest())
        self.boxes.append({"page": len(self.pages), "role": "picture", "text": str(path.name),
                           "x": x, "y": y, "w": w, "h": h})

    def arrow(self, x1, y1, x2, y2, color=BLUE[0]):
        self.path([(x1,y1), ((x1+x2)/2,(y1+y2)/2-.013), (x2,y2)], color, 1.75, opacity=.8)
        angle = math.atan2(y2-y1,x2-x1)
        for delta in [-.55,.55]:
            self.path([(x2-.105*math.cos(angle+delta),y2-.105*math.sin(angle+delta)),(x2,y2)],color,1.5)

    def start(self, source, index, part=None):
        self.source_index = index
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.pages.append({"page":len(self.pages)+1,"source_slide":index+1,"part":part})
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(PAPER)
        if source.note:
            self.slide.notes_slide.notes_text_frame.text = source.note
        if source.title:
            if source.section:
                self.text(source.section, 4.6, .18, 6.8, .36, 11, BLUE[0], bold=True,
                          role="section", min_size=10)
            self.text(source.title, .85, .63, 14.3, .64, 31, INK, bold=True, family=HAND,
                      role="title", min_size=24)
            if source.subtitle:
                self.text(source.subtitle, 1, 1.31, 14, .36, 15.5, MUTED,
                          role="subtitle", min_size=13)
            else:
                self.stroke(6.85, 1.51, 2.3, GOLD[0], 3.0)
            self.sparkle(15.30, .72, GOLD[0], .095)
            if part:
                self.text(f"{part[0]} / {part[1]}", 14.1, .2, 1.15, .32, 11, MUTED,
                          role="part", min_size=9)
            self.path([(.7,8.48),(8.0,8.49),(15.3,8.48)], PALE, .6)
            self.text(FOOTER, .75, 8.56, 13.65, .23, 8.0, MUTED, role="footer", min_size=7)
            self.text(str(len(self.pages)), 14.77, 8.55, .43, .27, 10, BLUE[0], role="page", min_size=9)
        return self.slide


def feasibility_slide(d, s, index):
    """Render the shared replay analysis without framed titles or dense pages."""
    cards, tables = s.of("card"), s.of("table")
    texts, bands = s.of("textbox"), s.of("band")
    d.start(s, index)
    if s.title == "Full incident replay starts with an access problem":
        for i, card in enumerate(cards):
            d.card(card, .7+i*4.98, 1.95, 4.64, 4.43, [BLUE, GOLD, CORAL][i], 20, 21)
        d.callout(bands[0]["text"], .7, 6.62, 14.6, 1.04, GOLD, 19)
        d.text(texts[0]["text"], .85, 7.84, 14.3, .41, 14.5, MUTED, min_size=14)
    elif s.title == "What the published OpenAI examples let us reproduce":
        d.table(tables[0], X, 1.95, CW, 4.83, [3.8, 5.05, 5.75], 20)
        d.callout(bands[0]["text"], .7, 7.05, 14.6, 1.08, BLUE, 20)
    elif s.title == "Three replay modes support different claims":
        for i, card in enumerate(cards):
            d.card(card, .7+i*4.98, 1.95, 4.64, 4.58, [BLUE, SAGE, GOLD][i], 20, 21)
        d.callout(bands[0]["text"], .7, 6.87, 14.6, 1.22, GOLD, 20)
    elif s.title == "The evaluator is part of the incident":
        d.table(tables[0], X, 1.95, CW, 6.17, [3.3, 5.6, 5.7], 20)
    elif s.title == "Where GPU, API and engineering costs become bottlenecks":
        d.table(tables[0], X, 1.95, CW, 4.34, [3.4, 5.6, 5.6], 19)
        d.text(texts[0]["text"], .85, 6.43, 14.3, .45, 18, MUTED, min_size=17)
        d.callout(bands[0]["text"], .7, 7.04, 14.6, 1.10, GOLD, 19)
    elif s.title == "When full replay is blocked, test the key mechanism":
        for i, card in enumerate(cards):
            d.card(card, .7+i*4.98, 1.95, 4.64, 4.62, [BLUE, GOLD, SAGE][i], 20, 21)
        d.callout(bands[0]["text"], .7, 6.85, 14.6, 1.28, GOLD, 20)
    elif s.title == "Choose the replay scope case by case":
        d.table(tables[0], X, 1.95, CW, 4.92, [4.7, 4.1, 5.8], 19)
        d.callout(bands[0]["text"], .7, 7.10, 14.6, 1.05, BLUE, 19)
    else:
        raise ValueError(f"No feasibility layout for {s.title}")


def build(sources):
    d = Deck()
    legacy_index = 0
    for index, s in enumerate(sources):
        if s.section == "7 · REPLAY FEASIBILITY":
            feasibility_slide(d, s, index)
            continue
        legacy_index += 1
        n = legacy_index
        cards, tables = s.of("card"), s.of("table")
        texts, bands = s.of("textbox"), s.of("band")
        if n == 5:
            data = tables[0]["rows"]
            chunks = [data[1:6], data[6:10], data[10:]]
            for part, chunk in enumerate(chunks,1):
                d.start(s,index,(part,3))
                d.table(tables[0],X,1.97,CW,6.12,[1.3,2.75,6.65,3.9],18.5,
                        [data[0]]+chunk,repeat_header=part>1)
            continue
        if n == 13:
            data=tables[0]["rows"]
            for part, chunk in enumerate([data[1:4],data[4:]],1):
                d.start(s,index,(part,2))
                d.table(tables[0],X,2.04,CW,5.95,[2.85,4.2,3.8,3.75],20,
                        [data[0]]+chunk,repeat_header=part>1)
            continue
        if n == 19:
            d.start(s,index,(1,2))
            d.callout(texts[0]["text"],1.06,2.03,13.88,1.51,BLUE,29,oval=True,
                      family="Arial Unicode MS",bold=True)
            d.card(cards[0],.7,4.04,7.12,3.84,BLUE,23)
            d.card(cards[1],8.18,4.04,7.12,3.84,CORAL,23)
            d.start(s,index,(2,2))
            d.table(tables[0],X,2.05,CW,5.90,[4.0,10.6],23)
            continue
        if n == 24:
            d.start(s,index,(1,2))
            d.table(tables[0],X,1.98,CW,6.12,[6.0,8.6],21)
            d.start(s,index,(2,2))
            d.card(cards[0],.7,2.67,7.12,4.47,BLUE,24)
            d.card(cards[1],8.18,2.67,7.12,4.47,SAGE,24)
            continue
        if n in (27,28):
            for part, ev in enumerate(s.of("bullets"),1):
                d.start(s,index,(part,2))
                items=ev["items"]
                row_h=6.2/len(items)
                for i,item in enumerate(items):
                    cy=1.94+i*row_h
                    d.shape(.9,cy+row_h/2-.07,.12,.12,COLORS[i%5][0],oval=True,sketch=False)
                    d.text(item,1.28,cy,13.7,row_h-.055,19.5 if n==27 else 20,
                           align="left",min_size=17)
                    if i<len(items)-1:
                        d.path([(1.28,cy+row_h-.018),(14.98,cy+row_h-.018)],PALE,.45)
            continue

        d.start(s,index)
        if n == 1:
            d.text(texts[0]["text"],2.05,1.55,11.9,.89,44,INK,bold=True,family=HAND,min_size=38)
            d.text(texts[1]["text"],2.80,2.60,10.40,1.2,24,BLUE[0],min_size=22)
            d.stroke(5.1,4.94,5.8,GOLD[0],5.0)
            d.text(texts[2]["text"],2.0,5.38,12,1.15,23,INK,min_size=20)
            d.text(texts[3]["text"],1.6,7.64,12.8,.5,15,MUTED,min_size=14)
            d.sparkle(1.05,1.10, GOLD[0],.19)
            d.sparkle(14.9,4.63,SAGE[0],.20)
            d.stroke(1.05,8.35,1.25,CORAL[0],4)
            d.stroke(13.70,.44,1.1,SAGE[0],4)
        elif n == 2:
            for i in range(6):
                num,head,body=texts[3*i:3*i+3]
                x=.7+(i%2)*7.48; y=1.97+(i//2)*2.1
                palette=COLORS[i%5]
                d.shape(x,y,7.12,1.80,palette[1],palette[0])
                d.text(num["text"],x+.3,y+.32,.46,.45,22,palette[0],bold=True,min_size=18)
                d.text(head["text"],x+1.0,y+.18,5.76,.51,23,palette[0],bold=True,family=HAND,min_size=20)
                d.text(body["text"],x+.95,y+.83,5.82,.84,20,align="center",min_size=18)
        elif n == 3:
            for i,card in enumerate(cards):
                x=.7;y=1.90+i*1.60
                palette=COLORS[i%5]
                d.shape(x,y,7.80,1.40,palette[1],palette[0])
                d.text(card["heading"],x+.23,y+.12,7.34,.45,19,palette[0],bold=True,family=HAND,min_size=17)
                d.text(card["body"],x+.30,y+.60,7.2,.75,18.5,min_size=17)
            d.table(tables[0],8.94,1.9,6.36,4.29,[1,5.1],17.5)
            d.callout(texts[0]["text"],8.94,6.49,6.36,1.58,GOLD,19)
        elif n == 4:
            d.table(tables[0],X,1.95,CW,4.65,[1.9,5.4,5.6],20)
            d.callout(bands[0]["text"],X,6.94,CW,1.19,BLUE,21)
        elif n == 6:
            d.card(cards[0],.7,1.93,7.12,4.55,CORAL,20)
            d.card(cards[1],8.18,1.93,7.12,4.55,BLUE,20)
            d.callout(bands[0]["text"],.7,6.80,14.6,1.32,GOLD,21)
        elif n == 7:
            d.callout(texts[0]["text"],.85,1.89,14.3,1.37,BLUE,27,bold=True)
            d.text(texts[1]["text"],1.05,3.44,13.9,.39,16.5,MUTED,min_size=16)
            for i,card in enumerate(cards):
                d.card(card,.7+i*4.98,4.1,4.64,4.04,[CORAL,SAGE,BLUE][i],19)
        elif n == 8:
            d.picture(s.of("picture")[0],.82,2.37,9.65,5.45)
            d.bullets(s.of("big_bullets")[0]["items"],10.97,1.98,4.30,6.05,22,min_size=18,gap=.30)
        elif n == 9:
            d.table(tables[0],X,2.05,CW,5.87,[5.1,7.2],23)
        elif n == 10:
            for i,card in enumerate(cards):
                d.card(card,.7+i*4.98,2.0,4.64,4.77,[BLUE,GOLD,SAGE][i],21)
            d.callout(bands[0]["text"],.7,7.10,14.6,1.03,BLUE,19)
        elif n == 11:
            d.table(tables[0],X,1.94,CW,2.93,[2.7,4.3,6.4],21)
            for i,card in enumerate(cards):
                d.card(card,.7+i*7.48,5.21,7.12,2.93,[BLUE,CORAL][i],20)
        elif n == 12:
            d.picture(s.of("picture")[0],2.9,1.89,10.20,6.32)
        elif n == 14:
            for i,card in enumerate(cards[:6]):
                d.card(card,.7+(i%3)*4.98,1.93+(i//3)*2.35,4.64,2.14,COLORS[i%5],18,18.5,True)
            card=cards[6]
            d.shape(.7,6.89,14.6,1.23,SAGE[1],SAGE[0])
            d.text(str(card["badge"]),1.00,7.26,.35,.37,16,SAGE[0],bold=True,min_size=14)
            d.text(card["heading"],1.7,7.00,5.0,.91,22,SAGE[0],bold=True,family=HAND,min_size=18)
            d.text(card["body"],7.0,7.00,7.8,.91,22,min_size=19)
        elif n == 15:
            for i,card in enumerate(cards[:3]):
                d.card(card,.7+i*5.07,1.94,4.46,3.84,[BLUE,GOLD,CORAL][i],19.5)
                if i<2:
                    d.arrow(5.27+i*5.07,3.78,5.64+i*5.07,3.78)
            d.card(cards[3],.7,5.96,7.12,2.17,SAGE,19)
            d.card(cards[4],8.18,5.96,7.12,2.17,CORAL,19)
        elif n == 16:
            d.card(cards[0],.7,1.95,7.12,6.18,BLUE,22)
            d.card(cards[1],8.18,1.95,7.12,2.88,GOLD,22)
            d.card(cards[2],8.18,5.12,7.12,3.01,SAGE,22)
        elif n == 17:
            for i in range(5):
                head,body=texts[2*i:2*i+2]
                cy=1.98+i*1.23
                palette=COLORS[i%5]
                d.shape(.7,cy,8.96,1.04,palette[1],palette[0])
                d.text(head["text"],.89,cy+.10,2.90,.84,19,palette[0],bold=True,min_size=17)
                d.text(body["text"],4.02,cy+.10,5.39,.84,19,align="left",min_size=17)
            d.card(cards[0],10.01,1.98,5.29,5.96,SAGE,22)
        elif n == 18:
            d.picture(s.of("picture")[0],.85,2.60,8.18,4.60)
            d.card(cards[0],9.49,1.95,5.81,2.93,BLUE,20)
            d.card(cards[1],9.49,5.19,5.81,2.94,SAGE,20)
        elif n == 20:
            for i,ev in enumerate(texts):
                x=.7+i*3.01
                palette=SAGE if i==4 else BLUE
                d.shape(x,1.95,2.56,1.43,palette[1],palette[0],oval=True)
                d.text(ev["text"],x+.28,2.19,2.00,.92,18.5,palette[0],bold=True,min_size=16)
                if i<4:
                    d.arrow(x+2.64,2.66,x+2.89,2.66)
            d.card(cards[0],.7,3.78,7.12,4.35,BLUE,20.5)
            d.card(cards[1],8.18,3.78,7.12,4.35,SAGE,20.5)
        elif n == 21:
            for i,ev in enumerate(texts):
                col=i if i<4 else 7-i
                x=.7+col*3.88; y=1.94 if i<4 else 3.48
                palette=CORAL if i in (5,6) else (SAGE if i==7 else BLUE)
                d.shape(x,y,2.96,1.11,palette[1],palette[0],oval=True)
                d.text(ev["text"],x+.28,y+.18,2.4,.75,21,palette[0],bold=True,family=HAND,min_size=18)
                if i<3:
                    d.arrow(x+3.09,y+.55,x+3.70,y+.55)
                elif 4<=i<7:
                    d.arrow(x-.15,y+.55,x-.75,y+.55)
            d.arrow(14.34,3.12,14.34,3.37)
            d.bullets(s.of("big_bullets")[0]["items"],1.1,5.03,13.8,3.07,22,gap=.2,min_size=20)
        elif n == 22:
            for i,card in enumerate(cards):
                d.card(card,.7+(i%4)*3.735,1.98+(i//4)*3.19,3.395,2.97,COLORS[i%5],19,18.5,True)
        elif n == 23:
            d.table(tables[0],X,1.98,CW,6.14,[4.5,9.0],22)
        elif n == 25:
            for i,card in enumerate(cards):
                d.card(card,.7+i*4.98,1.94,4.64,3.03,[SAGE,GOLD,BLUE][i],19,20)
            d.table(tables[0],X,5.15,CW,3.00,[1.2,10.0,3.4],19)
        elif n == 26:
            d.bullets(s.of("big_bullets")[0]["items"],.92,2.02,8.51,4.83,23,gap=.26,min_size=21)
            d.card(cards[0],9.92,2.02,5.38,4.83,SAGE,22)
            d.callout(bands[0]["text"],.7,7.17,14.6,.97,GOLD,20)
        else:
            raise ValueError(f"No layout for source slide {n}")
    return d


def check_original(sources):
    original = Presentation(ORIGINAL)
    assert len(original.slides) == len(sources)
    results=[]
    for i,(source,slide) in enumerate(zip(sources,original.slides)):
        skip=collections.Counter(t for t in [source.title,source.subtitle,source.section,FOOTER,str(i+1)] if t)
        actual=[]
        pictures=[]
        for shape in slide.shapes:
            if shape.has_table:
                actual.extend(cell.text for row in shape.table.rows for cell in row.cells)
            elif shape.has_text_frame and shape.text:
                if skip[shape.text]>0:
                    skip[shape.text]-=1
                    continue
                actual.append(shape.text)
            elif shape.shape_type == 13:
                pictures.append(hashlib.sha256(shape.image.blob).hexdigest())
        expected=collections.Counter(t for v in source_strings(source) for t in tokens(v))
        found=collections.Counter(t for v in actual for t in tokens(v))
        assert expected == found, (i+1,expected-found,found-expected)
        expected_pics=[hashlib.sha256(Path(ev["path"]).read_bytes()).hexdigest() for ev in source.of("picture")]
        assert pictures==expected_pics, f"Source image changed on slide {i+1}"
        note=slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        assert note==source.note, f"Source notes changed on slide {i+1}"
        results.append({"source_slide":i+1,"words_and_symbols":sum(expected.values()),
                        "original_matches_source":True})
    return results


def validate(d,sources):
    results=check_original(sources)
    for i,source in enumerate(sources):
        expected=collections.Counter(t for v in source_strings(source) for t in tokens(v))
        found=collections.Counter(t for v in d.ledger[i] for t in tokens(v))
        assert expected==found,(f"Restyled slide {i+1}",expected-found,found-expected)
        expected_pics=[hashlib.sha256(Path(ev["path"]).read_bytes()).hexdigest() for ev in source.of("picture")]
        assert d.images[i]==expected_pics
        results[i]["new_content_matches"]=True
        results[i]["images_identical"]=True
    collisions=[]
    for i,a in enumerate(d.boxes):
        assert a["x"]>=0 and a["y"]>=0 and a["x"]+a["w"]<=W+.005 and a["y"]+a["h"]<=H+.005, a
        for b in d.boxes[i+1:]:
            if a["page"]!=b["page"]:
                continue
            ix=min(a["x"]+a["w"],b["x"]+b["w"])-max(a["x"],b["x"])
            iy=min(a["y"]+a["h"],b["y"]+b["h"])-max(a["y"],b["y"])
            if ix>.01 and iy>.01:
                collisions.append({"page":a["page"],"a":a["text"],"b":b["text"]})
    assert not collisions, collisions
    return {"source":ORIGINAL.name,"output":OUT.name,"original_pages":len(sources),
            "new_pages":len(d.pages),"all_content_preserved":True,"text_or_figure_overlaps":collisions,
            "minimum_body_font_pt":min(v["size"] for v in d.sizes if v["role"]=="body"),
            "checks":results,"page_map":d.pages,
            "font_notes":"Chalkboard SE headings; Trebuchet MS body; Arial Unicode MS equation",
            "render_review":"Pending native PowerPoint PDF export and visual review"}


def main():
    sources=record_source()
    d=build(sources)
    report=validate(d,sources)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    d.prs.core_properties.title="Incident-Aware Optimization"
    d.prs.core_properties.subject="Crayon restyling of the original v2 deck; original content preserved"
    d.prs.core_properties.author="Jiawei Li"
    d.prs.save(OUT)
    # Validate the serialized text, including every line break written to PPTX.
    reopened=Presentation(OUT)
    for page,slide in zip(d.pages,reopened.slides):
        source=sources[page["source_slide"]-1]
        if source.note:
            assert slide.notes_slide.notes_text_frame.text==source.note
    written=collections.Counter(t for slide in reopened.slides for shape in slide.shapes
                                if shape.has_text_frame and shape.name.startswith("body: ")
                                for t in tokens(shape.text))
    expected=collections.Counter(t for source in sources for v in source_strings(source) for t in tokens(v))
    assert written==expected,(expected-written,written-expected)
    AUDIT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n")
    print(f"Saved {OUT.relative_to(ROOT)} ({len(reopened.slides)} slides)")
    print(f"All {len(sources)} source slides match the original PPTX; text, figures and notes preserved.")
    print("No text/figure overlaps or off-slide content in the geometry check.")
    print(f"Validation: {AUDIT.relative_to(ROOT)}")


if __name__=="__main__":
    main()
